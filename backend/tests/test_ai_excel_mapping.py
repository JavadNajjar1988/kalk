from openpyxl import Workbook
from fastapi import HTTPException
import pytest

from app.api.routes.data_import import _build_sheet_samples, _validate_ai_mapping
from app.services.excel_scenario_import import (
    inspect_source_table_layout,
    parse_excel_workbook,
    transform_workbook_with_mapping,
)


def _heterogeneous_workbook() -> Workbook:
    workbook = Workbook()
    scenario = workbook.active
    scenario.title = "اطلاعات عملیات"
    scenario.append(["عنوان عملیات", "شرح عملیات"])
    scenario.append(["عملیات آزمایشی", "شرح نمونه"])

    events = workbook.create_sheet("رخدادها")
    events.append(["نام رخداد", "زمان رخداد", "طول", "عرض", "یادداشت داخلی"])
    events.append(["شروع عملیات", "2026-09-07T10:00:00Z", 51.4, 35.7, "منتقل نشود"])
    return workbook


def _reviewed_mapping() -> dict:
    return {
        "sheetMappings": [
            {
                "sourceSheet": "اطلاعات عملیات",
                "targetSheet": "scenario",
                "confidence": 0.98,
            },
            {
                "sourceSheet": "رخدادها",
                "targetSheet": "events",
                "confidence": 0.91,
            },
        ],
        "columnMaps": [
            {
                "targetSheet": "scenario",
                "mappings": [
                    {"fromHeader": "عنوان عملیات", "toField": "name"},
                    {"fromHeader": "شرح عملیات", "toField": "description"},
                ],
            },
            {
                "targetSheet": "events",
                "mappings": [
                    {"fromHeader": "نام رخداد", "toField": "title"},
                    {"fromHeader": "زمان رخداد", "toField": "start_time"},
                    {"fromHeader": "طول", "toField": "lon"},
                    {"fromHeader": "عرض", "toField": "lat"},
                ],
            },
        ],
    }


def test_reviewed_mapping_is_validated_and_transformed() -> None:
    source = _heterogeneous_workbook()
    mapping = _validate_ai_mapping(_reviewed_mapping(), _build_sheet_samples(source))

    standardized = transform_workbook_with_mapping(source, mapping)
    content, errors = parse_excel_workbook(standardized)

    assert errors == []
    assert content["name"] == "عملیات آزمایشی"
    assert len(content["events"]) == 1
    assert content["events"][0]["title"] == "شروع عملیات"


@pytest.mark.parametrize(
    ("path", "value"),
    [
        (("sheetMappings", 0, "sourceSheet"), "برگه ساختگی"),
        (("columnMaps", 0, "mappings", 0, "toField"), "فیلد_ساختگی"),
    ],
)
def test_reviewed_mapping_rejects_invented_sources_and_fields(path, value) -> None:
    source = _heterogeneous_workbook()
    mapping = _reviewed_mapping()
    target = mapping
    for key in path[:-1]:
        target = target[key]
    target[path[-1]] = value

    with pytest.raises(HTTPException) as exc:
        _validate_ai_mapping(mapping, _build_sheet_samples(source))

    assert exc.value.status_code == 422


def test_detects_header_after_introductory_rows() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "رخدادهای عملیات"
    worksheet.append(["گزارش رخدادهای قرارگاه"])
    worksheet.append(["این دو ردیف جزو داده‌های جدول نیستند."])
    worksheet.append([])
    worksheet.append(["نام رخداد", "زمان رخداد", "طول", "عرض"])
    worksheet.append(["حرکت", "2026-09-07T10:00:00Z", 51.4, 35.7])

    layout = inspect_source_table_layout(worksheet)
    samples = _build_sheet_samples(workbook)

    assert layout["headerRow"] == 4
    assert layout["headerDepth"] == 1
    assert layout["dataStartRow"] == 5
    assert samples[0]["headerRow"] == 4
    assert samples[0]["sampleRows"][0][0] == "حرکت"


def test_transforms_merged_two_row_headers() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "رخدادهای عملیات"
    worksheet.append(["گزارش رخدادهای قرارگاه"])
    worksheet.append([])
    worksheet.append(["مشخصات رخداد", None, "موقعیت", None])
    worksheet.merge_cells("A3:B3")
    worksheet.merge_cells("C3:D3")
    worksheet.append(["نام", "زمان", "طول", "عرض"])
    worksheet.append(["حرکت", "2026-09-07T10:00:00Z", 51.4, 35.7])

    layout = inspect_source_table_layout(worksheet)
    assert layout["headerRow"] == 4
    assert layout["headerDepth"] == 2
    assert layout["headers"] == [
        "مشخصات رخداد / نام",
        "مشخصات رخداد / زمان",
        "موقعیت / طول",
        "موقعیت / عرض",
    ]

    mapping = {
        "sheetMappings": [
            {
                "sourceSheet": "رخدادهای عملیات",
                "targetSheet": "events",
                "confidence": 0.95,
            }
        ],
        "columnMaps": [
            {
                "targetSheet": "events",
                "mappings": [
                    {"fromHeader": "مشخصات رخداد / نام", "toField": "title"},
                    {"fromHeader": "مشخصات رخداد / زمان", "toField": "start_time"},
                    {"fromHeader": "موقعیت / طول", "toField": "lon"},
                    {"fromHeader": "موقعیت / عرض", "toField": "lat"},
                ],
            }
        ],
    }
    validated = _validate_ai_mapping(mapping, _build_sheet_samples(workbook))
    standardized = transform_workbook_with_mapping(workbook, validated)
    content, errors = parse_excel_workbook(standardized)

    assert not [error for error in errors if error["sheet"] == "حوادث"]
    assert len(content["events"]) == 1
    assert content["events"][0]["title"] == "حرکت"
