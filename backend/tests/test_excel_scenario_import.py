from openpyxl import Workbook, load_workbook
from datetime import datetime, timezone
import pytest
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.db.base import Base
from app.models.resource import Resource, ResourceMedia
from app.models.scenario import Scenario
from app.services.excel_scenario_import import (
    BUNDLED_COMPLETE_EXAMPLE,
    BUNDLED_GRAPHIC_TEMPLATE,
    build_template_workbook,
    parse_excel_workbook,
    parse_resources_workbook,
)
from app.services.resource_import import (
    link_scenario_content_to_resources,
    scenario_units_to_bulk_items,
    upsert_resource_import_items,
    workbook_resources_to_bulk_items,
)
from app.services.scenario_import import merge_scenario_content
from app.services.resource_usage import build_resource_usage_graph
from app.services.scenario_import_impact import build_scenario_import_impact
from app.services.unit_resource_reconciliation import (
    collect_unlinked_resource_occurrences,
    collect_unlinked_unit_occurrences,
    link_resource_occurrence,
    link_unit_occurrence,
)


def _workbook() -> Workbook:
    wb = Workbook()
    scenario = wb.active
    scenario.title = "سناریو"
    scenario.append(
        ["name", "description", "start_time", "time_zone", "symbology_standard"]
    )
    scenario.append(["آزمون", "", "2026-01-01T08:00:00+00:00", "Asia/Tehran", "app6"])

    units = wb.create_sheet("یگان")
    units.append(["id", "name", "parent_id", "side", "unit_type", "time", "lon", "lat"])
    units.append(
        ["u1", "گردان", "", "friend", "tank", "2026-01-01T08:00:00+00:00", 51.4, 35.7]
    )
    units.append(["u2", "گروهان", "u1", "friend", "infantry", "", "", ""])

    equipment = wb.create_sheet("تجهیزات")
    equipment.append(
        ["id", "name", "type", "quantity", "unit_id", "start_time", "lon", "lat"]
    )
    equipment.append(
        ["eq1", "تانک", "تی ۶۴", 2, "u1", "2026-01-01T09:00:00+00:00", 51.5, 35.8]
    )

    personnel = wb.create_sheet("پرسنل")
    personnel.append(
        ["id", "first_name", "last_name", "rank", "specialty", "national_id", "unit_id"]
    )
    personnel.append(["p1", "علی", "رضایی", "سروان", "فرمانده", "", "u1"])

    events = wb.create_sheet("حوادث")
    events.append(
        ["id", "title", "start_time", "side", "unit_ids", "equipment_ids", "lon", "lat"]
    )
    events.append(
        [
            "ev1",
            "حرکت",
            "2026-01-01T09:00:00+00:00",
            "friend",
            "u1,u2",
            "eq1",
            51.6,
            35.9,
        ]
    )

    features = wb.create_sheet("عوارض")
    features.append(
        [
            "id",
            "name",
            "type",
            "lon",
            "lat",
            "radius_m",
            "start_time",
            "end_time",
            "event_id",
        ]
    )
    features.append(["f1", "شهر نمونه", "شهر", 51.7, 36.0, 3000, "", "", "ev1"])
    return wb


def test_parser_builds_orbat_attachments_features_and_storyboard():
    content, errors = parse_excel_workbook(_workbook())

    assert errors == []
    assert content["sides"][0]["id"] == "excel-side-3"
    root = content["sides"][0]["groups"][0]["subUnits"][0]
    assert root["id"] == "u1"
    assert root["sidc"] == "10031000001205000000"
    assert root["subUnits"][0]["id"] == "u2"
    assert root["equipment"] == [
        {"name": "تانک - تی ۶۴", "count": 2, "description": "تی ۶۴"}
    ]
    assert root["personnel"][0]["name"] == "علی رضایی"
    assert content["equipment"][0]["name"] == "تانک - تی ۶۴"

    layers = {layer["id"]: layer for layer in content["layers"]}
    equipment_feature = layers["excel-independent-equipment"]["features"][0]
    assert equipment_feature["geometry"]["coordinates"] == [51.5, 35.8]
    assert equipment_feature["style"]["militarySymbolSidc"] == "10031500001202000000"
    city = layers["excel-imported-features"]["features"][0]
    assert city["meta"]["type"] == "Circle"
    assert city["meta"]["radius"] == 3000

    event = content["events"][0]
    assert event["involvedUnits"] == ["u1", "u2"]
    assert event["involvedEquipment"] == ["eq1"]
    assert content["storyboard"]["scenes"][0]["camera"]["type"] == "eventWhere"


def test_parser_adds_sorted_unit_movement_states_for_kalknegar_playback():
    workbook = _workbook()
    states = workbook.create_sheet("وضعیت‌های زمانی یگان‌ها")
    states.append(
        [
            "شناسه_وضعیت",
            "شناسه_یگان",
            "زمان_مقصد",
            "طول",
            "عرض",
            "روش_انتقال",
            "نوع_مسیر",
            "زمان_شروع_حرکت",
        ]
    )
    states.append(
        [
            "move-2",
            "u1",
            "2026-01-01T12:00:00+00:00",
            51.8,
            36.0,
            "پیوسته",
            "منحنی",
            "2026-01-01T10:00:00+00:00",
        ]
    )
    states.append(
        [
            "move-1",
            "u1",
            "2026-01-01T10:00:00+00:00",
            51.6,
            35.8,
            "پرش",
            "مستقیم",
            "",
        ]
    )

    content, errors = parse_excel_workbook(workbook)

    assert errors == []
    unit = content["sides"][0]["groups"][0]["subUnits"][0]
    assert [state["id"] for state in unit["state"]] == [
        "excel-state-u1-initial",
        "excel-state-u1-move-1",
        "excel-state-u1-move-2",
    ]
    assert unit["state"][1] == {
        "id": "excel-state-u1-move-1",
        "t": "2026-01-01T10:00:00+00:00",
        "location": [51.6, 35.8],
        "interpolate": False,
        "pathMode": "straight",
    }
    assert unit["state"][2]["interpolate"] is True
    assert unit["state"][2]["pathMode"] == "curved"
    assert unit["state"][2]["viaStartTime"] == "2026-01-01T10:00:00+00:00"


def test_parser_rejects_invalid_unit_movement_rows():
    workbook = _workbook()
    states = workbook.create_sheet("وضعیت‌های زمانی یگان‌ها")
    states.append(
        [
            "شناسه_وضعیت",
            "شناسه_یگان",
            "زمان_مقصد",
            "طول",
            "عرض",
            "روش_انتقال",
            "نوع_مسیر",
            "زمان_شروع_حرکت",
        ]
    )
    states.append(
        [
            "bad-1",
            "missing-unit",
            "2026-01-01T10:00:00+00:00",
            51.6,
            35.8,
            "پیوسته",
            "مستقیم",
            "",
        ]
    )
    states.append(
        [
            "bad-2",
            "u1",
            "2026-01-01T10:00:00+00:00",
            51.6,
            35.8,
            "ناشناخته",
            "مستقیم",
            "",
        ]
    )

    _, errors = parse_excel_workbook(workbook)

    assert [error["message"] for error in errors] == [
        "یگان با شناسه missing-unit یافت نشد",
        "روش انتقال باید «پیوسته» یا «پرش» باشد",
    ]


def test_merge_updates_matching_ids_and_preserves_absent_records():
    existing = {
        "id": "scenario-1",
        "name": "قدیمی",
        "events": [
            {"id": "ev-old", "title": "ماندگار"},
            {"id": "ev1", "title": "قدیمی"},
        ],
        "sides": [],
        "layers": [],
        "equipment": [],
        "personnel": [],
        "mapLayers": [{"id": "map-1"}],
        "storyboard": {"scenes": [], "settings": {}},
    }
    incoming = {
        "id": "temporary",
        "name": "جدید",
        "events": [{"id": "ev1", "title": "به‌روز"}],
        "sides": [],
        "layers": [],
        "equipment": [],
        "personnel": [],
        "mapLayers": [],
        "storyboard": {"scenes": [], "settings": {}},
    }

    merged = merge_scenario_content(existing, incoming)

    assert merged["id"] == "scenario-1"
    assert [event["id"] for event in merged["events"]] == ["ev-old", "ev1"]
    assert merged["events"][1]["title"] == "به‌روز"
    assert merged["mapLayers"] == [{"id": "map-1"}]


def test_merge_preserves_existing_unit_states_and_upserts_matching_state_ids():
    existing = {
        "id": "scenario-1",
        "sides": [
            {
                "id": "excel-side-3",
                "standardIdentity": "3",
                "groups": [
                    {
                        "id": "excel-group-3",
                        "subUnits": [
                            {
                                "id": "u1",
                                "state": [
                                    {"id": "old", "t": "2026-01-01T09:00:00Z"},
                                    {"id": "shared", "t": "2026-01-01T10:00:00Z"},
                                ],
                                "subUnits": [],
                            }
                        ],
                    }
                ],
            }
        ],
    }
    incoming = {
        "id": "temporary",
        "sides": [
            {
                "id": "excel-side-3",
                "standardIdentity": "3",
                "groups": [
                    {
                        "id": "excel-group-3",
                        "subUnits": [
                            {
                                "id": "u1",
                                "state": [
                                    {"id": "shared", "t": "2026-01-01T11:00:00Z"},
                                    {"id": "new", "t": "2026-01-01T12:00:00Z"},
                                ],
                                "subUnits": [],
                            }
                        ],
                    }
                ],
            }
        ],
    }

    merged = merge_scenario_content(existing, incoming)
    states = merged["sides"][0]["groups"][0]["subUnits"][0]["state"]

    assert [state["id"] for state in states] == ["old", "shared", "new"]
    assert states[1]["t"] == "2026-01-01T11:00:00Z"


def test_equipment_symbols_use_explicit_inferred_and_advanced_identities():
    workbook = _workbook()
    equipment = workbook["تجهیزات"]
    equipment.append(["eq2", "بلدزر", "دی ۸", 1, "", "", 51.51, 35.81])
    equipment.append(["eq3", "سنگر تانک", "سنگر تانک", 1, "", "", 51.52, 35.82])
    equipment.append(
        [
            "eq4",
            "هواپیما",
            "هواپیما",
            1,
            "",
            "",
            51.53,
            35.83,
            "دشمن",
            "10040100001101000000",
        ]
    )
    equipment.cell(1, 9, "side")
    equipment.cell(1, 10, "advanced_sidc")

    content, errors = parse_excel_workbook(workbook)

    assert errors == []
    features = {
        feature["properties"]["equipmentId"]: feature
        for layer in content["layers"]
        if layer["id"] == "excel-independent-equipment"
        for feature in layer["features"]
    }
    assert features["eq2"]["style"]["militarySymbolSidc"] == "10031500001311000000"
    assert "militarySymbolSidc" not in features["eq3"]["style"]
    assert features["eq4"]["style"]["militarySymbolSidc"] == "10040100001101000000"


def test_resource_ids_link_scenario_catalog_units_and_map_features():
    workbook = _workbook()
    content, _ = parse_excel_workbook(workbook)
    personnel, equipment, _, _ = parse_resources_workbook(workbook)

    link_scenario_content_to_resources(
        content,
        personnel,
        equipment,
        {
            "equipment": {"eq1": "equipment-resource-1"},
            "personnel": {"p1": "personnel-resource-1"},
        },
    )

    root = content["sides"][0]["groups"][0]["subUnits"][0]
    assert content["equipment"][0]["resourceId"] == "equipment-resource-1"
    assert content["personnel"][0]["resourceId"] == "personnel-resource-1"
    assert root["equipment"][0]["resourceId"] == "equipment-resource-1"
    assert root["personnel"][0]["resourceId"] == "personnel-resource-1"
    equipment_feature = next(
        layer
        for layer in content["layers"]
        if layer["id"] == "excel-independent-equipment"
    )["features"][0]
    assert equipment_feature["properties"]["resourceId"] == "equipment-resource-1"


def test_scenario_units_become_linked_catalog_resources():
    content, _ = parse_excel_workbook(_workbook())
    items = scenario_units_to_bulk_items(content)

    assert {item.code for item in items} == {"u1", "u2"}
    assert all(item.type == "units" for item in items)

    link_scenario_content_to_resources(
        content,
        [],
        [],
        {"equipment": {}, "personnel": {}, "units": {"u1": "unit-resource-1"}},
    )
    root = content["sides"][0]["groups"][0]["subUnits"][0]
    assert root["linkedResourceId"] == "unit-resource-1"


def test_usage_graph_reports_operations_assignments_and_statuses():
    equipment = Resource(
        id="equipment-resource-1",
        type="equipment",
        name="تانک تی-۷۲",
        code="EQ-72",
        status="available",
    )
    scenario = Scenario(
        id="operation-1",
        name="عملیات نمونه",
        created=datetime.now(timezone.utc),
        modified=datetime.now(timezone.utc),
        content={
            "status": "active",
            "startTime": "2026-01-01T08:00:00+00:00",
            "sides": [
                {
                    "name": "خودی",
                    "groups": [
                        {
                            "subUnits": [
                                {
                                    "id": "u1",
                                    "name": "گردان زرهی",
                                    "status": "درگیر",
                                    "equipment": [
                                        {
                                            "name": "تانک تی-۷۲",
                                            "count": 12,
                                            "onHand": 10,
                                            "resourceId": "equipment-resource-1",
                                        }
                                    ],
                                    "subUnits": [],
                                }
                            ]
                        }
                    ],
                }
            ],
            "layers": [],
        },
    )

    graph = build_resource_usage_graph(equipment, [scenario])

    assert graph["summary"] == {"operationsCount": 1, "assignmentsCount": 1}
    usage = graph["operations"][0]
    assert usage["scenarioName"] == "عملیات نمونه"
    assert usage["scenarioStatus"] == "active"
    assert usage["assignments"][0]["unitName"] == "گردان زرهی"
    assert usage["assignments"][0]["status"] == "درگیر"
    assert usage["assignments"][0]["quantity"] == 12


def test_usage_graph_reports_person_operational_participation_details():
    person = Resource(
        id="person-resource-1",
        type="personnel",
        name="علی صیاد شیرازی",
        code="PER-0001",
        status="active",
    )
    scenario = Scenario(
        id="operation-person-1",
        name="عملیات شخص نمونه",
        created=datetime.now(timezone.utc),
        modified=datetime.now(timezone.utc),
        content={
            "sides": [
                {
                    "name": "خودی",
                    "groups": [
                        {
                            "subUnits": [
                                {
                                    "id": "u1",
                                    "name": "قرارگاه عملیاتی",
                                    "personnel": [
                                        {
                                            "name": "علی صیاد شیرازی",
                                            "count": 1,
                                            "resourceId": "person-resource-1",
                                            "participationStatus": "active",
                                            "operationalRole": "فرمانده عملیات",
                                            "participationStartTime": 100,
                                            "participationEndTime": 200,
                                            "participationNotes": "هدایت محور جنوبی",
                                            "sourceReference": "گزارش روزانه، صفحه ۳۵",
                                        }
                                    ],
                                    "subUnits": [],
                                }
                            ]
                        }
                    ],
                }
            ],
            "layers": [],
        },
    )

    assignment = build_resource_usage_graph(person, [scenario])["operations"][0][
        "assignments"
    ][0]
    assert assignment["operationalRole"] == "فرمانده عملیات"
    assert assignment["status"] == "فعال در عملیات"
    assert assignment["startTime"] == 100
    assert assignment["endTime"] == 200
    assert assignment["notes"] == "هدایت محور جنوبی"
    assert assignment["sourceReference"] == "گزارش روزانه، صفحه ۳۵"


def test_usage_graph_reports_linked_logistics_supply():
    supply = Resource(
        id="logistics-resource-1",
        type="logistics",
        name="سوخت دیزل",
        code="LOG-0001",
        status="active",
    )
    scenario = Scenario(
        id="operation-supply-1",
        name="عملیات تدارکاتی نمونه",
        created=datetime.now(timezone.utc),
        modified=datetime.now(timezone.utc),
        content={
            "sides": [
                {
                    "name": "خودی",
                    "groups": [
                        {
                            "subUnits": [
                                {
                                    "id": "u1",
                                    "name": "گردان پشتیبانی",
                                    "supplies": [
                                        {
                                            "name": "سوخت دیزل",
                                            "count": 5000,
                                            "onHand": 4200,
                                            "resourceId": "logistics-resource-1",
                                        }
                                    ],
                                    "subUnits": [],
                                }
                            ]
                        }
                    ],
                }
            ],
            "layers": [],
        },
    )

    graph = build_resource_usage_graph(supply, [scenario])

    assert graph["summary"] == {"operationsCount": 1, "assignmentsCount": 1}
    assignment = graph["operations"][0]["assignments"][0]
    assert assignment["kind"] == "supply"
    assert assignment["unitName"] == "گردان پشتیبانی"
    assert assignment["quantity"] == 5000
    assert assignment["onHand"] == 4200


def test_usage_graph_reports_unit_operational_participation_details():
    unit_resource = Resource(
        id="unit-resource-1",
        type="units",
        name="گردان یک",
        code="UNIT-0001",
        status="active",
    )
    scenario = Scenario(
        id="operation-unit-1",
        name="عملیات طریق‌القدس",
        created=datetime.now(timezone.utc),
        modified=datetime.now(timezone.utc),
        content={
            "sides": [
                {
                    "name": "خودی",
                    "groups": [
                        {
                            "subUnits": [
                                {
                                    "id": "u1",
                                    "name": "گردان یک",
                                    "linkedResourceId": "unit-resource-1",
                                    "participationStatus": "completed",
                                    "operationalRole": "پدافند از محور شمالی",
                                    "participationStartTime": 100,
                                    "participationEndTime": 300,
                                    "participationNotes": "مأموریت با حفظ موضع پایان یافت",
                                    "sourceReference": "گزارش عملیات، صفحه ۱۲",
                                    "subUnits": [],
                                }
                            ]
                        }
                    ],
                }
            ],
            "layers": [],
        },
    )

    assignment = build_resource_usage_graph(unit_resource, [scenario])["operations"][
        0
    ]["assignments"][0]
    assert assignment["status"] == "پایان‌یافته"
    assert assignment["operationalRole"] == "پدافند از محور شمالی"
    assert assignment["startTime"] == 100
    assert assignment["endTime"] == 300
    assert assignment["notes"] == "مأموریت با حفظ موضع پایان یافت"
    assert assignment["sourceReference"] == "گزارش عملیات، صفحه ۱۲"


def test_download_template_includes_new_columns_and_features_sheet():
    wb = build_template_workbook()

    assert "عوارض" in wb.sheetnames
    assert "زمان‌بندی" in wb.sheetnames
    assert "وضعیت‌های زمانی یگان‌ها" in wb.sheetnames
    assert [cell.value for cell in wb["تجهیزات"][1]] == [
        "شناسه",
        "نام",
        "نوع",
        "تعداد",
        "شناسه_یگان",
        "شناسه_زمان_شروع",
        "طول",
        "عرض",
        "طرف",
        "کد_نماد_پیشرفته",
    ]
    assert "شناسه_یگان" in [cell.value for cell in wb["حوادث"][1]]
    assert [cell.value for cell in wb["یگان‌ها"][1]] == [
        "شناسه",
        "نام",
        "طرف",
        "رده_یگان",
        "نوع_نماد",
        "شناسه_والد",
        "وضعیت_نماد",
        "شناسه_زمان",
        "طول",
        "عرض",
        "کد_نماد_پیشرفته",
        "کد_مرجع_یگان",
    ]
    assert [cell.value for cell in wb["شناسنامه یگان‌ها"][1]] == [
        "کد_مرجع_یگان",
        "عنوان_کوتاه",
        "وابستگی_سازمانی",
        "کد_یگان_بالادست",
        "نام_یگان_بالادست",
        "استان",
        "شهر_استقرار",
        "نام_پادگان",
        "توانمندی‌ها",
        "استعداد_اسمی",
        "آمادگی_پایه",
        "محدوده_مسئولیت",
        "فرمانده",
        "اطلاعات_تماس",
        "تاریخ_تشکیل",
        "تاریخ_پایان_فعالیت",
        "توضیحات",
    ]
    assert len(wb["یگان‌ها"].data_validations.dataValidation) >= 4
    assert len(wb["تجهیزات"].data_validations.dataValidation) >= 2
    assert all(cell.value is None for cell in wb["سناریو"][2])
    assert all(cell.value is None for cell in wb["یگان‌ها"][2])
    assert all(cell.value is None for cell in wb["وضعیت‌های زمانی یگان‌ها"][2])
    echelon_values = [wb["فهرست‌های انتخاب"].cell(row, 8).value for row in range(2, 11)]
    assert echelon_values[:4] == ["ارتش", "سپاه", "لشکر", "تیپ"]
    assert wb["زمان‌بندی"]["H2"].value.startswith("=IFERROR")
    assert wb["زمان‌بندی"]["H3"].value.startswith("=IFERROR")
    assert [cell.value for cell in wb["وضعیت‌های زمانی یگان‌ها"][1]] == [
        "شناسه_وضعیت",
        "شناسه_یگان",
        "شناسه_زمان_مقصد",
        "طول",
        "عرض",
        "روش_انتقال",
        "نوع_مسیر",
        "شناسه_زمان_شروع_حرکت",
    ]
    assert len(wb["وضعیت‌های زمانی یگان‌ها"].data_validations.dataValidation) == 3

    example = build_template_workbook(include_example=True)
    assert example["یگان‌ها"][2][1].value == "لشکر ۱۶"
    assert example["یگان‌ها"][4][5].value == "u2"
    assert example["یگان‌ها"][5][2].value == "دشمن"
    assert example["راهنما"][10][2].value == "دارای داده نمونه"


def test_persian_symbol_selections_generate_sidc_and_advanced_code_overrides():
    wb = Workbook()
    scenario = wb.active
    scenario.title = "سناریو"
    scenario.append(["نام"])
    scenario.append(["آزمون نماد"])

    units = wb.create_sheet("یگان‌ها")
    units.append(
        ["شناسه", "نام", "طرف", "نوع_نماد", "رده_یگان", "وضعیت_نماد", "کد_نماد_پیشرفته"]
    )
    units.append(
        ["u1", "گردان زرهی", "دشمن", "زرهی", "گردان / اسکادران", "آسیب‌دیده", ""]
    )
    units.append(
        [
            "u2",
            "نماد ویژه",
            "خودی",
            "پیاده‌نظام",
            "دسته / جزء مستقل",
            "حاضر",
            "10031000181211000000",
        ]
    )

    content, errors = parse_excel_workbook(wb)

    assert errors == []
    hostile = next(side for side in content["sides"] if side["id"] == "excel-side-6")
    assert hostile["groups"][0]["subUnits"][0]["sidc"] == "10061030161205000000"
    friendly = next(side for side in content["sides"] if side["id"] == "excel-side-3")
    assert friendly["groups"][0]["subUnits"][0]["sidc"] == "10031000181211000000"


def test_persian_headers_and_jalali_time_reference_are_parsed():
    wb = Workbook()
    scenario = wb.active
    scenario.title = "سناریو"
    scenario.append(
        ["نام", "توضیحات", "شناسه_زمان_شروع", "ناحیه_زمانی", "استاندارد_نماد"]
    )
    scenario.append(["آزمون شمسی", "", "زمان-001", "Asia/Tehran", "app6"])

    times = wb.create_sheet("زمان‌بندی")
    times.append(
        [
            "شناسه_زمان",
            "سال_شمسی",
            "ماه_شمسی",
            "روز_شمسی",
            "ساعت",
            "دقیقه",
            "ناحیه_زمانی",
        ]
    )
    times.append(["زمان-001", 1405, 6, 7, 8, 15, "Asia/Tehran"])

    events = wb.create_sheet("حوادث")
    events.append(
        ["شناسه", "عنوان", "شناسه_زمان_شروع", "طرف", "طول_جغرافیایی", "عرض_جغرافیایی"]
    )
    events.append(["ev-1", "رویداد نمونه", "زمان-001", "خودی", 51.4, 35.7])

    content, errors = parse_excel_workbook(wb)

    assert errors == []
    assert content["startTime"] == "2026-08-29T08:15:00+03:30"
    assert content["events"][0]["startTime"] == "2026-08-29T08:15:00+03:30"


def test_invalid_jalali_day_is_rejected_and_leap_esfand_is_accepted():
    wb = Workbook()
    scenario = wb.active
    scenario.title = "سناریو"
    scenario.append(["نام", "شناسه_زمان_شروع"])
    scenario.append(["آزمون تاریخ شمسی", "زمان-معتبر"])

    times = wb.create_sheet("زمان‌بندی")
    times.append(
        [
            "شناسه_زمان",
            "سال_شمسی",
            "ماه_شمسی",
            "روز_شمسی",
            "ساعت",
            "دقیقه",
            "ناحیه_زمانی",
        ]
    )
    times.append(["زمان-معتبر", 1403, 12, 30, 8, 0, "Asia/Tehran"])
    times.append(["زمان-نامعتبر-اسفند", 1404, 12, 30, 8, 0, "Asia/Tehran"])
    times.append(["زمان-نامعتبر-ماه", 1404, 7, 31, 8, 0, "Asia/Tehran"])

    content, errors = parse_excel_workbook(wb)

    assert content["startTime"] == "2025-03-20T08:00:00+03:30"
    assert [error["row"] for error in errors if error["sheet"] == "زمان‌بندی"] == [3, 4]
    assert all(
        error["message"] == "اجزای تاریخ شمسی یا ساعت نامعتبر است" for error in errors
    )


def test_resource_parser_preserves_equipment_quantity_for_dashboard_persistence():
    personnel, equipment, units, errors = parse_resources_workbook(_workbook())

    assert errors == []
    assert len(personnel) == 1
    assert equipment[0]["quantity"] == 2
    assert {unit["unitCode"] for unit in units} == {"u1", "u2"}

    items = workbook_resources_to_bulk_items(personnel, equipment, units)
    equipment_item = next(item for item in items if item.type == "equipment")
    assert equipment_item.code == "eq1"
    assert equipment_item.metadata["quantity"] == 2
    personnel_item = next(item for item in items if item.type == "personnel")
    assert personnel_item.name == "علی رضایی"
    unit_item = next(
        item for item in items if item.type == "units" and item.code == "u2"
    )
    assert unit_item.name == "گروهان"
    assert unit_item.metadata["parentCode"] == "u1"


def test_rich_unit_profile_uses_stable_resource_code_and_structured_metadata():
    wb = Workbook()
    scenario = wb.active
    scenario.title = "سناریو"
    scenario.append(["نام"])
    scenario.append(["عملیات نمونه"])
    units = wb.create_sheet("یگان‌ها")
    units.append(
        [
            "شناسه",
            "نام",
            "طرف",
            "رده_یگان",
            "نوع_نماد",
            "شناسه_والد",
            "کد_مرجع_یگان",
        ]
    )
    units.append(
        [
            "scenario-u1",
            "تیپ یکم",
            "خودی",
            "تیپ",
            "زرهی",
            "",
            "UNIT-001",
        ]
    )
    profiles = wb.create_sheet("شناسنامه یگان‌ها")
    profiles.append(
        [
            "کد_مرجع_یگان",
            "عنوان_کوتاه",
            "وابستگی_سازمانی",
            "کد_یگان_بالادست",
            "نام_یگان_بالادست",
            "استان",
            "شهر_استقرار",
            "نام_پادگان",
            "توانمندی‌ها",
            "استعداد_اسمی",
            "آمادگی_پایه",
            "محدوده_مسئولیت",
            "فرمانده",
            "اطلاعات_تماس",
            "تاریخ_تشکیل",
            "تاریخ_پایان_فعالیت",
            "توضیحات",
        ]
    )
    profiles.append(
        [
            "UNIT-001",
            "تیپ ۱",
            "نیروی زمینی",
            "DIV-001",
            "لشکر نمونه",
            "تهران",
            "تهران",
            "پادگان نمونه",
            "زرهی، راکت‌انداز، ضدزره",
            "۳۵۰۰",
            "بالا",
            "محور شمال",
            "سرهنگ نمونه",
            "داخلی ۱۲۳",
            "1981-09-22",
            "",
            "یگان واکنش سریع",
        ]
    )

    content, errors = parse_excel_workbook(wb)
    assert errors == []
    unit = content["sides"][0]["groups"][0]["subUnits"][0]
    assert unit["id"] == "scenario-u1"
    assert unit["resourceCode"] == "UNIT-001"
    assert unit["capabilities"] == ["زرهی", "راکت‌انداز", "ضدزره"]
    assert unit["organizationalAffiliation"] == "نیروی زمینی"
    assert unit["parentCode"] == "DIV-001"
    assert unit["parentName"] == "لشکر نمونه"
    assert unit["formedOn"] == "1981-09-22"

    items = scenario_units_to_bulk_items(content)
    assert items[0].code == "UNIT-001"
    assert items[0].metadata["echelon"] == "تیپ"
    assert items[0].metadata["garrisonCity"] == "تهران"
    assert items[0].metadata["capabilities"] == ["زرهی", "راکت‌انداز", "ضدزره"]
    assert items[0].metadata["organizationalAffiliation"] == "نیروی زمینی"
    assert items[0].metadata["parentName"] == "لشکر نمونه"
    assert items[0].metadata["formedOn"] == "1981-09-22"

    _, _, parsed_units, resource_errors = parse_resources_workbook(wb)
    assert resource_errors == []
    assert parsed_units[0]["unitCode"] == "UNIT-001"
    assert parsed_units[0]["defaultReadiness"] == "بالا"
    assert parsed_units[0]["parentCode"] == "DIV-001"
    assert parsed_units[0]["organizationalAffiliation"] == "نیروی زمینی"


def test_scenario_import_impact_reports_merge_and_replace_consequences():
    existing = {
        "events": [{"id": "event-old"}, {"id": "event-shared"}],
        "equipment": [{"sourceId": "eq-old"}],
        "sides": [],
        "layers": [],
        "personnel": [],
        "storyboard": {"scenes": []},
    }
    incoming = {
        "events": [{"id": "event-new"}, {"id": "event-shared"}],
        "equipment": [],
        "sides": [],
        "layers": [],
        "personnel": [],
        "storyboard": {"scenes": []},
    }

    merged = build_scenario_import_impact(existing, incoming, "merge")
    replaced = build_scenario_import_impact(existing, incoming, "replace")

    assert merged["collections"]["events"] == {
        "current": 2,
        "incoming": 2,
        "added": 1,
        "updated": 1,
        "removed": 0,
        "preserved": 1,
        "result": 3,
    }
    assert replaced["collections"]["events"]["removed"] == 1
    assert replaced["collections"]["events"]["result"] == 2
    assert replaced["collections"]["equipmentCatalog"]["removed"] == 1


def test_legacy_unit_reconciliation_only_links_the_confirmed_occurrence():
    content = {
        "sides": [
            {
                "name": "خودی",
                "groups": [
                    {
                        "subUnits": [
                            {
                                "id": "unit-1",
                                "name": "تیپ یکم",
                                "sidc": "10031000220000000000",
                                "subUnits": [{"id": "unit-2", "name": "گردان یکم"}],
                            }
                        ]
                    }
                ],
            }
        ]
    }

    occurrences = collect_unlinked_unit_occurrences(
        "scenario-1", "عملیات نمونه", content
    )
    assert [item["unitId"] for item in occurrences] == ["unit-1", "unit-2"]

    assert link_unit_occurrence(content, "unit-2", "units-canonical", "گردان مرجع")
    child = content["sides"][0]["groups"][0]["subUnits"][0]["subUnits"][0]
    parent = content["sides"][0]["groups"][0]["subUnits"][0]
    assert child["linkedResourceId"] == "units-canonical"
    assert "linkedResourceId" not in parent


def test_legacy_equipment_reconciliation_groups_and_links_catalog_unit_and_map():
    content = {
        "equipment": [{"sourceId": "eq-1", "name": "تانک تی-۷۲"}],
        "sides": [
            {
                "name": "خودی",
                "groups": [
                    {
                        "subUnits": [
                            {
                                "id": "unit-1",
                                "name": "گردان زرهی",
                                "equipment": [
                                    {
                                        "sourceId": "eq-1",
                                        "name": "تانک تی-۷۲",
                                        "count": 12,
                                    }
                                ],
                                "subUnits": [],
                            }
                        ]
                    }
                ],
            }
        ],
        "layers": [
            {
                "features": [
                    {
                        "meta": {"name": "تانک تی-۷۲"},
                        "properties": {"equipmentId": "eq-1", "quantity": 2},
                    }
                ]
            }
        ],
    }

    rows = collect_unlinked_resource_occurrences(
        "scenario-1", "عملیات نمونه", content, "equipment"
    )
    assert len(rows) == 1
    assert rows[0]["occurrenceKey"] == "source:eq-1"
    assert rows[0]["occurrenceCount"] == 3
    assert rows[0]["unitNames"] == ["گردان زرهی", "موقعیت مستقل"]

    linked = link_resource_occurrence(
        content, "equipment", "source:eq-1", "equipment-canonical"
    )
    assert linked == 3
    assert content["equipment"][0]["resourceId"] == "equipment-canonical"
    assignment = content["sides"][0]["groups"][0]["subUnits"][0]["equipment"][0]
    assert assignment["resourceId"] == "equipment-canonical"
    assert (
        content["layers"][0]["features"][0]["properties"]["resourceId"]
        == "equipment-canonical"
    )


def test_legacy_personnel_reconciliation_keeps_different_names_separate():
    content = {
        "personnel": [{"name": "علی رضایی"}, {"name": "رضا احمدی"}],
        "sides": [
            {
                "groups": [
                    {
                        "subUnits": [
                            {
                                "id": "unit-1",
                                "name": "ستاد",
                                "personnel": [{"name": "علی رضایی", "count": 1}],
                                "subUnits": [],
                            }
                        ]
                    }
                ]
            }
        ],
    }
    rows = collect_unlinked_resource_occurrences(
        "scenario-1", "عملیات نمونه", content, "personnel"
    )
    assert {row["occurrenceKey"] for row in rows} == {
        "name:علی رضایی",
        "name:رضا احمدی",
    }
    linked = link_resource_occurrence(
        content, "personnel", "name:علی رضایی", "personnel-canonical"
    )
    assert linked == 2
    assert "resourceId" not in content["personnel"][1]


def test_bundled_blank_template_and_complete_example_are_distinct():
    assert BUNDLED_GRAPHIC_TEMPLATE.is_file()
    assert BUNDLED_COMPLETE_EXAMPLE.is_file()

    blank = load_workbook(BUNDLED_GRAPHIC_TEMPLATE, data_only=False)
    for sheet_name, max_column in [
        ("سناریو", 5),
        ("حوادث", 10),
        ("یگان‌ها", 11),
        ("تجهیزات", 10),
        ("پرسنل", 7),
        ("عوارض", 9),
        ("وضعیت‌های زمانی یگان‌ها", 8),
    ]:
        assert all(
            blank[sheet_name].cell(row, column).value is None
            for row in range(2, 501)
            for column in range(1, max_column + 1)
        )
    assert all(
        blank["زمان‌بندی"].cell(row, column).value is None
        for row in range(2, 501)
        for column in range(1, 8)
    )
    assert blank["راهنما"][17][2].value == "بدون داده نمونه"
    assert len(blank["وضعیت‌های زمانی یگان‌ها"].data_validations.dataValidation) == 5

    wb = load_workbook(BUNDLED_COMPLETE_EXAMPLE, data_only=True)
    content, errors = parse_excel_workbook(wb)
    personnel, equipment, _, resource_errors = parse_resources_workbook(wb)

    assert errors == []
    assert resource_errors == []
    assert len(content["events"]) == 49
    assert len(content["storyboard"]["scenes"]) == 49
    assert len(equipment) == 33
    assert sum(item["quantity"] for item in equipment) == 183
    assert len(personnel) == 1
    assert wb["راهنما"][17][2].value == "دارای داده نمونه"
    equipment_features = next(
        layer
        for layer in content["layers"]
        if layer["id"] == "excel-independent-equipment"
    )["features"]
    assert (
        sum("militarySymbolSidc" in feature["style"] for feature in equipment_features)
        == 31
    )
    assert (
        sum(
            "militarySymbolSidc" not in feature["style"]
            for feature in equipment_features
        )
        == 2
    )


@pytest.mark.asyncio
async def test_resource_upsert_persists_and_updates_quantity(tmp_path):
    engine = create_async_engine(
        f"sqlite+aiosqlite:///{(tmp_path / 'resources.db').as_posix()}"
    )
    sessions = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    async with engine.begin() as connection:
        await connection.run_sync(
            lambda sync_connection: Base.metadata.create_all(
                sync_connection,
                tables=[Resource.__table__, ResourceMedia.__table__],
            )
        )

    _, equipment, _, _ = parse_resources_workbook(_workbook())
    items = workbook_resources_to_bulk_items([], equipment)
    async with sessions() as session:
        first = await upsert_resource_import_items(session, items)
        assert first == {"created": 1, "updated": 0, "skipped": 0}

        items[0].metadata["quantity"] = 5
        second = await upsert_resource_import_items(session, items)
        assert second == {"created": 0, "updated": 1, "skipped": 0}

        result = await session.execute(select(Resource).where(Resource.code == "eq1"))
        stored = result.scalar_one()
        assert stored.metadata_["quantity"] == 5

    await engine.dispose()
