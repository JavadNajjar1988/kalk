import asyncio
from types import SimpleNamespace

import pytest

from app.services.document_extraction import (
    build_document_review_workbook,
    canonical_entity_name,
    merge_page_text,
    completion,
    document_model_status,
    normalize_native_text,
    normalize_proposal_semantics,
    pipeline_markdown,
    pipeline_ocr,
    prepare_ocr_image,
    proposal_entity_key,
    entity_draft_id,
    proposed_resource_reference_code,
    read_page,
    repetition_detected,
    validate_proposals,
    visual_page_analysis,
)

from PIL import Image, ImageDraw


def test_review_workbook_preserves_evidence_and_leaves_unknowns_blank():
    scenario = {
        'id': 'scenario-1', 'kind': 'scenario', 'name': 'عملیات نمونه',
        'evidence': 'عملیات نمونه آغاز شد', 'sourcePage': 2,
        'sourceMethod': 'native', 'documentId': 'doc-1', 'reviewStatus': 'accepted',
    }
    event = {
        'id': 'event-1', 'kind': 'event', 'name': 'عبور از رودخانه',
        'evidence': 'نیروها از رودخانه عبور کردند', 'sourcePage': 3,
        'sourceMethod': 'native', 'documentId': 'doc-1', 'reviewStatus': 'accepted',
    }
    rejected_unit = {
        'id': 'unit-1', 'kind': 'unit', 'name': 'تیپ نمونه',
        'evidence': 'تیپ نمونه در منطقه بود', 'sourcePage': 4,
        'sourceMethod': 'native', 'documentId': 'doc-1', 'reviewStatus': 'rejected',
    }
    workbook = build_document_review_workbook(
        filename='report.pdf', document_id='doc-1', page_count=4,
        main_scenario_id='scenario-1', items=[scenario, event, rejected_unit],
    )
    assert workbook['سناریو']['A2'].value == 'عملیات نمونه'
    assert workbook['سناریو']['C2'].value == ''
    assert workbook['حوادث']['B2'].value == 'عبور از رودخانه'
    assert workbook['حوادث']['D2'].value == ''
    assert workbook['یگان‌ها'].max_row == 1
    assert workbook['کالک‌یار']['D3'].value == event['evidence']
    assert workbook['کالک‌یار']['E3'].value == 3
    from app.services.excel_scenario_import import parse_excel_workbook
    _, errors = parse_excel_workbook(workbook)
    assert any('زمان شروع سناریو' in error['message'] for error in errors)
    assert any('زمان رویداد' in error['message'] for error in errors)


def test_review_workbook_requires_matching_accepted_scenario():
    item = {
        'id': 'scenario-1', 'kind': 'scenario', 'name': 'عملیات نمونه',
        'evidence': 'عملیات نمونه', 'sourcePage': 1,
        'sourceMethod': 'native', 'documentId': 'other-doc', 'reviewStatus': 'accepted',
    }
    with pytest.raises(ValueError, match='شناسه سند'):
        build_document_review_workbook(
            filename='report.pdf', document_id='doc-1', page_count=1,
            main_scenario_id='scenario-1', items=[item],
        )


def test_review_workbook_maps_event_fields_to_their_exact_columns():
    common = {
        'sourceMethod': 'native', 'documentId': 'doc-1',
        'reviewStatus': 'accepted',
    }
    scenario = {
        **common, 'id': 'scenario-1', 'kind': 'scenario',
        'name': 'عملیات نمونه', 'evidence': 'عملیات نمونه آغاز شد',
        'sourcePage': 1, 'startTime': '2026-01-01T08:00:00+03:30',
    }
    event = {
        **common, 'id': 'event-1', 'kind': 'event',
        'name': 'پیشروی یگان', 'evidence': 'یگان در محور پیشروی کرد',
        'sourcePage': 2, 'startTime': '2026-01-01T09:00:00+03:30',
        'side': 'خودی', 'unitId': 'unit-1', 'equipmentId': 'equipment-1',
        'longitude': 48.1234, 'latitude': 31.5678,
    }

    workbook = build_document_review_workbook(
        filename='report.pdf', document_id='doc-1', page_count=2,
        main_scenario_id='scenario-1', items=[scenario, event],
    )

    row = [workbook['حوادث'].cell(2, column).value for column in range(1, 11)]
    assert row[5:] == ['خودی', 'unit-1', 'equipment-1', 48.1234, 31.5678]


def test_completed_review_workbook_deduplicates_rows_and_keeps_all_evidence():
    common = {
        'sourceMethod': 'native', 'documentId': 'doc-1',
        'reviewStatus': 'accepted',
    }
    items = [
        {
            **common, 'id': 'scenario-1', 'kind': 'scenario',
            'name': 'عملیات نمونه', 'evidence': 'عملیات نمونه آغاز شد',
            'sourcePage': 1, 'startTime': '2026-01-01T08:00:00+03:30',
        },
        {
            **common, 'id': 'unit-1', 'kind': 'unit', 'name': 'تیپ نمونه',
            'evidence': 'تیپ نمونه وارد منطقه شد', 'sourcePage': 2,
            'side': 'خودی', 'unitType': 'زرهی', 'echelon': 'تیپ',
            'resourceCode': 'UNIT-EXAMPLE',
        },
        {
            **common, 'id': 'unit-2', 'kind': 'unit', 'name': 'تیپ‌ نمونه',
            'evidence': 'گزارش تیپ‌ نمونه ادامه یافت', 'sourcePage': 3,
        },
    ]
    workbook = build_document_review_workbook(
        filename='report.pdf', document_id='doc-1', page_count=3,
        main_scenario_id='scenario-1', items=items,
    )
    assert workbook['یگان‌ها'].max_row == 2
    assert workbook['یگان‌ها']['L2'].value == 'UNIT-EXAMPLE'
    assert workbook['کالک‌یار'].max_row >= 8
    assert workbook['کالک‌یار']['D3'].value == items[1]['evidence']
    assert workbook['کالک‌یار']['D4'].value == items[2]['evidence']
    from app.services.excel_scenario_import import parse_excel_workbook
    _, errors = parse_excel_workbook(workbook)
    assert errors == []


def test_document_workbook_route_returns_xlsx_without_writing_database():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from io import BytesIO
    from app.api.routes import data_import

    app = FastAPI()
    app.include_router(data_import.router)
    route = next(r for r in data_import.router.routes if r.path == '/data-import/document/workbook')
    for dependency in route.dependencies:
        app.dependency_overrides[dependency.dependency] = lambda: None
    payload = {
        'filename': 'report.pdf', 'documentId': 'doc-1', 'pageCount': 1,
        'mainScenarioId': 'scenario-1',
        'items': [{
            'id': 'scenario-1', 'kind': 'scenario', 'name': 'عملیات نمونه',
            'evidence': 'عملیات نمونه', 'sourcePage': 1,
            'sourceMethod': 'native', 'documentId': 'doc-1',
            'reviewStatus': 'accepted',
        }],
    }
    with TestClient(app) as client:
        response = client.post('/data-import/document/workbook', json=payload)
    assert response.status_code == 200
    assert response.headers['content-type'].startswith(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook['سناریو']['A2'].value == 'عملیات نمونه'
    assert workbook['کالک‌یار']['B2'].value == 'عملیات نمونه'


def test_legacy_visual_order_keeps_unit_numbers():
    assert normalize_native_text('92 ﺮﮕﺸﻟ') == 'لشگر 92'
    assert normalize_native_text('لشکر ۹۲ زرهی') == 'لشکر ۹۲ زرهی'


def test_source_evidence_required_and_identity_not_inferred():
    items, rejected = validate_proposals([
        {"kind": "unit", "name": "لشکر ۹۲", "evidence": "لشکر ۹۲ در متن", "lat": 32, "resourceId": "fake"},
        {"kind": "person", "name": "نام ساختگی", "evidence": "شاهد ساختگی"},
    ], 'لشکر ۹۲ در متن آمده است.', 'doc', 43, 'native')
    assert rejected == 1 and len(items) == 1
    assert items[0]['sourcePage'] == 43
    assert items[0]['reviewStatus'] == 'pending'
    assert 'lat' not in items[0] and 'resourceId' not in items[0]


def test_repeated_ocr_rejected():
    assert repetition_detected('مخالف\n' * 100)
    assert not repetition_detected('یک صفحه کوتاه')


def test_ocr_preparation_crops_scanner_whitespace_and_enlarges_content():
    page = Image.new('RGB', (1000, 1000), 'white')
    draw = ImageDraw.Draw(page)
    draw.rectangle((300, 250, 700, 750), fill='black')
    prepared = prepare_ocr_image(page, max_side=1200)
    try:
        assert 1000 < max(prepared.size) <= 1200
        assert prepared.width < prepared.height
        assert prepared.size != page.size
    finally:
        prepared.close()
        page.close()


def test_quote_whitespace_is_restored_and_changed_number_rejected():
    source = 'تیپ55 هوابرد که\nدر منطقه بود'
    items, rejected = validate_proposals([
        {'kind': 'unit', 'name': 'تیپ 55 هوابرد', 'evidence': 'تیپ 55 هوابرد که در منطقه بود'},
        {'kind': 'unit', 'name': 'تیپ 45 هوابرد', 'evidence': 'تیپ 45 هوابرد که در منطقه بود'},
    ], source, 'doc', 43, 'native')
    assert rejected == 1
    assert items[0]['evidence'] == source


def test_page_header_is_not_a_timeline_event():
    source = 'طرح ریزی عملیات طریق القدس/٤٣'
    items, rejected = validate_proposals([
        {'kind': 'event', 'name': 'طریق القدس/٤٣', 'evidence': source},
    ], source, 'doc', 43, 'native')
    assert items == [] and rejected == 1


def test_semantic_rules_keep_resources_and_named_operations_conservative():
    assert normalize_proposal_semantics(
        'event', 'عملیات طریق القدس', 'عملیات طریق القدس آغاز شد'
    ) == ('scenario', 'عملیات طریق القدس')
    assert normalize_proposal_semantics(
        'place', 'طریق القدس', 'تاریخ عملیات طریق القدس'
    ) == ('scenario', 'عملیات طریق القدس')
    assert normalize_proposal_semantics(
        'scenario', 'دشت آزادگان', 'شماتیک دشت آزادگان و منطقه عملیات'
    ) == ('place', 'دشت آزادگان')
    assert normalize_proposal_semantics(
        'unit', 'شکستن حصر آبادان', 'فرمان شکستن حصر آبادان صادر شد'
    ) is None
    assert normalize_proposal_semantics(
        'person', 'نیروهای ایرانی', 'حضور نیروهای ایرانی'
    ) is None
    assert normalize_proposal_semantics(
        'equipment', 'خاکریز سوّم', 'احداث خاکریز سوّم پایان یافت'
    ) == ('place', 'خاکریز سوّم')
    assert normalize_proposal_semantics(
        'equipment', 'گردان تانک', 'گردان تانک در منطقه مستقر شد'
    ) == ('unit', 'گردان تانک')
    assert normalize_proposal_semantics(
        'scenario', 'پایان عملیات طریق القدس', 'پایان عملیات طریق القدس اعلام شد'
    ) == ('event', 'پایان عملیات طریق القدس')
    assert normalize_proposal_semantics(
        'scenario', 'ستاد عملیات کربلا 1', 'ستاد عملیات کربلا 1 تشکیل شد'
    ) == ('unit', 'ستاد عملیات کربلا 1')
    assert normalize_proposal_semantics(
        'scenario', 'منطقه عملیات طریق القدس', 'منطقه عملیات طریق القدس مشخص شد'
    ) == ('place', 'منطقه عملیات طریق القدس')
    assert normalize_proposal_semantics(
        'scenario', 'طرح ریزی عملیات طریق القدس/٤١',
        'طرح ریزی عملیات طریق القدس/٤١',
    ) == ('scenario', 'عملیات طریق القدس')
    assert normalize_proposal_semantics(
        'scenario', 'عملیات آفندی', 'ادامه عملیات آفندی بررسی شد'
    ) is None
    assert normalize_proposal_semantics(
        'scenario', 'گسترش نیروهای دشمن در منطقه عملیات',
        'گسترش نیروهای دشمن در منطقه عملیات ادامه یافت',
    ) == ('event', 'گسترش نیروهای دشمن در منطقه عملیات')
    assert normalize_proposal_semantics(
        'scenario', 'عملیات جاده سازی در منطقه رملها',
        'عملیات جاده سازی در منطقه رملها انجام شد',
    ) == ('event', 'عملیات جاده سازی در منطقه رملها')


def test_invalid_document_and_page_rejected():
    with pytest.raises(ValueError):
        read_page(b'not pdf', 'test.pdf', 1)
    with pytest.raises(ValueError):
        read_page(b'text', 'test.txt', 2)
    assert read_page('متن'.encode(), 'test.txt', 1)['text'] == 'متن'


def test_native_and_image_text_are_both_preserved():
    merged = merge_page_text('توضیح بالای صفحه', 'نام روستا در نقشه')
    assert 'توضیح بالای صفحه' in merged
    assert 'نام روستا در نقشه' in merged
    assert merge_page_text('متن یکسان', 'متن  یکسان') == 'متن  یکسان'


def test_person_titles_share_one_canonical_entity_but_keep_original_names():
    first = {'kind': 'person', 'name': 'شهید علی صیاد شیرازی'}
    second = {'kind': 'person', 'name': 'سپهبد علی صیاد شیرازی'}
    assert canonical_entity_name(first['kind'], first['name']) == 'علی صیاد شیرازی'
    assert proposal_entity_key(first) == proposal_entity_key(second)

    first_draft = entity_draft_id('doc-1', first['kind'], first['name'])
    second_draft = entity_draft_id('doc-1', second['kind'], second['name'])
    assert first_draft == second_draft
    assert proposed_resource_reference_code('person', first_draft).startswith('PER-')

    common = {
        'sourceMethod': 'native', 'documentId': 'doc-1',
        'reviewStatus': 'accepted', 'resourceCode': 'PERSON-SAYYAD',
    }
    workbook = build_document_review_workbook(
        filename='report.pdf', document_id='doc-1', page_count=3,
        main_scenario_id='scenario-1',
        items=[
            {**common, 'id': 'scenario-1', 'kind': 'scenario',
             'name': 'عملیات نمونه', 'evidence': 'عملیات نمونه',
             'sourcePage': 1, 'startTime': '2026-01-01T08:00:00+03:30',
             'resourceCode': None},
            {**common, **first, 'id': 'person-1',
             'evidence': 'شهید علی صیاد شیرازی', 'sourcePage': 2},
            {**common, **second, 'id': 'person-2',
             'evidence': 'سپهبد علی صیاد شیرازی', 'sourcePage': 3},
        ],
    )
    assert workbook['پرسنل'].max_row == 2
    assert workbook['پرسنل']['B2'].value == 'علی صیاد شیرازی'
    assert workbook['پرسنل']['I2'].value == (
        'شهید علی صیاد شیرازی | سپهبد علی صیاد شیرازی'
    )
    assert workbook['کالک‌یار']['B3'].value == 'شهید علی صیاد شیرازی'
    assert workbook['کالک‌یار']['B4'].value == 'سپهبد علی صیاد شیرازی'
    assert workbook['کالک‌یار']['C3'].value == 'علی صیاد شیرازی'
    assert workbook['کالک‌یار']['C4'].value == 'علی صیاد شیرازی'
    from app.services.excel_scenario_import import parse_resources_workbook
    personnel, _equipment, _units, errors = parse_resources_workbook(workbook)
    assert errors == []
    assert personnel[0]['aliases'] == [
        'شهید علی صیاد شیرازی', 'سپهبد علی صیاد شیرازی'
    ]


def test_reviewed_canonical_name_changes_display_not_draft_identity_or_aliases():
    draft_id = entity_draft_id('doc-2', 'person', 'شهید علی صیاد شیرازی')
    common = {
        'kind': 'person', 'canonicalName': 'سپهبد علی صیاد شیرازی',
        'entityDraftId': draft_id, 'sourceMethod': 'native',
        'documentId': 'doc-2', 'reviewStatus': 'accepted',
    }
    workbook = build_document_review_workbook(
        filename='report.pdf', document_id='doc-2', page_count=2,
        main_scenario_id='scenario-1',
        items=[
            {'id': 'scenario-1', 'kind': 'scenario', 'name': 'عملیات نمونه',
             'evidence': 'عملیات نمونه', 'sourcePage': 1, 'sourceMethod': 'native',
             'documentId': 'doc-2', 'reviewStatus': 'accepted',
             'startTime': '2026-01-01T08:00:00+03:30'},
            {**common, 'id': 'person-1', 'name': 'شهید علی صیاد شیرازی',
             'evidence': 'شهید علی صیاد شیرازی', 'sourcePage': 1},
            {**common, 'id': 'person-2', 'name': 'سپهبد علی صیاد شیرازی',
             'evidence': 'سپهبد علی صیاد شیرازی', 'sourcePage': 2},
        ],
    )
    assert workbook['پرسنل']['B2'].value == 'سپهبد علی صیاد شیرازی'
    assert workbook['پرسنل']['H2'].value.startswith('PER-')
    assert workbook['کالک‌یار']['I3'].value == draft_id
    assert workbook['کالک‌یار']['I4'].value == draft_id


def test_visual_page_analysis_accepts_multiple_kinds_and_rotation():
    result = visual_page_analysis(
        '```json\n{"pageKinds":["text","military-map","unknown"],'
        '"mapConfidence":1.4,"rotationDegrees":90}\n```'
    )
    assert result == {
        'pageKinds': ['text', 'military-map'],
        'mapConfidence': 1.0,
        'rotationDegrees': 90,
    }


@pytest.mark.asyncio
async def test_mixed_map_page_keeps_native_text_when_image_ocr_fails(monkeypatch):
    from app.services import document_extraction as service

    async def fake_classify(client, settings, image_bytes):
        return {
            'pageKinds': ['text', 'image', 'military-map'],
            'mapConfidence': .92,
            'rotationDegrees': 90,
        }

    async def fake_completion(client, settings, model, messages, **kwargs):
        if kwargs.get('structured'):
            return (
                '{"items":[{"kind":"event","name":"آغاز عملیات",'
                '"evidence":"آغاز عملیات"}]}'
            )
        raise ValueError(
            'بازشناسی این صفحه تصویری کامل نشد؛ صفحه احتمالاً نقشه یا تصویر پرجزئیات است.'
        )

    monkeypatch.setattr(service, 'classify_visual_page', fake_classify)
    monkeypatch.setattr(service, 'completion', fake_completion)
    settings = SimpleNamespace(
        DOCUMENT_OCR_PIPELINE_URL='', DOCUMENT_OCR_MODEL='paddle',
        DOCUMENT_LLM_MODEL='qwen',
    )
    result = await service.extract_page(
        {
            'page': 35, 'pageCount': 100, 'text': 'آغاز عملیات',
            'nativeText': 'آغاز عملیات', 'ocrText': '',
            'pageKinds': ['text', 'image'],
            'image': Image.new('RGB', (400, 300), 'white'),
            'visualImage': Image.new('RGB', (200, 150), 'white'),
            'hasSignificantImage': True,
            'method': 'hybrid', 'warnings': [],
        },
        b'document', 'report.pdf', settings,
    )

    assert result['text'] == 'آغاز عملیات'
    assert result['items'][0]['kind'] == 'event'
    assert result['pageKinds'] == ['text', 'image', 'military-map']
    assert result['mapCandidate']['rotationDegrees'] == 90


def test_review_workbook_contains_page_coverage_ledger():
    scenario = {
        'id': 'scenario-1', 'kind': 'scenario', 'name': 'عملیات نمونه',
        'evidence': 'عملیات نمونه', 'sourcePage': 1,
        'sourceMethod': 'native', 'documentId': 'doc-1', 'reviewStatus': 'accepted',
    }
    workbook = build_document_review_workbook(
        filename='report.pdf', document_id='doc-1', page_count=2,
        main_scenario_id='scenario-1', items=[scenario],
        coverage=[
            {'page': 1, 'pageKinds': ['text'], 'status': 'reviewed',
             'itemCount': 1, 'acceptedCount': 1, 'rejectedCount': 0,
             'pendingCount': 0, 'warnings': []},
            {'page': 2, 'pageKinds': ['image', 'military-map'],
             'status': 'no_relevant_data', 'itemCount': 0, 'acceptedCount': 0,
             'rejectedCount': 0, 'pendingCount': 0, 'mapStatus': 'ignored',
             'warnings': ['بازبینی شد']},
        ],
        finalized=True,
    )
    assert workbook['پوشش سند']['A3'].value == 2
    assert workbook['پوشش سند']['H3'].value == 'ignored'
    assert workbook['کالک‌یار']['H2'].value == 'accepted'


def test_truncated_model_output_is_not_a_success():
    class Response:
        def raise_for_status(self):
            pass
        def json(self):
            return {'choices': [{'finish_reason': 'length', 'message': {'content': '{}'}}]}
    class Client:
        async def post(self, url, **kwargs):
            assert url == 'http://localhost:1234/v1/chat/completions'
            return Response()
    settings = SimpleNamespace(INTERNAL_LLM_BASE_URL='http://localhost:1234/v1', INTERNAL_LLM_API_KEY='')
    with pytest.raises(ValueError):
        asyncio.run(completion(Client(), settings, 'model', [], structured=True))


@pytest.mark.asyncio
async def test_document_model_status_checks_separate_services(monkeypatch):
    requested = []

    class Response:
        def raise_for_status(self):
            pass
        def json(self):
            return {'data': [{'id': 'qwen'}, {'id': 'paddle'}]}

    class Client:
        def __init__(self, *args, **kwargs):
            pass
        async def __aenter__(self):
            return self
        async def __aexit__(self, *args):
            pass
        async def get(self, url, **kwargs):
            requested.append(url)
            return Response()

    monkeypatch.setattr('app.services.document_extraction.httpx.AsyncClient', Client)
    settings = SimpleNamespace(
        DOCUMENT_LLM_BASE_URL='http://llm:1234', DOCUMENT_LLM_MODEL='qwen',
        DOCUMENT_OCR_BASE_URL='http://ocr:8000/v1', DOCUMENT_OCR_MODEL='paddle',
        DOCUMENT_OCR_PIPELINE_URL='',
        INTERNAL_LLM_BASE_URL='', INTERNAL_LLM_API_KEY='',
    )
    result = await document_model_status(settings)
    assert result['ready'] is True
    assert set(requested) == {'http://llm:1234/v1/models', 'http://ocr:8000/v1/models'}
    assert result['ocrMode'] == 'vlm-component'


@pytest.mark.asyncio
async def test_full_pipeline_contract_and_markdown_extraction():
    class Response:
        def raise_for_status(self):
            pass
        def json(self):
            return {'errorCode': 0, 'result': {'layoutParsingResults': [
                {'markdown': {'text': 'متن بخش نخست'}},
                {'markdown': {'text': 'متن بخش دوم'}},
            ]}}

    class Client:
        async def post(self, url, **kwargs):
            assert url == 'http://ocr-pipeline:8080/layout-parsing'
            assert kwargs['json']['fileType'] == 1
            assert kwargs['json']['useLayoutDetection'] is True
            assert kwargs['json']['returnMarkdownImages'] is False
            return Response()

    settings = SimpleNamespace(
        DOCUMENT_OCR_PIPELINE_URL='http://ocr-pipeline:8080',
        INTERNAL_LLM_API_KEY='',
    )
    text = await pipeline_ocr(Client(), settings, b'image')
    assert text == 'متن بخش نخست\n\nمتن بخش دوم'
    with pytest.raises(ValueError):
        pipeline_markdown({'errorCode': 1, 'errorMsg': 'failed'})


def test_document_preview_route_contract(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.api.routes import data_import

    app = FastAPI()
    app.include_router(data_import.router)
    route = next(r for r in data_import.router.routes if r.path == '/data-import/document/preview')
    # Isolated authorization override only; no production app or database access.
    for dependency in route.dependencies:
        app.dependency_overrides[dependency.dependency] = lambda: None
    monkeypatch.setattr(data_import.settings, 'DOCUMENT_LLM_BASE_URL', 'http://local-test')

    async def fake_extract(page_data, raw, filename, settings):
        assert page_data['text'] == 'متن آزمایش' and page_data['page'] == 1
        return {'items': [], 'persisted': False, 'filename': filename}
    monkeypatch.setattr(data_import, 'extract_page', fake_extract)
    with TestClient(app) as client:
        response = client.post('/data-import/document/preview', files={'file': ('report.txt', 'متن آزمایش'.encode())}, data={'page': '1'})
        assert response.status_code == 200
        assert response.json()['data']['persisted'] is False
        response = client.post('/data-import/document/preview', files={'file': ('report.pdf', b'not pdf')})
        assert response.status_code == 422


@pytest.mark.asyncio
async def test_whole_document_continues_after_page_failure_and_skips_completed(monkeypatch):
    from app.services import document_extraction as service
    visited = []
    def fake_read(raw, filename, page, force):
        visited.append(page)
        return {'page': page}
    async def fake_extract(data, *args):
        if data['page'] == 2:
            raise ValueError('خطای صفحه دوم')
        return {'page': data['page'], 'items': []}
    monkeypatch.setattr(service, 'read_page', fake_read)
    monkeypatch.setattr(service, 'extract_page', fake_extract)
    events = [event async for event in service.stream_document(b'sample', 'report.pdf', 4, {1}, False, None)]
    assert visited == [2, 3, 4]
    assert events[0]['type'] == 'metadata' and events[0]['pageCount'] == 4
    assert [e['result']['page'] for e in events if e['type'] == 'page'] == [3, 4]
    assert [e['page'] for e in events if e['type'] == 'page_error'] == [2]
    assert events[-1] == {'type': 'done', 'pageCount': 4, 'succeeded': 3, 'failed': 1}


def test_whole_document_route_streams_and_rejects_invalid_skip(monkeypatch):
    import json
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.api.routes import data_import
    app = FastAPI()
    app.include_router(data_import.router)
    route = next(r for r in data_import.router.routes if r.path == '/data-import/document/preview-all')
    for dependency in route.dependencies:
        app.dependency_overrides[dependency.dependency] = lambda: None
    monkeypatch.setattr(data_import.settings, 'DOCUMENT_LLM_BASE_URL', 'http://local-test')
    async def fake_stream(raw, filename, total, skip, force, settings):
        assert raw == b'sample' and total == 1
        yield {'type': 'metadata', 'pageCount': total}
        yield {'type': 'done', 'succeeded': 1, 'failed': 0}
    monkeypatch.setattr(data_import, 'stream_document', fake_stream)
    with TestClient(app) as client:
        response = client.post('/data-import/document/preview-all', files={'file': ('report.txt', b'sample')})
        assert response.status_code == 200
        assert response.headers['x-accel-buffering'] == 'no'
        assert [json.loads(line)['type'] for line in response.text.splitlines()] == ['metadata', 'done']
        response = client.post('/data-import/document/preview-all', files={'file': ('report.txt', b'sample')}, data={'skip_pages': '[2]'})
        assert response.status_code == 422
