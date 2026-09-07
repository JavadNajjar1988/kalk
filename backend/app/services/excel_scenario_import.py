"""
Parse standard multi-sheet Excel into ORBAT-mapper scenario dict + validation errors.
Sheet names (case-insensitive, trimmed): سناریو|scenario, حوادث|events, یگان‌ها|units,
تجهیزات|equipment, پرسنل|personnel, وضعیت‌های زمانی یگان‌ها|unit_states
"""

from __future__ import annotations

import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.worksheet import Worksheet

BUNDLED_GRAPHIC_TEMPLATE = (
    Path(__file__).resolve().parents[1]
    / "templates"
    / "scenario_import_template_graphic_fa_main_orbat.xlsx"
)
BUNDLED_COMPLETE_EXAMPLE = (
    Path(__file__).resolve().parents[1]
    / "templates"
    / "scenario_import_example_complete_fa.xlsx"
)

# Default 20-digit APP-6D / MIL-STD-2525D SIDC: friendly land infantry,
# present status, unknown echelon.
DEFAULT_SIDC = "10031000001211000000"

SIDE_MAP = {
    "friend": "3",
    "friendly": "3",
    "دوست": "3",
    "خودی": "3",
    "blue": "3",
    "hostile": "6",
    "دشمن": "6",
    "enemy": "6",
    "red": "6",
    "neutral": "4",
    "خنثی": "4",
    "unknown": "0",
    "نامشخص": "0",
}

UNIT_SYMBOL_TYPES = {
    "infantry": ("10", "121100"),
    "پیاده": ("10", "121100"),
    "پیاده_نظام": ("10", "121100"),
    "mechanized_infantry": ("10", "121102"),
    "پیاده_نظام_مکانیزه": ("10", "121102"),
    "پیاده_مکانیزه": ("10", "121102"),
    "reconnaissance": ("10", "121300"),
    "شناسایی": ("10", "121300"),
    "artillery": ("10", "130300"),
    "توپخانه": ("10", "130300"),
    "armor": ("10", "120500"),
    "armour": ("10", "120500"),
    "tank": ("10", "120500"),
    "زرهی": ("10", "120500"),
    "combat_service_support": ("10", "160600"),
    "پشتیبانی_خدمات_رزمی": ("10", "160600"),
    "air_defense": ("10", "130100"),
    "پدافند_هوایی": ("10", "130100"),
    "combat_engineer": ("10", "140700"),
    "مهندسی_رزمی": ("10", "140700"),
}

# APP-6D entity symbols used for point equipment imported from Excel.
# The symbol set is kept alongside the entity code because aircraft and land
# equipment use different sets.
EQUIPMENT_SYMBOL_TYPES: dict[str, tuple[str, str]] = {
    "تانک": ("15", "120200"),
    "تانک_تی_54": ("15", "120200"),
    "تانک_تی_64": ("15", "120200"),
    "تی_54": ("15", "120200"),
    "تی54": ("15", "120200"),
    "t_54": ("15", "120200"),
    "t54": ("15", "120200"),
    "تی_64": ("15", "120200"),
    "تی64": ("15", "120200"),
    "t_64": ("15", "120200"),
    "t64": ("15", "120200"),
    "پی_ام_پی": ("15", "120101"),
    "پی_ام_پی_1": ("15", "120101"),
    "bmp": ("15", "120101"),
    "bmp_1": ("15", "120101"),
    "نفربر_زرهی": ("15", "120101"),
    "خودروی_رزمی_زرهی": ("15", "120101"),
    "دی_8": ("15", "131100"),
    "دی8": ("15", "131100"),
    "d_8": ("15", "131100"),
    "d8": ("15", "131100"),
    "بلدوزر": ("15", "131100"),
    "بلدزر": ("15", "131100"),
    "توپ_155": ("15", "110903"),
    "توپ_155_م_م": ("15", "110903"),
    "هویتزر_155": ("15", "110903"),
    "هویتزر_155_میلی_متری": ("15", "110903"),
    "خودرو": ("15", "140100"),
    "خودروی_نظامی": ("15", "140100"),
    "وسیله_نقلیه": ("15", "140100"),
    "هواپیما": ("01", "110100"),
    "هواپیمای_بال_ثابت": ("01", "110100"),
    "aircraft": ("01", "110100"),
    "fixed_wing": ("01", "110100"),
    "بالگرد": ("01", "110200"),
    "هلیکوپتر": ("01", "110200"),
    "helicopter": ("01", "110200"),
}

ECHELON_MAP = {
    "": "00",
    "unknown": "00",
    "نامشخص": "00",
    "team": "11",
    "crew": "11",
    "تیم": "11",
    "خدمه": "11",
    "تیم_خدمه": "11",
    "squad": "12",
    "گروه": "12",
    "section": "13",
    "جوخه": "13",
    "بخش": "13",
    "جوخه_بخش": "13",
    "platoon": "14",
    "دسته": "14",
    "جزء_مستقل": "14",
    "دسته_جزء_مستقل": "14",
    "company": "15",
    "battery": "15",
    "گروهان": "15",
    "آتشبار": "15",
    "گروهان_آتشبار": "15",
    "battalion": "16",
    "squadron": "16",
    "گردان": "16",
    "اسکادران": "16",
    "گردان_اسکادران": "16",
    "regiment": "17",
    "هنگ": "17",
    "هنگ_گروه": "17",
    "brigade": "18",
    "تیپ": "18",
    "division": "21",
    "لشکر": "21",
    "corps": "22",
    "سپاه": "22",
    "army": "23",
    "ارتش": "23",
    "army_group": "24",
    "front": "24",
    "گروه_ارتش": "24",
    "جبهه": "24",
    "گروه_ارتش_جبهه": "24",
    "region": "25",
    "theater": "25",
    "منطقه": "25",
    "صحنه_عملیات": "25",
    "منطقه_صحنه_عملیات": "25",
    "command": "26",
    "فرماندهی": "26",
}

SYMBOL_STATUS_MAP = {
    "": "0",
    "present": "0",
    "حاضر": "0",
    "planned": "1",
    "anticipated": "1",
    "برنامه_ریزی_شده": "1",
    "پیش_بینی_شده": "1",
    "damaged": "3",
    "آسیب_دیده": "3",
    "destroyed": "4",
    "نابود_شده": "4",
    "نابودشده": "4",
}

MAIN_ECHELONS = [
    "ارتش",
    "سپاه",
    "لشکر",
    "تیپ",
    "هنگ / گروه",
    "گردان / اسکادران",
    "گروهان / آتشبار",
    "دسته / جزء مستقل",
    "نامشخص",
]


def _build_unit_sidc(
    side_id: str, unit_type: str, echelon: str = "", status: str = ""
) -> str:
    symbol_set, main_icon = UNIT_SYMBOL_TYPES.get(
        _norm_key(unit_type), UNIT_SYMBOL_TYPES["infantry"]
    )
    echelon_code = ECHELON_MAP.get(_norm_key(echelon), "00")
    status_code = SYMBOL_STATUS_MAP.get(_norm_key(status), "0")
    return f"100{side_id}{symbol_set}{status_code}0{echelon_code}{main_icon}0000"


def _equipment_symbol_key(value: str | None) -> str:
    translated = (
        _norm(value)
        .lower()
        .translate(str.maketrans("كي۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "کی01234567890123456789"))
    )
    return re.sub(r"[^\w]+", "_", translated, flags=re.UNICODE).strip("_")


def _equipment_symbol_spec(
    equipment_type: str, equipment_name: str
) -> tuple[str, str] | None:
    for candidate in (
        equipment_type,
        equipment_name,
        f"{equipment_name} {equipment_type}",
    ):
        spec = EQUIPMENT_SYMBOL_TYPES.get(_equipment_symbol_key(candidate))
        if spec:
            return spec
    return None


def _build_equipment_sidc(
    side_id: str, equipment_type: str, equipment_name: str
) -> str | None:
    spec = _equipment_symbol_spec(equipment_type, equipment_name)
    if not spec:
        return None
    symbol_set, main_icon = spec
    return f"100{side_id}{symbol_set}0000{main_icon}0000"


def _clean_sidc(value: str) -> str:
    return _norm(value).lstrip("'").replace("‌", "").replace(" ", "")


def _norm(s: str | None) -> str:
    if s is None:
        return ""
    return str(s).strip()


def _norm_key(s: str) -> str:
    t = _norm(s).lower()
    t = re.sub(r"[\s\u200c\-/–—]+", "_", t)
    return t.strip("_")


def _find_sheet(wb, *aliases: str) -> Worksheet | None:
    alias_set = {_norm_key(a) for a in aliases}
    for name in wb.sheetnames:
        if _norm_key(name) in alias_set or _norm(name).lower() in {
            a.lower() for a in aliases
        }:
            return wb[name]
    return None


def _header_map(ws: Worksheet) -> tuple[dict[str, int], list[dict[str, Any]]]:
    """Row 1 = headers. Return lowercase_underscore -> col index (1-based), and errors."""
    errors: list[dict[str, Any]] = []
    if ws.max_row is None or ws.max_row < 1:
        return {}, [{"sheet": ws.title, "row": 0, "message": "ورق خالی است"}]
    headers: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=col).value
        key = _norm_key(str(val) if val is not None else "")
        if key:
            headers[key] = col
    return headers, errors


def _cell(row: tuple[Any, ...], headers: dict[str, int], *keys: str) -> str:
    for k in keys:
        k2 = _norm_key(k)
        if k2 in headers:
            c = headers[k2]
            if c <= len(row):
                v = row[c - 1]
                if v is None:
                    return ""
                return str(v).strip()
    return ""


def _cell_raw(row: tuple[Any, ...], headers: dict[str, int], *keys: str) -> Any:
    for k in keys:
        k2 = _norm_key(k)
        if k2 in headers:
            c = headers[k2]
            if c <= len(row):
                return row[c - 1]
    return None


def _parse_iso_or_ms(
    val: Any, errors: list, sheet: str, row: int, field: str
) -> str | int | None:
    if isinstance(val, datetime):
        return val.isoformat()
    if isinstance(val, (int, float)):
        if val > 10_000_000_000:
            return int(val)
        if 1_000 <= val < 100_000:
            return from_excel(val).isoformat()
    v = _norm(str(val) if val is not None else "")
    if not v:
        return None
    try:
        # numeric epoch ms
        if v.isdigit() and len(v) > 10:
            return int(v)
        # ISO-ish
        datetime.fromisoformat(v.replace("Z", "+00:00"))
        return v.replace("Z", "+00:00") if "Z" in v else v
    except Exception:
        errors.append(
            {
                "sheet": sheet,
                "row": row,
                "message": f"تاریخ/زمان نامعتبر ({field}): {v}",
            }
        )
        return None


def _validate_lon_lat(
    lon_s: str, lat_s: str, errors: list, sheet: str, row: int
) -> tuple[float | None, float | None]:
    if not lon_s and not lat_s:
        return None, None
    try:
        lon = float(lon_s)
        lat = float(lat_s)
    except ValueError:
        errors.append({"sheet": sheet, "row": row, "message": "مختصات باید عدد باشند"})
        return None, None
    if not (-180 <= lon <= 180 and -90 <= lat <= 90):
        errors.append(
            {"sheet": sheet, "row": row, "message": "مختصات خارج از محدوده مجاز"}
        )
        return None, None
    return lon, lat


def _read_rows(ws: Worksheet) -> list[tuple[Any, ...]]:
    rows: list[tuple[Any, ...]] = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r is None or all(x is None or str(x).strip() == "" for x in r):
            continue
        rows.append(tuple(r))
    return rows


def _split_ids(value: str) -> list[str]:
    return [item.strip() for item in re.split(r"[,،;\n]+", value or "") if item.strip()]


def _equipment_display_name(name: str, equipment_type: str) -> str:
    if not equipment_type or _norm_key(equipment_type) in {
        "general",
        "عمومی",
        _norm_key(name),
    }:
        return name
    return f"{name} - {equipment_type}"


def _feature_style(color: str, symbol: str = "circle") -> dict[str, Any]:
    return {
        "marker-color": color,
        "marker-size": "medium",
        "marker-symbol": symbol,
        "stroke": color,
        "stroke-width": 2,
        "fill": color,
        "fill-opacity": 0.18,
    }


def _jalali_to_gregorian(jy: int, jm: int, jd: int) -> tuple[int, int, int]:
    jy += 1595
    days = -355668 + (365 * jy) + ((jy // 33) * 8) + (((jy % 33) + 3) // 4) + jd
    days += (jm - 1) * 31 if jm < 7 else ((jm - 7) * 30) + 186
    gy = 400 * (days // 146097)
    days %= 146097
    if days > 36524:
        days -= 1
        gy += 100 * (days // 36524)
        days %= 36524
        if days >= 365:
            days += 1
    gy += 4 * (days // 1461)
    days %= 1461
    if days > 365:
        gy += (days - 1) // 365
        days = (days - 1) % 365
    gd = days + 1
    month_days = [
        0,
        31,
        29 if gy % 4 == 0 and (gy % 100 != 0 or gy % 400 == 0) else 28,
        31,
        30,
        31,
        30,
        31,
        31,
        30,
        31,
        30,
        31,
    ]
    gm = 1
    while gm <= 12 and gd > month_days[gm]:
        gd -= month_days[gm]
        gm += 1
    return gy, gm, gd


def _jalali_month_length(year: int, month: int) -> int:
    if month <= 6:
        return 31
    if month <= 11:
        return 30
    current_year = datetime(*_jalali_to_gregorian(year, 1, 1))
    next_year = datetime(*_jalali_to_gregorian(year + 1, 1, 1))
    return 30 if (next_year - current_year).days == 366 else 29


def _parse_time_references(wb, errors: list[dict[str, Any]]) -> dict[str, str]:
    ws = _find_sheet(wb, "زمان‌بندی", "زمان بندی", "زمانبندی", "times", "time_lookup")
    if not ws:
        return {}
    headers, _ = _header_map(ws)
    result: dict[str, str] = {}
    for row_index, row in enumerate(_read_rows(ws), start=2):
        time_id = _cell(row, headers, "id", "time_id", "شناسه", "شناسه_زمان")
        if not time_id:
            continue
        gregorian = _cell_raw(
            row,
            headers,
            "gregorian_datetime",
            "datetime",
            "تاریخ_و_ساعت_میلادی",
            "تاریخ_میلادی",
        )
        if gregorian is not None and str(gregorian).strip():
            parsed = _parse_iso_or_ms(
                gregorian,
                errors,
                ws.title,
                row_index,
                "gregorian_datetime",
            )
            if parsed is not None:
                result[time_id] = str(parsed)
            continue
        try:
            jy = int(float(_cell(row, headers, "jalali_year", "سال_شمسی", "سال")))
            jm = int(float(_cell(row, headers, "jalali_month", "ماه_شمسی", "ماه")))
            jd = int(float(_cell(row, headers, "jalali_day", "روز_شمسی", "روز")))
            hour = int(float(_cell(row, headers, "hour", "ساعت") or "0"))
            minute = int(float(_cell(row, headers, "minute", "دقیقه") or "0"))
            if not (
                1 <= jm <= 12
                and 1 <= jd <= _jalali_month_length(jy, jm)
                and 0 <= hour <= 23
                and 0 <= minute <= 59
            ):
                raise ValueError
            gy, gm, gd = _jalali_to_gregorian(jy, jm, jd)
            zone_name = (
                _cell(
                    row, headers, "time_zone", "timezone", "ناحیه_زمانی", "منطقه_زمانی"
                )
                or "Asia/Tehran"
            )
            try:
                tz = ZoneInfo(zone_name)
            except ZoneInfoNotFoundError:
                errors.append(
                    {
                        "sheet": ws.title,
                        "row": row_index,
                        "message": f"ناحیه زمانی نامعتبر است: {zone_name}",
                    }
                )
                tz = timezone.utc
            result[time_id] = datetime(gy, gm, gd, hour, minute, tzinfo=tz).isoformat()
        except (TypeError, ValueError):
            errors.append(
                {
                    "sheet": ws.title,
                    "row": row_index,
                    "message": "اجزای تاریخ شمسی یا ساعت نامعتبر است",
                }
            )
    return result


def _resolve_row_time(
    row: tuple[Any, ...],
    headers: dict[str, int],
    errors: list[dict[str, Any]],
    sheet: str,
    row_number: int,
    field: str,
    time_refs: dict[str, str],
    *,
    direct_keys: tuple[str, ...],
    reference_keys: tuple[str, ...],
) -> str | int | None:
    direct = _cell_raw(row, headers, *direct_keys)
    if direct is not None and str(direct).strip() != "":
        return _parse_iso_or_ms(direct, errors, sheet, row_number, field)
    reference = _cell(row, headers, *reference_keys)
    if not reference:
        return None
    if reference not in time_refs:
        errors.append(
            {
                "sheet": sheet,
                "row": row_number,
                "message": f"شناسه زمان یافت نشد: {reference}",
            }
        )
        return None
    return time_refs[reference]


def _scenario_time_sort_key(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value or "").strip()
    try:
        return float(text)
    except ValueError:
        try:
            return datetime.fromisoformat(text.replace("Z", "+00:00")).timestamp() * 1000
        except ValueError:
            return float("inf")


def _read_unit_profiles(wb) -> dict[str, dict[str, Any]]:
    ws = _find_sheet(
        wb,
        "شناسنامه یگان‌ها",
        "شناسنامه یگان ها",
        "مشخصات یگان‌ها",
        "unit_profiles",
        "unit profiles",
    )
    if not ws:
        return {}
    headers, _ = _header_map(ws)
    profiles: dict[str, dict[str, Any]] = {}
    for row in _read_rows(ws):
        code = _cell(
            row, headers, "resource_code", "unit_code", "کد_مرجع_یگان", "کد_مرجع"
        )
        if not code:
            continue
        formed_on_raw = _cell_raw(
            row, headers, "formed_on", "formation_date", "تاریخ_تشکیل"
        )
        deactivated_on_raw = _cell_raw(
            row, headers, "deactivated_on", "end_date", "تاریخ_پایان_فعالیت"
        )
        formed_on = (
            formed_on_raw.date().isoformat()
            if isinstance(formed_on_raw, datetime)
            else _norm(formed_on_raw)
        )
        deactivated_on = (
            deactivated_on_raw.date().isoformat()
            if isinstance(deactivated_on_raw, datetime)
            else _norm(deactivated_on_raw)
        )
        profiles[code] = {
            "shortName": _cell(row, headers, "short_name", "عنوان_کوتاه", "نام_کوتاه")
            or None,
            "organizationalAffiliation": _cell(
                row,
                headers,
                "organizational_affiliation",
                "affiliation",
                "وابستگی_سازمانی",
            )
            or None,
            "parentCode": _cell(row, headers, "parent_code", "کد_یگان_بالادست") or None,
            "parentName": _cell(row, headers, "parent_name", "نام_یگان_بالادست")
            or None,
            "formedOn": formed_on or None,
            "deactivatedOn": deactivated_on or None,
            "province": _cell(row, headers, "province", "استان") or None,
            "garrisonCity": _cell(row, headers, "garrison_city", "city", "شهر_استقرار")
            or None,
            "baseName": _cell(
                row, headers, "base_name", "نام_پادگان", "پادگان", "قرارگاه"
            )
            or None,
            "capabilities": _split_ids(
                _cell(row, headers, "capabilities", "توانمندی‌ها", "توانمندی ها")
            ),
            "nominalPersonnelStrength": _cell(
                row, headers, "nominal_strength", "استعداد_اسمی"
            )
            or None,
            "defaultReadiness": _cell(row, headers, "default_readiness", "آمادگی_پایه")
            or None,
            "areaOfOps": _cell(row, headers, "area_of_ops", "محدوده_مسئولیت") or None,
            "commanderName": _cell(row, headers, "commander", "فرمانده") or None,
            "contact": _cell(row, headers, "contact", "اطلاعات_تماس") or None,
            "description": _cell(row, headers, "description", "توضیح", "توضیحات")
            or None,
        }
    return profiles


def parse_excel_workbook(wb) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    errors: list[dict[str, Any]] = []
    now = datetime.now(timezone.utc).isoformat()
    unit_profiles = _read_unit_profiles(wb)
    time_refs = _parse_time_references(wb, errors)

    scenario_ws = _find_sheet(wb, "سناریو", "scenario")
    events_ws = _find_sheet(wb, "حوادث", "events")
    units_ws = _find_sheet(wb, "یگان", "یگان‌ها", "یگان ها", "units", "orbat")
    unit_states_ws = _find_sheet(
        wb,
        "وضعیت‌های زمانی یگان‌ها",
        "وضعیت های زمانی یگان ها",
        "مسیر یگان‌ها",
        "مسیر یگان ها",
        "unit_states",
        "unit states",
        "unit_tracks",
        "unit tracks",
    )
    equip_ws = _find_sheet(wb, "تجهیزات", "equipment")
    pers_ws = _find_sheet(wb, "پرسنل", "personnel")
    features_ws = _find_sheet(wb, "عوارض", "عارضه", "features", "map_features")

    review_ws = _find_sheet(wb, "کالک‌یار")
    is_document_draft = bool(
        review_ws
        and any(
            str(cell.value or "").strip() == "DOCUMENT_REVIEW_DRAFT_V1"
            for row in review_ws.iter_rows()
            for cell in row
        )
    )

    def require_draft_value(ws, row_number, row, headers, keys, message):
        if is_document_draft and not _cell(row, headers, *keys):
            errors.append({"sheet": ws.title, "row": row_number, "message": message})

    if not scenario_ws:
        errors.append(
            {"sheet": "-", "row": 0, "message": "شیت «سناریو» یا scenario یافت نشد"}
        )
        return {}, errors

    sh, _ = _header_map(scenario_ws)
    srows = _read_rows(scenario_ws)
    if not srows:
        errors.append(
            {
                "sheet": scenario_ws.title,
                "row": 0,
                "message": "حداقل یک ردیف داده در شیت سناریو لازم است",
            }
        )
    row1 = srows[0] if srows else tuple()

    if row1:
        require_draft_value(
            scenario_ws,
            2,
            row1,
            sh,
            (
                "start_time",
                "starttime",
                "زمان_شروع",
                "زمان_شروع_میلادی",
                "تاریخ_شروع",
                "start_time_id",
                "شناسه_زمان_شروع",
            ),
            "زمان شروع سناریو در پیش‌نویس کالک‌یار باید از روی سند تکمیل شود",
        )

    name = _cell(row1, sh, "name", "نام", "عنوان")
    if not name:
        errors.append(
            {"sheet": scenario_ws.title, "row": 2, "message": "نام سناریو الزامی است"}
        )

    description = _cell(row1, sh, "description", "توضیح", "توضیحات")
    start_raw = _resolve_row_time(
        row1,
        sh,
        errors,
        scenario_ws.title,
        2,
        "start_time",
        time_refs,
        direct_keys=(
            "start_time",
            "starttime",
            "زمان_شروع",
            "زمان_شروع_میلادی",
            "تاریخ_شروع",
        ),
        reference_keys=("start_time_id", "شناسه_زمان_شروع"),
    )
    time_zone = (
        _cell(row1, sh, "time_zone", "timezone", "منطقه_زمانی", "ناحیه_زمانی") or "UTC"
    )
    sym = _cell(row1, sh, "symbology_standard", "symbology", "استاندارد_نماد") or "app6"

    start_time = start_raw
    if start_time is None and start_raw:
        pass
    elif start_time is None:
        start_time = int(datetime.now(timezone.utc).timestamp() * 1000)

    events: list[dict[str, Any]] = []
    if events_ws:
        eh, _ = _header_map(events_ws)
        for i, er in enumerate(_read_rows(events_ws), start=2):
            require_draft_value(
                events_ws,
                i,
                er,
                eh,
                (
                    "start_time",
                    "starttime",
                    "زمان_شروع",
                    "زمان_شروع_میلادی",
                    "زمان",
                    "تاریخ",
                    "start_time_id",
                    "شناسه_زمان_شروع",
                ),
                "زمان رویداد در پیش‌نویس کالک‌یار باید از روی سند تکمیل شود",
            )
            eid = _cell(er, eh, "id", "شناسه") or f"ev-{uuid.uuid4().hex[:12]}"
            title = _cell(er, eh, "title", "عنوان")
            if not title:
                errors.append(
                    {
                        "sheet": events_ws.title,
                        "row": i,
                        "message": "عنوان رویداد خالی است",
                    }
                )
                continue
            sub = _cell(er, eh, "subtitle", "sub_title", "زیرعنوان")
            st_parsed = _resolve_row_time(
                er,
                eh,
                errors,
                events_ws.title,
                i,
                "start_time",
                time_refs,
                direct_keys=(
                    "start_time",
                    "starttime",
                    "زمان_شروع",
                    "زمان_شروع_میلادی",
                    "زمان",
                    "تاریخ",
                ),
                reference_keys=("start_time_id", "شناسه_زمان_شروع"),
            )
            if st_parsed is None:
                st_parsed = now
            end_parsed = _resolve_row_time(
                er,
                eh,
                errors,
                events_ws.title,
                i,
                "end_time",
                time_refs,
                direct_keys=("end_time", "endtime", "زمان_پایان", "زمان_پایان_میلادی"),
                reference_keys=("end_time_id", "شناسه_زمان_پایان"),
            )
            lon_s = _cell(er, eh, "lon", "longitude", "طول", "طول_جغرافیایی")
            lat_s = _cell(er, eh, "lat", "latitude", "عرض", "عرض_جغرافیایی")
            lon, lat = _validate_lon_lat(lon_s, lat_s, errors, events_ws.title, i)
            ev: dict[str, Any] = {"id": eid, "title": title, "startTime": st_parsed}
            if sub:
                ev["subTitle"] = sub
            if end_parsed is not None:
                ev["endTime"] = end_parsed
            side = _cell(er, eh, "side", "طرف", "جبهه")
            if side:
                ev["side"] = side
            involved_units = _split_ids(
                _cell(er, eh, "unit_ids", "unit_id", "شناسه_یگان", "یگان‌ها")
            )
            involved_equipment = _split_ids(
                _cell(er, eh, "equipment_ids", "equipment_id", "شناسه_تجهیزات")
            )
            if involved_units:
                ev["involvedUnits"] = involved_units
            if involved_equipment:
                ev["involvedEquipment"] = involved_equipment
            if lon is not None and lat is not None:
                ev["where"] = {
                    "type": "geometry",
                    "geometry": {
                        "type": "Point",
                        "coordinates": [lon, lat],
                    },
                }
            events.append(ev)

    sides: list[dict[str, Any]] = []
    raw_units: list[dict[str, Any]] = []
    if units_ws:
        uh, _ = _header_map(units_ws)
        for i, ur in enumerate(_read_rows(units_ws), start=2):
            require_draft_value(
                units_ws,
                i,
                ur,
                uh,
                ("side", "طرف", "جبهه"),
                "طرف یگان در پیش‌نویس کالک‌یار باید تعیین شود",
            )
            require_draft_value(
                units_ws,
                i,
                ur,
                uh,
                ("unit_type", "type", "نوع", "نوع_واحد", "نوع_نماد"),
                "نوع یگان در پیش‌نویس کالک‌یار باید تعیین شود",
            )
            require_draft_value(
                units_ws,
                i,
                ur,
                uh,
                ("echelon", "unit_echelon", "رده", "رده_یگان", "رده_سازمانی"),
                "رده یگان در پیش‌نویس کالک‌یار باید تعیین شود",
            )
            require_draft_value(
                units_ws,
                i,
                ur,
                uh,
                ("resource_code", "unit_code", "کد_مرجع_یگان", "کد_مرجع"),
                "کد مرجع یگان برای تطبیق کاتالوگ باید تعیین شود",
            )
            uid = _cell(ur, uh, "id", "شناسه")
            uname = _cell(ur, uh, "name", "نام")
            if not uname:
                errors.append(
                    {"sheet": units_ws.title, "row": i, "message": "نام یگان خالی است"}
                )
                continue
            if not uid:
                uid = f"u-{uuid.uuid4().hex[:10]}"
            resource_code = (
                _cell(ur, uh, "resource_code", "unit_code", "کد_مرجع_یگان", "کد_مرجع")
                or uid
            )
            profile = unit_profiles.get(resource_code, {})
            parent = _cell(
                ur, uh, "parent_id", "parentid", "والد", "id_والد", "شناسه_والد"
            )
            side_raw = _cell(ur, uh, "side", "طرف", "جبهه") or "friend"
            sk = _norm_key(side_raw)
            std_id = SIDE_MAP.get(sk) or SIDE_MAP.get(side_raw.strip().lower()) or "3"
            utype = _cell(
                ur,
                uh,
                "symbol_type",
                "unit_type",
                "type",
                "نوع",
                "نوع_واحد",
                "نوع_نماد",
            )
            echelon = _cell(
                ur, uh, "echelon", "unit_echelon", "رده", "رده_یگان", "رده_سازمانی"
            )
            symbol_status = _cell(
                ur, uh, "symbol_status", "status", "وضعیت_نماد", "وضعیت"
            )
            advanced_sidc = _clean_sidc(
                _cell(
                    ur,
                    uh,
                    "advanced_sidc",
                    "sidc",
                    "نماد",
                    "کد_نماد",
                    "کد_نماد_پیشرفته",
                )
            )
            if advanced_sidc and not re.fullmatch(r"\d{20}", advanced_sidc):
                errors.append(
                    {
                        "sheet": units_ws.title,
                        "row": i,
                        "message": "کد نماد پیشرفته باید دقیقاً ۲۰ رقم باشد؛ نماد از انتخاب‌های فارسی ساخته شد",
                    }
                )
                advanced_sidc = ""
            sidc = advanced_sidc or _build_unit_sidc(
                std_id, utype, echelon, symbol_status
            )
            t_parsed = _resolve_row_time(
                ur,
                uh,
                errors,
                units_ws.title,
                i,
                "time",
                time_refs,
                direct_keys=("time", "t", "زمان", "زمان_میلادی"),
                reference_keys=("time_id", "شناسه_زمان"),
            )
            lon_s = _cell(ur, uh, "lon", "longitude", "طول")
            lat_s = _cell(ur, uh, "lat", "latitude", "عرض")
            lon, lat = _validate_lon_lat(lon_s, lat_s, errors, units_ws.title, i)
            st_list: list[dict[str, Any]] = []
            if t_parsed or lon is not None:
                loc = [lon, lat] if lon is not None and lat is not None else None
                st_list.append(
                    {
                        "t": t_parsed or now,
                        "location": loc,
                        "id": f"excel-state-{uid}-initial",
                    }
                )
            raw_units.append(
                {
                    "id": uid,
                    "name": uname,
                    "resourceCode": resource_code,
                    "shortName": profile.get("shortName")
                    or _cell(ur, uh, "short_name", "عنوان_کوتاه", "نام_کوتاه")
                    or None,
                    "description": profile.get("description")
                    or _cell(ur, uh, "description", "توضیح", "توضیحات")
                    or None,
                    "parent_id": parent or None,
                    "parentCode": profile.get("parentCode"),
                    "parentName": profile.get("parentName"),
                    "standard_identity": std_id,
                    "sidc": sidc,
                    "echelon": echelon or None,
                    "unitType": utype or None,
                    "organizationalAffiliation": profile.get(
                        "organizationalAffiliation"
                    ),
                    "formedOn": profile.get("formedOn"),
                    "deactivatedOn": profile.get("deactivatedOn"),
                    "province": profile.get("province")
                    or _cell(ur, uh, "province", "استان")
                    or None,
                    "garrisonCity": profile.get("garrisonCity")
                    or _cell(ur, uh, "garrison_city", "city", "شهر_استقرار")
                    or None,
                    "baseName": profile.get("baseName")
                    or _cell(ur, uh, "base_name", "نام_پادگان", "پادگان", "قرارگاه")
                    or None,
                    "capabilities": profile.get("capabilities")
                    or _split_ids(
                        _cell(ur, uh, "capabilities", "توانمندی‌ها", "توانمندی ها")
                    ),
                    "nominalPersonnelStrength": profile.get("nominalPersonnelStrength")
                    or _cell(ur, uh, "nominal_strength", "استعداد_اسمی")
                    or None,
                    "defaultReadiness": profile.get("defaultReadiness")
                    or _cell(ur, uh, "default_readiness", "آمادگی_پایه")
                    or None,
                    "areaOfOps": profile.get("areaOfOps")
                    or _cell(ur, uh, "area_of_ops", "محدوده_مسئولیت")
                    or None,
                    "commanderName": profile.get("commanderName")
                    or _cell(ur, uh, "commander", "فرمانده")
                    or None,
                    "contact": profile.get("contact")
                    or _cell(ur, uh, "contact", "اطلاعات_تماس")
                    or None,
                    "state": st_list,
                    "subUnits": [],
                    "equipment": [],
                    "personnel": [],
                }
            )

        unit_by_id = {str(unit["id"]): unit for unit in raw_units}
        if unit_states_ws:
            state_headers, _ = _header_map(unit_states_ws)
            for row_number, state_row in enumerate(
                _read_rows(unit_states_ws), start=2
            ):
                unit_id = _cell(
                    state_row,
                    state_headers,
                    "unit_id",
                    "شناسه_یگان",
                    "یگان",
                )
                if not unit_id:
                    errors.append(
                        {
                            "sheet": unit_states_ws.title,
                            "row": row_number,
                            "message": "شناسه یگان برای وضعیت زمانی الزامی است",
                        }
                    )
                    continue
                unit = unit_by_id.get(unit_id)
                if unit is None:
                    errors.append(
                        {
                            "sheet": unit_states_ws.title,
                            "row": row_number,
                            "message": f"یگان با شناسه {unit_id} یافت نشد",
                        }
                    )
                    continue

                state_time = _resolve_row_time(
                    state_row,
                    state_headers,
                    errors,
                    unit_states_ws.title,
                    row_number,
                    "time",
                    time_refs,
                    direct_keys=("time", "t", "زمان", "زمان_میلادی", "زمان_مقصد"),
                    reference_keys=(
                        "time_id",
                        "شناسه_زمان",
                        "شناسه_زمان_مقصد",
                    ),
                )
                if state_time is None:
                    errors.append(
                        {
                            "sheet": unit_states_ws.title,
                            "row": row_number,
                            "message": "زمان مقصد برای وضعیت یگان الزامی است",
                        }
                    )

                lon_raw = _cell(
                    state_row,
                    state_headers,
                    "lon",
                    "longitude",
                    "طول",
                    "طول_جغرافیایی",
                )
                lat_raw = _cell(
                    state_row,
                    state_headers,
                    "lat",
                    "latitude",
                    "عرض",
                    "عرض_جغرافیایی",
                )
                lon, lat = _validate_lon_lat(
                    lon_raw,
                    lat_raw,
                    errors,
                    unit_states_ws.title,
                    row_number,
                )
                if lon is None or lat is None or state_time is None:
                    if lon is None and lat is None and not lon_raw and not lat_raw:
                        errors.append(
                            {
                                "sheet": unit_states_ws.title,
                                "row": row_number,
                                "message": "مختصات وضعیت زمانی یگان الزامی است",
                            }
                        )
                    continue

                transfer_raw = _norm_key(
                    _cell(
                        state_row,
                        state_headers,
                        "transfer_mode",
                        "movement_mode",
                        "روش_انتقال",
                        "نوع_حرکت",
                    )
                    or "پیوسته"
                )
                continuous_values = {
                    "پیوسته",
                    "حرکت_پیوسته",
                    "continuous",
                    "interpolate",
                    "بله",
                    "yes",
                    "true",
                    "1",
                }
                jump_values = {
                    "پرش",
                    "جابجایی_آنی",
                    "جابه_جایی_آنی",
                    "jump",
                    "instant",
                    "خیر",
                    "no",
                    "false",
                    "0",
                }
                if transfer_raw not in continuous_values | jump_values:
                    errors.append(
                        {
                            "sheet": unit_states_ws.title,
                            "row": row_number,
                            "message": "روش انتقال باید «پیوسته» یا «پرش» باشد",
                        }
                    )
                    continue

                path_raw = _norm_key(
                    _cell(
                        state_row,
                        state_headers,
                        "path_mode",
                        "نوع_مسیر",
                    )
                    or "مستقیم"
                )
                path_modes = {
                    "مستقیم": "straight",
                    "straight": "straight",
                    "خطی": "straight",
                    "منحنی": "curved",
                    "curved": "curved",
                }
                if path_raw not in path_modes:
                    errors.append(
                        {
                            "sheet": unit_states_ws.title,
                            "row": row_number,
                            "message": "نوع مسیر باید «مستقیم» یا «منحنی» باشد",
                        }
                    )
                    continue

                movement_start = _resolve_row_time(
                    state_row,
                    state_headers,
                    errors,
                    unit_states_ws.title,
                    row_number,
                    "movement_start_time",
                    time_refs,
                    direct_keys=(
                        "movement_start_time",
                        "زمان_شروع_حرکت",
                        "زمان_شروع_حرکت_میلادی",
                    ),
                    reference_keys=(
                        "movement_start_time_id",
                        "شناسه_زمان_شروع_حرکت",
                    ),
                )
                state_key = _cell(
                    state_row,
                    state_headers,
                    "id",
                    "state_id",
                    "شناسه",
                    "شناسه_وضعیت",
                ) or str(row_number - 1)
                state_id = f"excel-state-{unit_id}-{state_key}"
                state: dict[str, Any] = {
                    "id": state_id,
                    "t": state_time,
                    "location": [lon, lat],
                    "interpolate": transfer_raw in continuous_values,
                    "pathMode": path_modes[path_raw],
                }
                if movement_start is not None:
                    if _scenario_time_sort_key(movement_start) > _scenario_time_sort_key(
                        state_time
                    ):
                        errors.append(
                            {
                                "sheet": unit_states_ws.title,
                                "row": row_number,
                                "message": "زمان شروع حرکت نمی‌تواند بعد از زمان مقصد باشد",
                            }
                        )
                        continue
                    state["viaStartTime"] = movement_start
                unit.setdefault("state", []).append(state)

            for unit in raw_units:
                unit["state"] = sorted(
                    unit.get("state") or [],
                    key=lambda state: _scenario_time_sort_key(state.get("t")),
                )

        by_id = {u["id"]: u for u in raw_units}
        for u in raw_units:
            parent_unit = by_id.get(u.get("parent_id"))
            if parent_unit:
                u["parentCode"] = parent_unit.get("resourceCode") or parent_unit.get(
                    "id"
                )
                u["parentName"] = parent_unit.get("name")
        roots: list[dict[str, Any]] = []
        for u in raw_units:
            pid = u.get("parent_id")
            if pid and pid in by_id:
                by_id[pid].setdefault("subUnits", []).append(u)
            else:
                roots.append(u)
        for u in raw_units:
            u.pop("parent_id", None)
        sides_map: dict[str, list[dict[str, Any]]] = {}
        for r in roots:
            sid = r["standard_identity"]
            sides_map.setdefault(sid, []).append(r)
        for std_id, subs in sides_map.items():
            sides.append(
                {
                    "id": f"excel-side-{std_id}",
                    "name": "خودی"
                    if std_id == "3"
                    else "دشمن"
                    if std_id == "6"
                    else "خنثی"
                    if std_id == "4"
                    else "نامشخص",
                    "standardIdentity": std_id,
                    "groups": [
                        {
                            "id": f"excel-group-{std_id}",
                            "name": "یگان‌های واردشده",
                            "subUnits": subs,
                        }
                    ],
                }
            )

    unit_by_id = {u["id"]: u for u in raw_units}
    equipment_by_name: dict[str, dict[str, Any]] = {}
    independent_equipment_features: list[dict[str, Any]] = []
    if equip_ws:
        event_identities_by_equipment: dict[str, set[str]] = {}
        scenario_event_identities: set[str] = set()
        for event in events:
            event_identity = SIDE_MAP.get(_norm_key(event.get("side") or ""))
            if not event_identity:
                continue
            scenario_event_identities.add(event_identity)
            for equipment_id in event.get("involvedEquipment", []):
                event_identities_by_equipment.setdefault(equipment_id, set()).add(
                    event_identity
                )
        scenario_event_identity = (
            next(iter(scenario_event_identities))
            if len(scenario_event_identities) == 1
            else "0"
        )
        eh, _ = _header_map(equip_ws)
        for i, er in enumerate(_read_rows(equip_ws), start=2):
            require_draft_value(
                equip_ws,
                i,
                er,
                eh,
                ("type", "نوع"),
                "نوع تجهیز در پیش‌نویس کالک‌یار باید تعیین شود",
            )
            require_draft_value(
                equip_ws,
                i,
                er,
                eh,
                ("quantity", "تعداد", "count"),
                "تعداد تجهیز در پیش‌نویس کالک‌یار باید از روی سند تکمیل شود",
            )
            require_draft_value(
                equip_ws,
                i,
                er,
                eh,
                ("equipment_code", "کد", "equipmentcode"),
                "کد مرجع تجهیز برای تطبیق کاتالوگ باید تعیین شود",
            )
            ename = _cell(er, eh, "name", "نام")
            if not ename:
                continue
            eid = _cell(er, eh, "id", "شناسه") or f"eq-{uuid.uuid4().hex[:10]}"
            equipment_type = _cell(er, eh, "type", "نوع") or "عمومی"
            display_name = _equipment_display_name(ename, equipment_type)
            qty_raw = _cell(er, eh, "quantity", "تعداد", "count") or "1"
            try:
                quantity = int(float(qty_raw))
            except ValueError:
                quantity = 1
                errors.append(
                    {
                        "sheet": equip_ws.title,
                        "row": i,
                        "message": f"تعداد نامعتبر: {qty_raw}",
                    }
                )
            equipment_by_name.setdefault(
                display_name,
                {"name": display_name, "description": equipment_type, "sourceId": eid},
            )
            unit_id = _cell(er, eh, "unit_id", "id_یگان", "شناسه_یگان", "یگان")
            explicit_side = _cell(er, eh, "side", "طرف", "جبهه")
            side_identity = (
                SIDE_MAP.get(_norm_key(explicit_side)) if explicit_side else None
            )
            if not side_identity and unit_id and unit_id in unit_by_id:
                side_identity = unit_by_id[unit_id].get("standard_identity")
            event_identities = event_identities_by_equipment.get(eid, set())
            if not side_identity and len(event_identities) == 1:
                side_identity = next(iter(event_identities))
            side_identity = side_identity or scenario_event_identity

            advanced_sidc = _clean_sidc(
                _cell(
                    er,
                    eh,
                    "advanced_sidc",
                    "sidc",
                    "نماد",
                    "کد_نماد",
                    "کد_نماد_پیشرفته",
                )
            )
            if advanced_sidc and not re.fullmatch(r"\d{20}", advanced_sidc):
                errors.append(
                    {
                        "sheet": equip_ws.title,
                        "row": i,
                        "message": "کد نماد پیشرفته تجهیزات باید دقیقاً ۲۰ رقم باشد؛ نماد از نوع و طرف ساخته شد",
                    }
                )
                advanced_sidc = ""
            equipment_sidc = advanced_sidc or _build_equipment_sidc(
                side_identity, equipment_type, ename
            )
            if unit_id and unit_id in unit_by_id:
                attached = unit_by_id[unit_id].setdefault("equipment", [])
                current = next(
                    (item for item in attached if item.get("name") == display_name),
                    None,
                )
                if current:
                    current["count"] = int(current.get("count") or 0) + quantity
                else:
                    attached.append(
                        {
                            "name": display_name,
                            "count": quantity,
                            "description": equipment_type,
                        }
                    )

            lon_s = _cell(er, eh, "lon", "longitude", "طول", "طول_جغرافیایی")
            lat_s = _cell(er, eh, "lat", "latitude", "عرض", "عرض_جغرافیایی")
            lon, lat = _validate_lon_lat(lon_s, lat_s, errors, equip_ws.title, i)
            start_parsed = _resolve_row_time(
                er,
                eh,
                errors,
                equip_ws.title,
                i,
                "start_time",
                time_refs,
                direct_keys=(
                    "start_time",
                    "starttime",
                    "زمان_شروع",
                    "زمان_شروع_میلادی",
                    "زمان",
                ),
                reference_keys=("start_time_id", "شناسه_زمان_شروع"),
            )
            if lon is not None and lat is not None:
                meta: dict[str, Any] = {"type": "Point", "name": display_name}
                if start_parsed is not None:
                    meta["visibleFromT"] = start_parsed
                equipment_style = _feature_style("#B45309", "square")
                if equipment_sidc:
                    equipment_style["militarySymbolSidc"] = equipment_sidc
                independent_equipment_features.append(
                    {
                        "type": "Feature",
                        "id": f"equipment-{eid}",
                        "geometry": {"type": "Point", "coordinates": [lon, lat]},
                        "properties": {
                            "equipmentId": eid,
                            "unitId": unit_id or None,
                            "quantity": quantity,
                            "equipmentType": equipment_type,
                        },
                        "meta": meta,
                        "style": equipment_style,
                    }
                )

    personnel_by_name: dict[str, dict[str, Any]] = {}
    if pers_ws:
        ph, _ = _header_map(pers_ws)
        for i, pr in enumerate(_read_rows(pers_ws), start=2):
            require_draft_value(
                pers_ws,
                i,
                pr,
                ph,
                ("personal_code", "کد_پرسنلی", "personalcode"),
                "کد مرجع شخص برای تطبیق کاتالوگ باید تعیین شود",
            )
            fn = _cell(pr, ph, "first_name", "نام")
            ln = _cell(pr, ph, "last_name", "نام_خانوادگی", "فامیلی")
            full = _cell(pr, ph, "full_name", "نام_کامل")
            if not fn and not ln and full:
                parts = full.split(maxsplit=1)
                fn = parts[0]
                ln = parts[1] if len(parts) > 1 else ""
            if not fn and not ln:
                continue
            pid = _cell(pr, ph, "id", "شناسه") or f"p-{uuid.uuid4().hex[:10]}"
            full_name = " ".join(part for part in (fn, ln) if part).strip()
            rank = _cell(pr, ph, "rank", "درجه")
            specialty = _cell(pr, ph, "specialty", "position", "تخصص", "سمت")
            description_parts = [part for part in (rank, specialty) if part]
            description = "، ".join(description_parts)
            personnel_by_name.setdefault(
                full_name,
                {"name": full_name, "description": description, "sourceId": pid},
            )
            unit_id = _cell(pr, ph, "unit_id", "id_یگان", "شناسه_یگان", "یگان")
            if unit_id and unit_id in unit_by_id:
                unit_by_id[unit_id].setdefault("personnel", []).append(
                    {"name": full_name, "count": 1, "description": description}
                )

    imported_features: list[dict[str, Any]] = []
    if features_ws:
        fh, _ = _header_map(features_ws)
        for i, fr in enumerate(_read_rows(features_ws), start=2):
            require_draft_value(
                features_ws,
                i,
                fr,
                fh,
                ("type", "نوع"),
                "نوع عارضه در پیش‌نویس کالک‌یار باید تعیین شود",
            )
            feature_id = (
                _cell(fr, fh, "id", "شناسه") or f"feature-{uuid.uuid4().hex[:10]}"
            )
            feature_name = _cell(fr, fh, "name", "نام", "عنوان")
            if not feature_name:
                errors.append(
                    {
                        "sheet": features_ws.title,
                        "row": i,
                        "message": "نام عارضه خالی است",
                    }
                )
                continue
            lon_s = _cell(fr, fh, "lon", "longitude", "طول", "طول_جغرافیایی")
            lat_s = _cell(fr, fh, "lat", "latitude", "عرض", "عرض_جغرافیایی")
            lon, lat = _validate_lon_lat(lon_s, lat_s, errors, features_ws.title, i)
            if lon is None or lat is None:
                errors.append(
                    {
                        "sheet": features_ws.title,
                        "row": i,
                        "message": "مختصات عارضه الزامی است",
                    }
                )
                continue
            radius_raw = _cell(fr, fh, "radius_m", "radius", "شعاع", "شعاع_متر")
            radius: float | None = None
            if radius_raw:
                try:
                    radius = float(radius_raw)
                    if radius <= 0:
                        raise ValueError
                except ValueError:
                    errors.append(
                        {
                            "sheet": features_ws.title,
                            "row": i,
                            "message": "شعاع باید عددی بزرگ‌تر از صفر باشد",
                        }
                    )
                    radius = None
            start_parsed = _resolve_row_time(
                fr,
                fh,
                errors,
                features_ws.title,
                i,
                "start_time",
                time_refs,
                direct_keys=(
                    "start_time",
                    "starttime",
                    "زمان_شروع",
                    "زمان_شروع_میلادی",
                ),
                reference_keys=("start_time_id", "شناسه_زمان_شروع"),
            )
            end_parsed = _resolve_row_time(
                fr,
                fh,
                errors,
                features_ws.title,
                i,
                "end_time",
                time_refs,
                direct_keys=("end_time", "endtime", "زمان_پایان", "زمان_پایان_میلادی"),
                reference_keys=("end_time_id", "شناسه_زمان_پایان"),
            )
            meta: dict[str, Any] = {
                "type": "Circle" if radius else "Point",
                "name": feature_name,
            }
            if radius:
                meta["radius"] = radius
            if start_parsed is not None:
                meta["visibleFromT"] = start_parsed
            if end_parsed is not None:
                meta["visibleUntilT"] = end_parsed
            imported_features.append(
                {
                    "type": "Feature",
                    "id": feature_id,
                    "geometry": {"type": "Point", "coordinates": [lon, lat]},
                    "properties": {
                        "category": _cell(fr, fh, "type", "نوع") or "عمومی",
                        "eventId": _cell(fr, fh, "event_id", "شناسه_رویداد") or None,
                    },
                    "meta": meta,
                    "style": _feature_style("#2563EB"),
                }
            )

    layers: list[dict[str, Any]] = []
    if imported_features:
        layers.append(
            {
                "id": "excel-imported-features",
                "name": "عوارض واردشده",
                "features": imported_features,
            }
        )
    if independent_equipment_features:
        layers.append(
            {
                "id": "excel-independent-equipment",
                "name": "تجهیزات مختصات‌دار",
                "features": independent_equipment_features,
            }
        )

    storyboard_scenes = []
    for order, event in enumerate(events):
        storyboard_scenes.append(
            {
                "id": f"excel-scene-{event['id']}",
                "title": event.get("title") or "",
                "body": event.get("subTitle") or "",
                "startTime": event.get("startTime"),
                "order": order,
                "linkedEventId": event.get("id"),
                "camera": {"type": "eventWhere", "maxZoom": 15}
                if event.get("where")
                else {"type": "none"},
            }
        )

    scenario_id = str(uuid.uuid4())
    content: dict[str, Any] = {
        "id": scenario_id,
        "type": "ORBAT-mapper",
        "version": "0.40.0",
        "name": name or "Imported scenario",
        "description": description or "",
        "startTime": start_time,
        "timeZone": time_zone,
        "symbologyStandard": sym,
        "sides": sides,
        "events": events,
        "layers": layers,
        "mapLayers": [],
        "equipment": list(equipment_by_name.values()),
        "personnel": list(personnel_by_name.values()),
        "storyboard": {
            "enabled": bool(storyboard_scenes),
            "scenes": storyboard_scenes,
            "settings": {
                "defaultSceneDurationMs": 5000,
                "autoDuration": True,
                "showMode": "cinematic",
            },
        },
        "supplyCategories": [],
        "settings": {
            "rangeRingGroups": [],
            "statuses": [],
            "supplyClasses": [
                {"name": "Class I"},
                {"name": "Class II"},
                {"name": "Class III"},
                {"name": "Class IV"},
                {"name": "Class V"},
            ],
            "supplyUoMs": [
                {"name": "Kilogram", "code": "KG", "type": "weight"},
                {"name": "Each", "code": "EA", "type": "quantity"},
            ],
            "map": {"baseMapId": "osm"},
        },
        "meta": {"createdDate": now, "lastModifiedDate": now},
        "metadata": {"source": "excel-import"},
    }
    return content, errors


def build_template_workbook(*, include_example: bool = False) -> Workbook:
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "سناریو"
    ws0.append(["نام", "توضیحات", "شناسه_زمان_شروع", "ناحیه_زمانی", "استاندارد_نماد"])
    ws0.append(
        ["نمونه عملیات", "پیش از ورود ویرایش شود", "زمان-001", "Asia/Tehran", "app6"]
    )

    ws1 = wb.create_sheet("حوادث")
    ws1.append(
        [
            "شناسه",
            "عنوان",
            "زیرعنوان",
            "شناسه_زمان_شروع",
            "شناسه_زمان_پایان",
            "طرف",
            "شناسه_یگان",
            "شناسه_تجهیزات",
            "طول_جغرافیایی",
            "عرض_جغرافیایی",
        ]
    )
    ws1.append(
        [
            "ev-1",
            "شروع عملیات",
            "نمونه",
            "زمان-001",
            "",
            "خودی",
            "u3",
            "eq1",
            51.4,
            35.7,
        ]
    )

    ws2 = wb.create_sheet("یگان‌ها")
    ws2.append(
        [
            "شناسه",
            "نام",
            "طرف",
            "رده_یگان",
            "نوع_نماد",
            "شناسه_والد",
            "وضعیت_نماد",
            "شناسه_زمان",
            "طول",
            "عرض",
            "کد_نماد_پیشرفته",
            "کد_مرجع_یگان",
        ]
    )

    ws_unit_states = wb.create_sheet("وضعیت‌های زمانی یگان‌ها")
    ws_unit_states.append(
        [
            "شناسه_وضعیت",
            "شناسه_یگان",
            "شناسه_زمان_مقصد",
            "طول",
            "عرض",
            "روش_انتقال",
            "نوع_مسیر",
            "شناسه_زمان_شروع_حرکت",
        ]
    )
    ws_unit_states.append(
        [
            "حرکت-u3-01",
            "u3",
            "زمان-002",
            51.46,
            35.74,
            "پیوسته",
            "مستقیم",
            "زمان-001",
        ]
    )
    ws2.append(
        [
            "u1",
            "لشکر ۱۶",
            "خودی",
            "لشکر",
            "پیاده‌نظام",
            "",
            "حاضر",
            "",
            "",
            "",
            "",
            "UNIT-16",
        ]
    )
    ws2.append(
        [
            "u2",
            "تیپ ۱",
            "خودی",
            "تیپ",
            "پیاده‌نظام مکانیزه",
            "u1",
            "حاضر",
            "",
            "",
            "",
            "",
            "UNIT-16-BDE-1",
        ]
    )
    ws2.append(
        [
            "u3",
            "گردان ۲۳۲",
            "خودی",
            "گردان / اسکادران",
            "زرهی",
            "u2",
            "حاضر",
            "زمان-001",
            51.41,
            35.71,
            "",
            "UNIT-232",
        ]
    )
    ws2.append(
        [
            "d1",
            "لشکر ۵",
            "دشمن",
            "لشکر",
            "زرهی",
            "",
            "حاضر",
            "",
            "",
            "",
            "",
            "HOSTILE-DIV-5",
        ]
    )
    ws2.append(
        [
            "d2",
            "تیپ ۱۲",
            "دشمن",
            "تیپ",
            "پیاده‌نظام",
            "d1",
            "حاضر",
            "زمان-001",
            51.60,
            35.80,
            "",
            "HOSTILE-BDE-12",
        ]
    )

    ws_unit_profiles = wb.create_sheet("شناسنامه یگان‌ها")
    ws_unit_profiles.append(
        [
            "کد_مرجع_یگان",
            "عنوان_کوتاه",
            "وابستگی_سازمانی",
            "کد_یگان_بالادست",
            "نام_یگان_بالادست",
            "استان",
            "شهر_استقرار",
            "نام_پادگان",
            "توانمندی‌ها",
            "استعداد_اسمی",
            "آمادگی_پایه",
            "محدوده_مسئولیت",
            "فرمانده",
            "اطلاعات_تماس",
            "تاریخ_تشکیل",
            "تاریخ_پایان_فعالیت",
            "توضیحات",
        ]
    )
    ws_unit_profiles.append(
        [
            "UNIT-16",
            "لشکر ۱۶",
            "نیروی زمینی",
            "",
            "",
            "",
            "",
            "",
            "فرماندهی، پشتیبانی",
            "",
            "بالا",
            "",
            "",
            "",
            "1981-09-22",
            "",
            "",
        ]
    )
    ws_unit_profiles.append(
        [
            "UNIT-16-BDE-1",
            "تیپ ۱",
            "نیروی زمینی",
            "UNIT-16",
            "لشکر ۱۶",
            "",
            "",
            "",
            "مکانیزه، ضدزره",
            "",
            "بالا",
            "",
            "",
            "",
            "",
            "",
            "",
        ]
    )
    ws_unit_profiles.append(
        [
            "UNIT-232",
            "گردان ۲۳۲",
            "نیروی زمینی",
            "UNIT-16-BDE-1",
            "تیپ ۱",
            "تهران",
            "تهران",
            "پادگان نمونه",
            "زرهی، راکت‌انداز، تحرک بالا",
            "۵۰۰",
            "متوسط",
            "محور عملیاتی نمونه",
            "",
            "",
            "",
            "",
            "نمونه قابل ویرایش",
        ]
    )
    ws_unit_profiles.append(
        [
            "HOSTILE-DIV-5",
            "لشکر ۵",
            "نیروی زمینی دشمن",
            "",
            "",
            "",
            "",
            "",
            "زرهی",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
        ]
    )
    ws_unit_profiles.append(
        [
            "HOSTILE-BDE-12",
            "تیپ ۱۲",
            "نیروی زمینی دشمن",
            "HOSTILE-DIV-5",
            "لشکر ۵",
            "",
            "",
            "",
            "پیاده‌نظام",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
        ]
    )

    ws3 = wb.create_sheet("تجهیزات")
    ws3.append(
        [
            "شناسه",
            "نام",
            "نوع",
            "تعداد",
            "شناسه_یگان",
            "شناسه_زمان_شروع",
            "طول",
            "عرض",
            "طرف",
            "کد_نماد_پیشرفته",
        ]
    )
    ws3.append(["eq1", "تانک", "تی ۶۴", 2, "u3", "زمان-001", 51.42, 35.72, "خودی", ""])

    ws4 = wb.create_sheet("پرسنل")
    ws4.append(["شناسه", "نام", "نام_خانوادگی", "درجه", "تخصص", "کد_ملی", "شناسه_یگان"])
    ws4.append(["p1", "علی", "رضایی", "سروان", "فرمانده", "00‌11111111", "u3"])

    ws5 = wb.create_sheet("عوارض")
    ws5.append(
        [
            "شناسه",
            "نام",
            "نوع",
            "طول",
            "عرض",
            "شعاع_متر",
            "شناسه_زمان_شروع",
            "شناسه_زمان_پایان",
            "شناسه_رویداد",
        ]
    )
    ws5.append(
        [
            "عارضه-نمونه-۱",
            "شهر نمونه؛ پیش از ورود ویرایش شود",
            "شهر",
            51.39,
            35.69,
            3000,
            "زمان-001",
            "",
            "ev-1",
        ]
    )

    ws_time = wb.create_sheet("زمان‌بندی")
    ws_time.append(
        [
            "شناسه_زمان",
            "سال_شمسی",
            "ماه_شمسی",
            "روز_شمسی",
            "ساعت",
            "دقیقه",
            "ناحیه_زمانی",
            "تاریخ_و_ساعت_میلادی",
        ]
    )
    ws_time.append(["زمان-001", 1405, 6, 7, 8, 0, "Asia/Tehran", None])
    ws_time.append(["زمان-002", 1405, 6, 7, 10, 0, "Asia/Tehran", None])

    ws_guide = wb.create_sheet("راهنما")
    ws_guide.append(["بخش", "کار کاربر", "نمونه", "نتیجه در کالک‌نگار"])
    ws_guide.append(
        [
            "مسیر پایه یگان",
            "برای هر مقصد یک ردیف زمان‌دار وارد کنید؛ موقعیت اولیه در برگه یگان‌ها باقی می‌ماند.",
            "u3، زمان-002، پیوسته، مستقیم",
            "حرکت پایه ساخته می‌شود و اصلاح دقیق مسیر در کالک‌نگار انجام می‌شود.",
        ]
    )
    ws_guide.append(
        [
            "زمان‌بندی",
            "اجزای تاریخ و ساعت را از فهرست انتخاب کنید.",
            "۱۴۰۵/۰۶/۰۷",
            "تاریخ میلادی محاسبه و هنگام ورود استاندارد می‌شود.",
        ]
    )
    ws_guide.append(
        [
            "نماد یگان",
            "نوع، رده و وضعیت نماد را از فهرست انتخاب کنید.",
            "زرهی، گردان، حاضر",
            "کد نماد هنگام ورود به‌صورت خودکار ساخته می‌شود.",
        ]
    )
    ws_guide.append(
        [
            "نماد تجهیزات",
            "نوع و طرف تجهیزات را وارد کنید؛ کد پیشرفته اختیاری است.",
            "تی ۶۴، دشمن",
            "نماد تانک، نفربر، بلدوزر، توپ، خودرو، هواپیما یا بالگرد هنگام ورود ساخته می‌شود.",
        ]
    )
    ws_guide.append(
        [
            "کد مرجع یگان",
            "برای یک یگان واقعی در همه عملیات‌ها یک کد ثابت وارد کنید.",
            "UNIT-232",
            "سابقه همان یگان در چند سناریو به هم متصل می‌شود.",
        ]
    )
    ws_guide.append(
        [
            "اطلاعات پایه یگان",
            "وابستگی، یگان بالادست، استقرار، توانمندی، استعداد و تاریخچه را در شناسنامه وارد کنید.",
            "نیروی زمینی، تهران، زرهی، ۵۰۰",
            "شناسنامه یگان در مدیریت منابع تکمیل می‌شود.",
        ]
    )
    ws_guide.append(
        [
            "کد نماد پیشرفته",
            "فقط برای نماد خارج از فهرست آن را وارد کنید.",
            "۲۰ رقم",
            "این کد بر انتخاب‌های فارسی اولویت دارد.",
        ]
    )
    ws_guide.append(
        [
            "عوارض",
            "مرکز و شعاع اختیاری را وارد کنید.",
            "۳۰۰۰ متر",
            "عارضه اولیه روی نقشه ساخته می‌شود.",
        ]
    )

    ws_guide.append(
        [
            "نوع فایل",
            (
                "نمونه تکمیل‌شده آموزشی است؛ برای ورود واقعی از قالب خالی استفاده کنید."
                if include_example
                else "قالب خالی برای ورود واقعی است؛ داده‌های سناریو را از ردیف دوم وارد کنید."
            ),
            "دارای داده نمونه" if include_example else "بدون داده نمونه",
            (
                "این فایل را بدون جایگزینی داده‌های نمونه به‌عنوان سناریوی واقعی وارد نکنید."
                if include_example
                else "پس از تکمیل، ابتدا پیش‌نمایش ورود را بررسی و سپس ثبت را تأیید کنید."
            ),
        ]
    )

    ws_symbol_guide = wb.create_sheet("راهنمای نمادها")
    ws_symbol_guide.append(["بخش", "عنوان", "کاربرد پیشنهادی", "کد پایه نمونه"])
    for echelon_name in MAIN_ECHELONS:
        echelon_code = ECHELON_MAP[_norm_key(echelon_name)]
        ws_symbol_guide.append(
            [
                "رده سازمانی",
                echelon_name,
                "برای ساخت اسکلت اصلی آرایش نبرد",
                f"10‌031000{echelon_code}1211000000",
            ]
        )
    for name, main_icon, use in [
        ("پیاده‌نظام", "121100", "یگان‌های پیاده"),
        ("پیاده‌نظام مکانیزه", "121102", "پیاده مجهز به نفربر یا خودروی رزمی"),
        ("شناسایی", "121300", "یگان‌های شناسایی"),
        ("توپخانه", "130300", "توپخانه صحرایی و آتش پشتیبانی"),
        ("زرهی", "120500", "یگان‌های تانک و زرهی"),
        ("پشتیبانی خدمات رزمی", "160600", "تدارکات و پشتیبانی"),
        ("پدافند هوایی", "130100", "یگان‌های پدافند هوایی"),
        ("مهندسی رزمی", "140700", "یگان‌های مهندسی رزمی"),
    ]:
        ws_symbol_guide.append(["نوع یگان", name, use, f"10‌03100016{main_icon}0000"])

    ws_lists = wb.create_sheet("فهرست‌های انتخاب")
    ws_lists.append(
        [
            "سال",
            "ماه",
            "روز",
            "ساعت",
            "دقیقه",
            "ناحیه_زمانی",
            "نوع_نماد",
            "رده_یگان",
            "وضعیت_نماد",
            "روش_انتقال",
            "نوع_مسیر",
            "کلید_تاریخ_شمسی",
            "تاریخ_میلادی",
        ]
    )
    years = list(range(1350, 1431))
    for index, value in enumerate(years, start=2):
        ws_lists.cell(index, 1, value)
    for index, value in enumerate(range(1, 13), start=2):
        ws_lists.cell(index, 2, value)
    for index, value in enumerate(range(1, 32), start=2):
        ws_lists.cell(index, 3, value)
    for index, value in enumerate(range(24), start=2):
        ws_lists.cell(index, 4, value)
    for index, value in enumerate(range(60), start=2):
        ws_lists.cell(index, 5, value)
    for index, value in enumerate(
        ["Asia/Tehran", "UTC", "Asia/Baghdad", "Asia/Kabul", "Asia/Dubai"], start=2
    ):
        ws_lists.cell(index, 6, value)
    for index, value in enumerate(
        [
            "پیاده‌نظام",
            "پیاده‌نظام مکانیزه",
            "شناسایی",
            "توپخانه",
            "زرهی",
            "پشتیبانی خدمات رزمی",
            "پدافند هوایی",
            "مهندسی رزمی",
        ],
        start=2,
    ):
        ws_lists.cell(index, 7, value)
    for index, value in enumerate(MAIN_ECHELONS, start=2):
        ws_lists.cell(index, 8, value)
    for index, value in enumerate(
        ["حاضر", "برنامه‌ریزی‌شده", "آسیب‌دیده", "نابودشده"], start=2
    ):
        ws_lists.cell(index, 9, value)
    for index, value in enumerate(["پیوسته", "پرش"], start=2):
        ws_lists.cell(index, 10, value)
    for index, value in enumerate(["مستقیم", "منحنی"], start=2):
        ws_lists.cell(index, 11, value)

    calendar_row = 2
    for year in years:
        for month in range(1, 13):
            for day in range(1, _jalali_month_length(year, month) + 1):
                gy, gm, gd = _jalali_to_gregorian(year, month, day)
                ws_lists.cell(calendar_row, 12, f"{year:04d}/{month:02d}/{day:02d}")
                ws_lists.cell(calendar_row, 13, datetime(gy, gm, gd))
                ws_lists.cell(calendar_row, 13).number_format = "yyyy-mm-dd"
                calendar_row += 1

    ws_time["H2"] = (
        f'=IFERROR(INDEX(\'فهرست‌های انتخاب\'!$M$2:$M${calendar_row - 1},MATCH(TEXT(B2,"0000")&"/"&TEXT(C2,"00")&"/"&TEXT(D2,"00"),\'فهرست‌های انتخاب\'!$L$2:$L${calendar_row - 1},0))+TIME(E2,F2,0),"")'
    )
    ws_time["H3"] = ws_time["H2"].value.replace("B2", "B3").replace("C2", "C3").replace("D2", "D3").replace("E2", "E3").replace("F2", "F3")
    ws_time["H3"].number_format = "yyyy-mm-dd hh:mm"
    ws_time["H2"].number_format = "yyyy-mm-dd hh:mm"

    for cell_range, formula in [
        ("B2:B500", "'فهرست‌های انتخاب'!$A$2:$A$82"),
        ("C2:C500", "'فهرست‌های انتخاب'!$B$2:$B$13"),
        ("D2:D500", "'فهرست‌های انتخاب'!$C$2:$C$32"),
        ("E2:E500", "'فهرست‌های انتخاب'!$D$2:$D$25"),
        ("F2:F500", "'فهرست‌های انتخاب'!$E$2:$E$61"),
        ("G2:G500", "'فهرست‌های انتخاب'!$F$2:$F$6"),
    ]:
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws_time.add_data_validation(validation)
        validation.add(cell_range)

    for sheet, ranges in [
        (ws0, ["C2:C500"]),
        (ws1, ["D2:E500"]),
        (ws2, ["H2:H500"]),
        (ws_unit_states, ["C2:C500", "H2:H500"]),
        (ws3, ["F2:F500"]),
        (ws5, ["G2:H500"]),
    ]:
        time_validation = DataValidation(
            type="list", formula1="'زمان‌بندی'!$A$2:$A$500", allow_blank=True
        )
        sheet.add_data_validation(time_validation)
        for cell_range in ranges:
            time_validation.add(cell_range)

    movement_validation = DataValidation(
        type="list", formula1="'فهرست‌های انتخاب'!$J$2:$J$3", allow_blank=False
    )
    ws_unit_states.add_data_validation(movement_validation)
    movement_validation.add("F2:F500")
    path_validation = DataValidation(
        type="list", formula1="'فهرست‌های انتخاب'!$K$2:$K$3", allow_blank=False
    )
    ws_unit_states.add_data_validation(path_validation)
    path_validation.add("G2:G500")

    for cell_range, formula in [
        ("E2:E500", "'فهرست‌های انتخاب'!$G$2:$G$9"),
        ("D2:D500", "'فهرست‌های انتخاب'!$H$2:$H$10"),
        ("G2:G500", "'فهرست‌های انتخاب'!$I$2:$I$5"),
    ]:
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws2.add_data_validation(validation)
        validation.add(cell_range)

    side_validation = DataValidation(
        type="list", formula1='"خودی,دشمن,خنثی,نامشخص"', allow_blank=True
    )
    ws2.add_data_validation(side_validation)
    side_validation.add("C2:C500")
    equipment_side_validation = DataValidation(
        type="list", formula1='"خودی,دشمن,خنثی,نامشخص"', allow_blank=True
    )
    ws3.add_data_validation(equipment_side_validation)
    equipment_side_validation.add("I2:I500")
    readiness_validation = DataValidation(
        type="list", formula1='"خیلی بالا,بالا,متوسط,پایین"', allow_blank=True
    )
    ws_unit_profiles.add_data_validation(readiness_validation)
    readiness_validation.add("K2:K500")
    friendly_fill = PatternFill("solid", fgColor="DBEAFE")
    hostile_fill = PatternFill("solid", fgColor="FEE2E2")
    ws2.conditional_formatting.add(
        "C2:C500", FormulaRule(formula=['$C2="خودی"'], fill=friendly_fill)
    )
    ws2.conditional_formatting.add(
        "C2:C500", FormulaRule(formula=['$C2="دشمن"'], fill=hostile_fill)
    )
    ws3.conditional_formatting.add(
        "I2:I500", FormulaRule(formula=['$I2="خودی"'], fill=friendly_fill)
    )
    ws3.conditional_formatting.add(
        "I2:I500", FormulaRule(formula=['$I2="دشمن"'], fill=hostile_fill)
    )

    header_fill = PatternFill("solid", fgColor="173A5E")
    time_fill = PatternFill("solid", fgColor="0F766E")
    input_fill = PatternFill("solid", fgColor="FFF8E1")
    output_fill = PatternFill("solid", fgColor="E8F5E9")
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = time_fill if ws.title == "زمان‌بندی" else header_fill
            cell.font = Font(color="FFFFFF", bold=True, name="Arial")
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
        ws.row_dimensions[1].height = 30
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(
                    horizontal="right", vertical="center", wrap_text=True
                )
        ws.auto_filter.ref = ws.dimensions

    for row in ws_time.iter_rows(min_row=2, max_row=500, min_col=2, max_col=7):
        for cell in row:
            cell.fill = input_fill
    for cell in ws_time["H"][1:500]:
        cell.fill = output_fill

    widths = {
        "سناریو": [30, 55, 24, 20, 20],
        "حوادث": [15, 65, 28, 24, 24, 14, 22, 24, 18, 18],
        "یگان‌ها": [15, 28, 14, 24, 28, 18, 20, 20, 16, 16, 28, 22],
        "وضعیت‌های زمانی یگان‌ها": [20, 18, 24, 16, 16, 18, 18, 28],
        "شناسنامه یگان‌ها": [
            22,
            22,
            24,
            24,
            26,
            18,
            20,
            24,
            34,
            18,
            18,
            30,
            24,
            24,
            18,
            20,
            42,
        ],
        "تجهیزات": [15, 24, 22, 12, 18, 24, 16, 16, 14, 28],
        "پرسنل": [15, 18, 24, 18, 26, 20, 18],
        "عوارض": [18, 38, 18, 16, 16, 16, 24, 24, 20],
        "زمان‌بندی": [18, 14, 14, 14, 12, 12, 20, 26],
        "راهنما": [22, 60, 30, 60],
        "فهرست‌های انتخاب": [14, 10, 10, 10, 10, 20, 28, 28, 22, 18, 18, 22, 22],
        "راهنمای نمادها": [20, 30, 50, 30],
    }
    for sheet_name, sheet_widths in widths.items():
        ws = wb[sheet_name]
        for index, width in enumerate(sheet_widths, start=1):
            ws.column_dimensions[chr(64 + index)].width = width

    if not include_example:
        for data_sheet, max_column in [
            (ws0, 5),
            (ws1, 10),
            (ws2, 12),
            (ws_unit_states, 8),
            (ws_unit_profiles, 17),
            (ws3, 10),
            (ws4, 7),
            (ws5, 9),
        ]:
            for row in data_sheet.iter_rows(
                min_row=2,
                max_row=data_sheet.max_row,
                max_col=max_column,
            ):
                for cell in row:
                    cell.value = None
        for row in ws_time.iter_rows(
            min_row=2,
            max_row=ws_time.max_row,
            min_col=1,
            max_col=7,
        ):
            for cell in row:
                cell.value = None

    return wb


# ---------------------------------------------------------------------------
# Standard sheet definitions (used by transform + export)
# ---------------------------------------------------------------------------

_STANDARD_SHEETS: dict[str, list[str]] = {
    "scenario": [
        "name",
        "description",
        "start_time",
        "time_zone",
        "symbology_standard",
    ],
    "events": [
        "id",
        "title",
        "subtitle",
        "start_time",
        "end_time",
        "side",
        "unit_ids",
        "equipment_ids",
        "lon",
        "lat",
    ],
    "units": [
        "id",
        "name",
        "parent_id",
        "side",
        "symbol_type",
        "echelon",
        "symbol_status",
        "advanced_sidc",
        "time",
        "lon",
        "lat",
        "resource_code",
    ],
    "unit_states": [
        "id",
        "unit_id",
        "time",
        "lon",
        "lat",
        "transfer_mode",
        "path_mode",
        "movement_start_time",
    ],
    "equipment": [
        "id",
        "name",
        "type",
        "quantity",
        "unit_id",
        "start_time",
        "lon",
        "lat",
        "side",
        "advanced_sidc",
    ],
    "personnel": [
        "id",
        "first_name",
        "last_name",
        "rank",
        "specialty",
        "national_id",
        "unit_id",
    ],
    "features": [
        "id",
        "name",
        "type",
        "lon",
        "lat",
        "radius_m",
        "start_time",
        "end_time",
        "event_id",
    ],
}

# Maps AI targetSheet names (lowercase) to canonical Persian sheet names
_TARGET_TO_SHEET_TITLE: dict[str, str] = {
    "scenario": "سناریو",
    "events": "حوادث",
    "units": "یگان‌ها",
    "unit_states": "وضعیت‌های زمانی یگان‌ها",
    "equipment": "تجهیزات",
    "personnel": "پرسنل",
    "features": "عوارض",
}


def inspect_source_table_layout(
    worksheet: Any,
    *,
    max_scan_rows: int = 15,
    max_columns: int = 50,
) -> dict[str, Any]:
    """Find a likely table header and return stable labels for mapped columns."""

    column_limit = min(max(worksheet.max_column, 1), max_columns)
    row_limit = min(max(worksheet.max_row, 1), max_scan_rows)

    def row_values(row: int) -> list[Any]:
        return [worksheet.cell(row=row, column=column).value for column in range(1, column_limit + 1)]

    def present(value: Any) -> bool:
        return value is not None and str(value).strip() != ""

    candidates: list[tuple[float, int]] = []
    for row in range(1, row_limit + 1):
        values = row_values(row)
        nonempty = sum(1 for value in values if present(value))
        text_cells = sum(
            1 for value in values if isinstance(value, str) and value.strip()
        )
        following_support = 0
        for next_row in range(row + 1, min(worksheet.max_row, row + 3) + 1):
            following_nonempty = sum(
                1 for value in row_values(next_row) if present(value)
            )
            if following_nonempty >= max(1, min(2, nonempty)):
                following_support += 1
        score = nonempty * 4 + text_cells * 2 + following_support * 3 + row * 0.05
        if nonempty < 2:
            score -= 20
        candidates.append((score, row))

    header_row = max(candidates, key=lambda item: item[0])[1]
    child_values = row_values(header_row)
    child_nonempty = sum(1 for value in child_values if present(value))

    parent_row = header_row - 1
    parent_values = row_values(parent_row) if parent_row >= 1 else []
    parent_nonempty = sum(1 for value in parent_values if present(value))
    merged_parent = False
    merged_ranges = getattr(getattr(worksheet, "merged_cells", None), "ranges", [])
    for merged_range in merged_ranges:
        if (
            merged_range.min_row <= parent_row <= merged_range.max_row
            and merged_range.max_col > merged_range.min_col
        ):
            merged_parent = True
            break
    header_depth = 2 if (
        parent_row >= 1
        and parent_nonempty > 0
        and (merged_parent or 2 <= parent_nonempty < child_nonempty)
    ) else 1

    expanded_parents: list[str] = []
    current_parent = ""
    if header_depth == 2:
        for value in parent_values:
            if present(value):
                current_parent = str(value).strip()
            expanded_parents.append(current_parent)

    last_header_column = 0
    for index, value in enumerate(child_values, start=1):
        if present(value) or (
            header_depth == 2 and index <= len(expanded_parents) and expanded_parents[index - 1]
        ):
            last_header_column = index

    headers: list[str] = []
    columns: dict[str, int] = {}
    duplicate_counts: dict[str, int] = {}
    for column in range(1, last_header_column + 1):
        child = str(child_values[column - 1]).strip() if present(child_values[column - 1]) else ""
        parent = expanded_parents[column - 1] if header_depth == 2 else ""
        label = (
            f"{parent} / {child}"
            if parent and child and parent != child
            else child or parent
        )
        if not label:
            headers.append("")
            continue
        count = duplicate_counts.get(label, 0) + 1
        duplicate_counts[label] = count
        unique_label = label if count == 1 else f"{label} [ستون {column}]"
        headers.append(unique_label)
        columns[_norm_key(unique_label)] = column

    return {
        "headerRow": header_row,
        "headerDepth": header_depth,
        "dataStartRow": header_row + 1,
        "headers": headers,
        "columns": columns,
    }


def transform_workbook_with_mapping(wb, mapping: dict[str, Any]) -> Workbook:
    """
    Given an arbitrary workbook and an AI mapping dict (sheetMappings + columnMaps),
    produce a new standard-format workbook that parse_excel_workbook() can consume.

    mapping shape:
      {
        "sheetMappings": [{"sourceSheet": "...", "targetSheet": "scenario|events|units|equipment|personnel", "confidence": 0-1}],
        "columnMaps": [{"targetSheet": "...", "mappings": [{"fromHeader": "...", "toField": "..."}]}]
      }
    """
    sheet_mappings: list[dict[str, Any]] = mapping.get("sheetMappings") or []
    column_maps: list[dict[str, Any]] = mapping.get("columnMaps") or []

    # Build lookup: targetSheet (lower) -> sourceSheetName
    source_for_target: dict[str, str] = {}
    for sm in sheet_mappings:
        t = _norm_key(sm.get("targetSheet") or "")
        s = _norm(sm.get("sourceSheet") or "")
        if t and s:
            source_for_target[t] = s

    # Build lookup: targetSheet -> {toField (lower) -> fromHeader (original)}
    col_map_for_target: dict[str, dict[str, str]] = {}
    for cm in column_maps:
        t = _norm_key(cm.get("targetSheet") or "")
        if not t:
            continue
        field_map: dict[str, str] = {}
        for m in cm.get("mappings") or []:
            to = _norm_key(m.get("toField") or "")
            frm = _norm(m.get("fromHeader") or "")
            if to and frm:
                field_map[to] = frm
        col_map_for_target[t] = field_map

    new_wb = Workbook()
    first = True

    for target_key, std_headers in _STANDARD_SHEETS.items():
        sheet_title = _TARGET_TO_SHEET_TITLE[target_key]
        if first:
            ws_new = new_wb.active
            ws_new.title = sheet_title
            first = False
        else:
            ws_new = new_wb.create_sheet(sheet_title)

        ws_new.append(std_headers)

        source_name = source_for_target.get(target_key)
        if not source_name:
            continue

        # Find source worksheet by name (case-insensitive)
        source_ws = None
        for sn in wb.sheetnames:
            if _norm(sn).lower() == _norm(source_name).lower():
                source_ws = wb[sn]
                break
        if source_ws is None:
            continue

        source_layout = inspect_source_table_layout(source_ws)
        src_headers: dict[str, int] = source_layout["columns"]

        field_map = col_map_for_target.get(target_key) or {}

        # For each source data row, build a standard row
        for row_cells in source_ws.iter_rows(
            min_row=source_layout["dataStartRow"], values_only=True
        ):
            if all(x is None or str(x).strip() == "" for x in row_cells):
                continue
            row_list = list(row_cells)

            std_row: list[Any] = []
            for field in std_headers:
                field_key = _norm_key(field)
                # Look up the original header that maps to this standard field
                from_header = field_map.get(field_key)
                value = None
                if from_header:
                    from_key = _norm_key(from_header)
                    if from_key in src_headers:
                        col_idx = src_headers[from_key] - 1
                        value = row_list[col_idx] if col_idx < len(row_list) else None
                std_row.append(value)

            ws_new.append(std_row)

    return new_wb


def export_content_to_workbook(content: dict[str, Any]) -> Workbook:
    """
    Convert a parsed scenario content dict back into a standard Excel workbook.
    Useful for downloading the standardized version after AI import.
    """
    wb = Workbook()

    # --- سناریو ---
    ws0 = wb.active
    ws0.title = "سناریو"
    ws0.append(["name", "description", "start_time", "time_zone", "symbology_standard"])
    ws0.append(
        [
            content.get("name") or "",
            content.get("description") or "",
            content.get("startTime") or "",
            content.get("timeZone") or "UTC",
            content.get("symbologyStandard") or "app6",
        ]
    )

    # --- حوادث ---
    ws1 = wb.create_sheet("حوادث")
    ws1.append(["id", "title", "subtitle", "start_time", "lon", "lat"])
    for ev in content.get("events") or []:
        geo = ev.get("where", {})
        coords = (
            geo.get("geometry", {}).get("coordinates")
            if isinstance(geo, dict)
            else None
        )
        lon = coords[0] if coords and len(coords) >= 2 else ""
        lat = coords[1] if coords and len(coords) >= 2 else ""
        ws1.append(
            [
                ev.get("id") or "",
                ev.get("title") or "",
                ev.get("subTitle") or "",
                ev.get("startTime") or "",
                lon,
                lat,
            ]
        )

    # --- یگان‌ها ---
    ws2 = wb.create_sheet("یگان‌ها")
    ws2.append(["id", "name", "parent_id", "side", "unit_type", "time", "lon", "lat"])
    ws_states = wb.create_sheet("وضعیت‌های زمانی یگان‌ها")
    ws_states.append(
        [
            "id",
            "unit_id",
            "time",
            "lon",
            "lat",
            "transfer_mode",
            "path_mode",
            "movement_start_time",
        ]
    )
    _side_std_to_label = {"3": "friend", "6": "hostile", "4": "neutral", "0": "unknown"}

    def _flatten_units(units: list, parent_id: str = "") -> None:
        for u in units or []:
            std_id = u.get("standard_identity") or u.get("standardIdentity") or "3"
            side_label = _side_std_to_label.get(std_id, "friend")
            # Find symbol type from SIDC (best-effort reverse)
            sidc = u.get("sidc") or ""
            symbol_set = sidc[4:6] if len(sidc) >= 16 else ""
            main_icon = sidc[10:16] if len(sidc) >= 16 else ""
            unit_type = next(
                (
                    k
                    for k, v in UNIT_SYMBOL_TYPES.items()
                    if v == (symbol_set, main_icon) and k.isascii()
                ),
                "infantry",
            )
            state0 = u.get("state", [{}])[0] if u.get("state") else {}
            loc = state0.get("location") or []
            lon = loc[0] if len(loc) >= 2 else ""
            lat = loc[1] if len(loc) >= 2 else ""
            ws2.append(
                [
                    u.get("id") or "",
                    u.get("name") or "",
                    parent_id,
                    side_label,
                    unit_type,
                    state0.get("t") or "",
                    lon,
                    lat,
                ]
            )
            for state in (u.get("state") or [])[1:]:
                state_location = state.get("location") or []
                ws_states.append(
                    [
                        state.get("id") or "",
                        u.get("id") or "",
                        state.get("t") or "",
                        state_location[0] if len(state_location) >= 2 else "",
                        state_location[1] if len(state_location) >= 2 else "",
                        "continuous"
                        if state.get("interpolate") is not False
                        else "jump",
                        state.get("pathMode") or "straight",
                        state.get("viaStartTime") or "",
                    ]
                )
            _flatten_units(u.get("subUnits") or [], u.get("id") or "")

    for side in content.get("sides") or []:
        for group in side.get("groups") or []:
            _flatten_units(group.get("subUnits") or [])

    # --- تجهیزات ---
    ws3 = wb.create_sheet("تجهیزات")
    ws3.append(["id", "name", "type", "quantity", "unit_id"])
    for eq in content.get("equipment") or []:
        ws3.append(
            [
                eq.get("id") or "",
                eq.get("name") or "",
                eq.get("type") or "",
                eq.get("quantity") or 1,
                eq.get("unitId") or "",
            ]
        )

    # --- پرسنل ---
    ws4 = wb.create_sheet("پرسنل")
    ws4.append(
        ["id", "first_name", "last_name", "rank", "specialty", "national_id", "unit_id"]
    )
    for p in content.get("personnel") or []:
        ws4.append(
            [
                p.get("id") or "",
                p.get("firstName") or "",
                p.get("lastName") or "",
                p.get("rank") or "",
                p.get("specialty") or "",
                p.get("nationalId") or "",
                p.get("unitId") or "",
            ]
        )

    return wb


def parse_resources_workbook(
    wb,
) -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
    list[dict[str, Any]],
    list[dict[str, Any]],
]:
    """
    شیت‌های یگان، تجهیزات و پرسنل را برای ثبت مستقل در کاتالوگ منابع می‌خواند.
    """
    errors: list[dict[str, Any]] = []
    now = datetime.now(timezone.utc).isoformat()
    unit_profiles = _read_unit_profiles(wb)
    equip_ws = _find_sheet(wb, "تجهیزات", "equipment")
    pers_ws = _find_sheet(wb, "پرسنل", "personnel")
    units_ws = _find_sheet(wb, "یگان", "یگان‌ها", "یگان ها", "units", "orbat")
    if not equip_ws and not pers_ws and not units_ws:
        errors.append(
            {
                "sheet": "-",
                "row": 0,
                "message": "هیچ شیت یگان، تجهیزات یا پرسنل یافت نشد",
            }
        )
        return [], [], [], errors

    units_out: list[dict[str, Any]] = []
    if units_ws:
        uh, _ = _header_map(units_ws)
        for i, ur in enumerate(_read_rows(units_ws), start=2):
            name = _cell(ur, uh, "name", "نام")
            if not name:
                errors.append(
                    {"sheet": units_ws.title, "row": i, "message": "نام یگان خالی است"}
                )
                continue
            unit_id = _cell(ur, uh, "id", "شناسه")
            if not unit_id:
                unit_id = f"unit-{uuid.uuid4().hex[:10]}"
                errors.append(
                    {
                        "sheet": units_ws.title,
                        "row": i,
                        "message": "شناسه یگان خالی بود؛ یک شناسه موقت ساخته شد",
                    }
                )
            side_raw = _cell(ur, uh, "side", "طرف", "جبهه") or "خودی"
            side_key = _norm_key(side_raw)
            side_id = (
                SIDE_MAP.get(side_key) or SIDE_MAP.get(side_raw.strip().lower()) or "3"
            )
            unit_type = _cell(
                ur,
                uh,
                "symbol_type",
                "unit_type",
                "type",
                "نوع",
                "نوع_واحد",
                "نوع_نماد",
            )
            echelon = _cell(
                ur, uh, "echelon", "unit_echelon", "رده", "رده_یگان", "رده_سازمانی"
            )
            resource_code = (
                _cell(ur, uh, "resource_code", "unit_code", "کد_مرجع_یگان", "کد_مرجع")
                or unit_id
            )
            profile = unit_profiles.get(resource_code, {})
            status_value = (
                _cell(ur, uh, "symbol_status", "status", "وضعیت_نماد", "وضعیت")
                or "حاضر"
            )
            advanced_sidc = _clean_sidc(
                _cell(
                    ur,
                    uh,
                    "advanced_sidc",
                    "sidc",
                    "نماد",
                    "کد_نماد",
                    "کد_نماد_پیشرفته",
                )
            )
            if advanced_sidc and not re.fullmatch(r"\d{20}", advanced_sidc):
                errors.append(
                    {
                        "sheet": units_ws.title,
                        "row": i,
                        "message": "کد نماد پیشرفته باید دقیقاً ۲۰ رقم باشد؛ نماد از انتخاب‌های فارسی ساخته شد",
                    }
                )
                advanced_sidc = ""
            units_out.append(
                {
                    "id": unit_id,
                    "unitCode": resource_code,
                    "name": name,
                    "shortName": profile.get("shortName")
                    or _cell(ur, uh, "short_name", "عنوان_کوتاه", "نام_کوتاه")
                    or None,
                    "description": profile.get("description")
                    or _cell(ur, uh, "description", "توضیح", "توضیحات")
                    or None,
                    "status": status_value,
                    "sidc": advanced_sidc
                    or _build_unit_sidc(side_id, unit_type, echelon, status_value),
                    "side": side_raw,
                    "parentCode": _cell(
                        ur, uh, "parent_id", "parentid", "والد", "id_والد", "شناسه_والد"
                    )
                    or profile.get("parentCode"),
                    "parentName": profile.get("parentName"),
                    "echelon": echelon or None,
                    "unitType": unit_type or None,
                    "organizationalAffiliation": profile.get(
                        "organizationalAffiliation"
                    ),
                    "formedOn": profile.get("formedOn"),
                    "deactivatedOn": profile.get("deactivatedOn"),
                    "province": profile.get("province")
                    or _cell(ur, uh, "province", "استان")
                    or None,
                    "garrisonCity": profile.get("garrisonCity")
                    or _cell(ur, uh, "garrison_city", "city", "شهر_استقرار")
                    or None,
                    "baseName": profile.get("baseName")
                    or _cell(ur, uh, "base_name", "نام_پادگان", "پادگان", "قرارگاه")
                    or None,
                    "capabilities": profile.get("capabilities")
                    or _split_ids(
                        _cell(ur, uh, "capabilities", "توانمندی‌ها", "توانمندی ها")
                    ),
                    "nominalPersonnelStrength": profile.get("nominalPersonnelStrength")
                    or _cell(ur, uh, "nominal_strength", "استعداد_اسمی")
                    or None,
                    "defaultReadiness": profile.get("defaultReadiness")
                    or _cell(ur, uh, "default_readiness", "آمادگی_پایه")
                    or None,
                    "areaOfOps": profile.get("areaOfOps")
                    or _cell(ur, uh, "area_of_ops", "محدوده_مسئولیت")
                    or None,
                    "commanderName": profile.get("commanderName")
                    or _cell(ur, uh, "commander", "فرمانده")
                    or None,
                    "contact": profile.get("contact")
                    or _cell(ur, uh, "contact", "اطلاعات_تماس")
                    or None,
                }
            )

        resource_code_by_row_id = {
            str(unit.get("id")): str(unit.get("unitCode"))
            for unit in units_out
            if unit.get("id") and unit.get("unitCode")
        }
        for unit in units_out:
            parent_id = str(unit.get("parentCode") or "")
            if parent_id in resource_code_by_row_id:
                unit["parentCode"] = resource_code_by_row_id[parent_id]

    equipment_out: list[dict[str, Any]] = []
    if equip_ws:
        eh, _ = _header_map(equip_ws)
        for i, er in enumerate(_read_rows(equip_ws), start=2):
            ename = _cell(er, eh, "name", "نام")
            if not ename:
                continue
            eid = _cell(er, eh, "id", "شناسه") or f"EQ-{uuid.uuid4().hex[:8]}"
            code = _cell(er, eh, "equipment_code", "کد", "equipmentcode") or eid
            unit_ref = _cell(
                er, eh, "unit_id", "id_یگان", "شناسه_یگان", "یگان", "assigned_to"
            )
            quantity_raw = _cell(er, eh, "quantity", "تعداد", "count") or "1"
            try:
                quantity = max(0, int(float(quantity_raw)))
            except ValueError:
                quantity = 1
                errors.append(
                    {
                        "sheet": equip_ws.title,
                        "row": i,
                        "message": f"تعداد نامعتبر: {quantity_raw}",
                    }
                )
            equipment_out.append(
                {
                    "id": eid,
                    "equipmentCode": code,
                    "name": ename,
                    "type": _cell(er, eh, "type", "نوع") or "عمومی",
                    "quantity": quantity,
                    "model": _cell(er, eh, "model", "مدل") or None,
                    "manufacturer": _cell(er, eh, "manufacturer", "سازنده") or None,
                    "serialNumber": _cell(er, eh, "serial_number", "سریال") or None,
                    "acquisitionDate": _cell(er, eh, "acquisition_date", "تاریخ_خرید")
                    or now[:10],
                    "condition": "excellent",
                    "location": _cell(er, eh, "location", "مکان") or unit_ref or "",
                    "assignedTo": unit_ref or None,
                    "status": "available",
                    "createdAt": now,
                    "updatedAt": now,
                }
            )

    personnel_out: list[dict[str, Any]] = []
    if pers_ws:
        ph, _ = _header_map(pers_ws)
        for i, pr in enumerate(_read_rows(pers_ws), start=2):
            fn = _cell(pr, ph, "first_name", "نام")
            ln = _cell(pr, ph, "last_name", "نام_خانوادگی", "فامیلی")
            full = _cell(pr, ph, "full_name", "نام_کامل")
            if not fn and not ln and full:
                parts = full.split(maxsplit=1)
                fn = parts[0]
                ln = parts[1] if len(parts) > 1 else ""
            if not fn and not ln:
                continue
            pid = _cell(pr, ph, "id", "شناسه") or f"person-{uuid.uuid4().hex[:8]}"
            pcode = _cell(pr, ph, "personal_code", "کد_پرسنلی", "personalcode") or pid
            personnel_out.append(
                {
                    "id": pid,
                    "personalCode": pcode,
                    "firstName": fn,
                    "lastName": ln or "",
                    "nationalId": _cell(pr, ph, "national_id", "کد_ملی", "nationalid")
                    .lstrip("'")
                    .replace("‌", "")
                    or "",
                    "rank": _cell(pr, ph, "rank", "درجه") or "",
                    "unit": _cell(
                        pr, ph, "unit_id", "id_یگان", "شناسه_یگان", "یگان", "unit"
                    )
                    or "",
                    "position": _cell(pr, ph, "specialty", "position", "تخصص", "سمت")
                    or None,
                    "phoneNumber": _cell(pr, ph, "phone", "تلفن") or None,
                    "email": _cell(pr, ph, "email", "ایمیل") or None,
                    "status": "active",
                    "startDate": _cell(pr, ph, "start_date", "تاریخ_شروع") or now[:10],
                    "createdAt": now,
                    "updatedAt": now,
                }
            )

    return personnel_out, equipment_out, units_out, errors
