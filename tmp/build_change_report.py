from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("output/documents/گزارش تغییرات پروژه کالک - 15 تیر تا 10 مرداد 1405.docx")

FONT = "Tahoma"
NAVY = "16324F"
BLUE = "2E74B5"
TEAL = "168C83"
INK = "243447"
MUTED = "667085"
LIGHT_BLUE = "EAF2F8"
LIGHT_TEAL = "EAF7F5"
LIGHT_GRAY = "F7F9FC"
BORDER = "D7DEE7"
WHITE = "FFFFFF"


def set_cell_text_direction(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    text_dir = tc_pr.find(qn("w:textDirection"))
    if text_dir is None:
        text_dir = OxmlElement("w:textDirection")
        tc_pr.append(text_dir)
    text_dir.set(qn("w:val"), "rtl")


def set_paragraph_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_run_font(run, size=10.5, color=INK, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT)
    rtl = r_pr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        r_pr.append(rtl)
    rtl.set(qn("w:val"), "1")


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.1, keep_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    bidi = tbl_pr.find(qn("w:bidiVisual"))
    if bidi is None:
        bidi = OxmlElement("w:bidiVisual")
        tbl_pr.append(bidi)
    bidi.set(qn("w:val"), "1")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def add_page_field(paragraph):
    run = paragraph.add_run("صفحه ")
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def add_cell_text(cell, text, *, size=10, color=INK, bold=False, center=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_text_direction(cell)
    p = cell.paragraphs[0]
    set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT)
    set_paragraph_spacing(p, before=0, after=0, line=1.15)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    doc.core_properties.title = "گزارش تغییرات پروژه کالک"
    doc.core_properties.subject = "گزارش تغییرات ثبت‌شده از ۱۵ تیر تا ۱۰ مرداد ۱۴۰۵"
    doc.core_properties.author = "تیم توسعه پروژه کالک"
    doc.core_properties.last_modified_by = "تیم توسعه پروژه کالک"

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # Running header
    hp = section.header.paragraphs[0]
    set_paragraph_rtl(hp)
    set_paragraph_spacing(hp, after=0, line=1)
    hr = hp.add_run("گزارش پیشرفت پروژه  |  کالک")
    set_run_font(hr, size=8.5, color=MUTED, bold=True)
    p_pr = hp._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), TEAL)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    # Footer
    fp = section.footer.paragraphs[0]
    set_paragraph_rtl(fp, WD_ALIGN_PARAGRAPH.CENTER)
    set_paragraph_spacing(fp, after=0, line=1)
    add_page_field(fp)

    # Kicker
    kicker = doc.add_paragraph()
    set_paragraph_rtl(kicker)
    set_paragraph_spacing(kicker, before=7, after=5, line=1)
    kr = kicker.add_run("گزارش مدیریتی دوره‌ای")
    set_run_font(kr, size=9.5, color=TEAL, bold=True)

    # Title and subtitle
    title = doc.add_paragraph()
    set_paragraph_rtl(title)
    set_paragraph_spacing(title, before=0, after=5, line=1)
    tr = title.add_run("گزارش تغییرات پروژه «کالک»")
    set_run_font(tr, size=23, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    set_paragraph_rtl(subtitle)
    set_paragraph_spacing(subtitle, before=0, after=13, line=1.1)
    sr = subtitle.add_run("مرور کوتاه اقدامات انجام‌شده برای ارائه به کارفرما")
    set_run_font(sr, size=11.5, color=MUTED)

    # Metadata block: factual key-value data.
    meta = doc.add_table(rows=2, cols=2)
    set_table_geometry(meta, [4680, 4680], indent_dxa=120)
    set_table_borders(meta, color="D8E5EC", size=5)
    meta_data = [
        ("بازه گزارش", "۱۵ تیر تا ۱۰ مرداد ۱۴۰۵"),
        ("آخرین تغییر بررسی‌شده", "۹ مرداد ۱۴۰۵"),
        ("تهیه‌کننده", "تیم توسعه پروژه کالک"),
        ("مبنای گزارش", "تغییرات ثبت‌شده در مخزن Git"),
    ]
    for idx, cell in enumerate([c for row in meta.rows for c in row.cells]):
        label, value = meta_data[idx]
        shade_cell(cell, LIGHT_BLUE if idx % 2 == 0 else LIGHT_TEAL)
        set_cell_margins(cell, top=95, start=120, bottom=95, end=120)
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_text_direction(cell)
        p = cell.paragraphs[0]
        set_paragraph_rtl(p)
        set_paragraph_spacing(p, after=1, line=1.05)
        r = p.add_run(label)
        set_run_font(r, size=8.5, color=MUTED, bold=True)
        p2 = cell.add_paragraph()
        set_paragraph_rtl(p2)
        set_paragraph_spacing(p2, after=0, line=1.1)
        r2 = p2.add_run(value)
        set_run_font(r2, size=10.2, color=NAVY, bold=True)

    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, after=0, line=0.2)

    # Executive summary callout.
    callout = doc.add_table(rows=1, cols=1)
    set_table_geometry(callout, [9360], indent_dxa=120)
    set_table_borders(callout, color="BFD8D4", size=7)
    c = callout.cell(0, 0)
    shade_cell(c, LIGHT_TEAL)
    set_cell_margins(c, top=120, start=120, bottom=120, end=120)
    c.text = ""
    set_cell_text_direction(c)
    p = c.paragraphs[0]
    set_paragraph_rtl(p)
    set_paragraph_spacing(p, after=2, line=1.05)
    r = p.add_run("خلاصه گزارش")
    set_run_font(r, size=9.2, color=TEAL, bold=True)
    p2 = c.add_paragraph()
    set_paragraph_rtl(p2)
    set_paragraph_spacing(p2, after=0, line=1.18)
    r2 = p2.add_run(
        "در این دوره، تمرکز تیم بر تکمیل جریان سناریو، بهبود ابزارهای نقشه و نمادها، "
        "ارتقای داشبورد و پایدارسازی ذخیره‌سازی و زیرساخت بوده است."
    )
    set_run_font(r2, size=10.2, color=INK)

    heading = doc.add_paragraph(style="Heading 1")
    set_paragraph_rtl(heading)
    hr = heading.add_run("تغییرات انجام‌شده به تفکیک بخش")
    set_run_font(hr, size=16, color=BLUE, bold=True)

    changes = [
        ("۱", "داشبورد و نمای مدیریتی", "صفحه اصلی و کارت‌های سناریو بازطراحی شد؛ پیش‌نمایش نقشه، فضای کاری قابل ویرایش و مسیرهای دسترسی صفحات نیز بهبود یافت."),
        ("۲", "مدیریت و جزئیات سناریو", "فرایند ساخت و ویرایش سناریو اصلاح شد. زمان‌بندی با تقویم فارسی، وضعیت سناریو و تب‌های معرفی، مراحل، محیط و تاریخچه تکمیل شدند."),
        ("۳", "ذخیره‌سازی و نسخه‌ها", "ذخیره خودکار، بازیابی پیش‌نویس، تاریخچه تغییرات و نسخه‌بندی سناریو اضافه شد تا احتمال از دست رفتن اطلاعات کاهش یابد."),
        ("۴", "کالک و نقشه", "پنل لایه‌ها و تنظیمات نقشه اصلاح شد؛ نقشه پایه و ترجیحات نمای کاربر حفظ می‌شود و ابزارهای ترسیم، پاک‌کردن و بازگشت تغییرات بهبود یافته‌اند."),
        ("۵", "نمادهای تاکتیکی", "عناوین فارسی بازبینی شد و عملکرد جست‌وجو، علاقه‌مندی‌ها و ترسیم نمادها بهبود یافت. ترسیم بادبزن برد و انتخاب و ویرایش رده مرزها نیز اضافه شد."),
        ("۶", "هواشناسی و شرایط محیطی", "کتابخانه و انتخابگر نمادهای هواشناسی ایجاد شد. نمایش نمادها روی نقشه و داشبورد، پیش‌نمایش زنده و امکان حذف عوارض محیطی نیز اضافه شد."),
        ("۷", "زمان و روشنایی نقشه", "زمان نقشه به موقعیت و منطقه زمانی متصل شد؛ نمایش تاریخ و ساعت اصلاح شد و لایه روز و شب با خط زمانی سناریو هماهنگ می‌ماند."),
        ("۸", "کاربران، دسترسی و اعلان‌ها", "مدیریت کاربران، سطح دسترسی و ثبت سوابق فعالیت بهبود یافت. انتخاب تصویر کاربر و نمایش اعلان‌های سیستمی نیز تکمیل شد."),
        ("۹", "زیرساخت و کنترل کیفیت", "APIها و ساختار پایگاه داده به‌روزرسانی شد؛ اجرای Docker و بسته آفلاین اصلاح شد و آزمون‌های خودکار برای بخش‌های اصلی افزایش یافت."),
    ]

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [780, 2150, 6430], indent_dxa=120)
    set_table_borders(table, color=BORDER, size=6)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    headers = ["ردیف", "بخش پروژه", "شرح تغییرات"]
    for i, (cell, text) in enumerate(zip(header.cells, headers)):
        shade_cell(cell, NAVY)
        set_cell_margins(cell, top=120, start=120, bottom=120, end=120)
        add_cell_text(cell, text, size=9.5, color=WHITE, bold=True, center=(i == 0))

    for row_idx, row_data in enumerate(changes, start=1):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        fill = WHITE if row_idx % 2 else LIGHT_GRAY
        for c in cells:
            shade_cell(c, fill)
            set_cell_margins(c, top=115, start=120, bottom=115, end=120)
        add_cell_text(cells[0], row_data[0], size=9.5, color=TEAL, bold=True, center=True)
        add_cell_text(cells[1], row_data[1], size=9.4, color=NAVY, bold=True)
        add_cell_text(cells[2], row_data[2], size=9.25, color=INK)

    # Re-apply geometry after all rows have been created.
    set_table_geometry(table, [780, 2150, 6430], indent_dxa=120)

    note_heading = doc.add_paragraph(style="Heading 2")
    set_paragraph_rtl(note_heading)
    nr = note_heading.add_run("یادداشت گزارش")
    set_run_font(nr, size=13, color=BLUE, bold=True)

    note = doc.add_paragraph()
    set_paragraph_rtl(note)
    set_paragraph_spacing(note, before=0, after=0, line=1.18)
    note_run = note.add_run(
        "این گزارش بر اساس تغییرات ثبت‌شده در مخزن پروژه تهیه شده است. نخستین ثبت موجود در بازه مورد بررسی "
        "مربوط به ۲۹ تیر و آخرین ثبت بررسی‌شده مربوط به ۹ مرداد ۱۴۰۵ است."
    )
    set_run_font(note_run, size=9.4, color=MUTED)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    build_document()
    print("change report created")
