from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\AI\Documents\GitHub\kalk")
SOURCE = next((ROOT / "output" / "documents").glob("*ساجد*C2.docx"))
TARGET = ROOT / "output" / "documents" / "گزارش مقایسه قابلیت‌های سامانه ما با سامانه C2.docx"
OLD_NAME = "ساجد"
NEW_NAME = "سامانه ما"
FONT = "B Nazanin"
DOUBLE_NAME = "سامانه سامانه ما"


def ensure_font(run) -> None:
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        r_fonts.set(qn(f"w:{attr}"), FONT)
    if r_pr.find(qn("w:rtl")) is None:
        r_pr.append(OxmlElement("w:rtl"))


def revise_paragraph(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is None:
        p_pr.append(OxmlElement("w:bidi"))
    for run in paragraph.runs:
        if OLD_NAME in run.text:
            run.text = run.text.replace(OLD_NAME, NEW_NAME)
        if DOUBLE_NAME in run.text:
            run.text = run.text.replace(DOUBLE_NAME, NEW_NAME)
        ensure_font(run)


def revise_table(table) -> None:
    tbl_pr = table._tbl.tblPr
    if tbl_pr.find(qn("w:bidiVisual")) is None:
        tbl_pr.append(OxmlElement("w:bidiVisual"))
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                revise_paragraph(paragraph)


doc = Document(SOURCE)

for style in doc.styles:
    if not hasattr(style, "font"):
        continue
    style.font.name = FONT
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        r_fonts.set(qn(f"w:{attr}"), FONT)

for paragraph in doc.paragraphs:
    revise_paragraph(paragraph)
for table in doc.tables:
    revise_table(table)

for section in doc.sections:
    for paragraph in section.header.paragraphs:
        revise_paragraph(paragraph)
    for table in section.header.tables:
        revise_table(table)
    for paragraph in section.footer.paragraphs:
        revise_paragraph(paragraph)
    for table in section.footer.tables:
        revise_table(table)

for field in ("title", "subject", "comments", "keywords", "author"):
    value = getattr(doc.core_properties, field) or ""
    value = value.replace(OLD_NAME, NEW_NAME).replace(DOUBLE_NAME, NEW_NAME)
    setattr(doc.core_properties, field, value)

doc.save(TARGET)

with ZipFile(TARGET) as archive:
    xml = "\n".join(
        archive.read(name).decode("utf-8", errors="ignore")
        for name in archive.namelist()
        if name.endswith(".xml")
    )
if OLD_NAME in xml:
    raise RuntimeError("نام قبلی هنوز در فایل Word وجود دارد")
if DOUBLE_NAME in xml:
    raise RuntimeError("ترکیب تکراری نام سامانه هنوز در فایل Word وجود دارد")
if FONT not in xml:
    raise RuntimeError("فونت B Nazanin در فایل Word ثبت نشده است")

print(TARGET)
