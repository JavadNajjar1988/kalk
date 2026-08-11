from __future__ import annotations

import re
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips


ROOT = Path(r"C:\Users\AI\Documents\GitHub\kalk")
OUT_DOCX = ROOT / "output" / "documents"
OUT_DOCX.mkdir(parents=True, exist_ok=True)

PROPOSAL_PATH = OUT_DOCX / "پروپوزال سامانه مدیریت و رصد میدان.docx"
SPEAKER_PATH = OUT_DOCX / "متن مطالعه و ارائه سامانه مدیریت و رصد میدان.docx"

SKILL_ROOT = Path(
    r"C:\Users\AI\.codex\plugins\cache\openai-primary-runtime"
    r"\documents\26.727.11326\skills\documents"
)
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
from table_geometry import apply_table_geometry, section_content_width_dxa  # noqa: E402


FONT = "B Nazanin"
FONT_BOLD = "B Nazanin"
NAVY = "243B53"
TEAL = "2F6F5E"
PALE_TEAL = "EAF2EF"
PALE_BLUE = "EAF0F6"
PALE_GOLD = "F7F1DF"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "697386"
DARK = "1E293B"
WHITE = "FFFFFF"
RED = "8C2D2D"
GREEN = "2E664B"

# A4 + Persian business-document override, based on narrative_proposal /
# compact_reference_guide. Exact values are applied rather than inherited.
PAGE_WIDTH = Cm(21)
PAGE_HEIGHT = Cm(29.7)
MARGIN_TOP = Cm(1.8)
MARGIN_BOTTOM = Cm(1.8)
MARGIN_SIDE = Cm(2.0)
HEADER_DISTANCE = Cm(0.9)
FOOTER_DISTANCE = Cm(0.9)
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 100, "bottom": 100, "start": 120, "end": 120}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "CBD5E1", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_run_font(run, size: float | None = None, *, bold=None, color=None, italic=None) -> None:
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        r_fonts.set(qn(f"w:{attr}"), FONT_BOLD if bold else FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if r_pr.find(qn("w:rtl")) is None:
        r_pr.append(OxmlElement("w:rtl"))


def set_paragraph_rtl(paragraph, *, align=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    paragraph.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is None:
        p_pr.append(OxmlElement("w:bidi"))


def set_keep(paragraph, *, next_=False, lines=True, page_break=False) -> None:
    fmt = paragraph.paragraph_format
    fmt.keep_with_next = next_
    fmt.keep_together = lines
    fmt.page_break_before = page_break
    fmt.widow_control = True


def add_page_number(paragraph) -> None:
    paragraph.add_run("صفحه ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_custom_numbering(doc: Document, *, decimal: bool = False) -> int:
    numbering = doc.part.numbering_part.element
    existing_abs = [
        int(x.get(qn("w:abstractNumId")))
        for x in numbering.findall(qn("w:abstractNum"))
        if x.get(qn("w:abstractNumId"))
    ]
    existing_nums = [
        int(x.get(qn("w:numId")))
        for x in numbering.findall(qn("w:num"))
        if x.get(qn("w:numId"))
    ]
    abstract_id = max(existing_abs, default=0) + 1
    num_id = max(existing_nums, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if decimal else "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if decimal else "•")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "right")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:right"), "540")
    ind.set(qn("w:hanging"), "260")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "290")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    bidi = OxmlElement("w:bidi")
    p_pr.append(bidi)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        r_fonts.set(qn(f"w:{attr}"), FONT)
    r_pr.append(r_fonts)
    rtl = OxmlElement("w:rtl")
    r_pr.append(rtl)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), str(abstract_id))
    num.append(abs_id)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)


class PersianDoc:
    def __init__(self, *, preset: str, running_title: str):
        self.doc = Document()
        self.preset = preset
        self.running_title = running_title
        self._setup_page()
        self._setup_styles()
        self.bullet_num_id = add_custom_numbering(self.doc, decimal=False)
        self.decimal_num_id = add_custom_numbering(self.doc, decimal=True)
        self._setup_header_footer()

    def _setup_page(self) -> None:
        for section in self.doc.sections:
            section.page_width = PAGE_WIDTH
            section.page_height = PAGE_HEIGHT
            section.top_margin = MARGIN_TOP
            section.bottom_margin = MARGIN_BOTTOM
            section.left_margin = MARGIN_SIDE
            section.right_margin = MARGIN_SIDE
            section.header_distance = HEADER_DISTANCE
            section.footer_distance = FOOTER_DISTANCE

    def _setup_styles(self) -> None:
        styles = self.doc.styles
        body_after = 8 if self.preset == "narrative_proposal" else 6
        body_line = 1.28 if self.preset == "narrative_proposal" else 1.20
        values = {
            "Normal": (13, DARK, 0, body_after, body_line, False),
            "Title": (25, NAVY, 0, 5, 1.0, True),
            "Subtitle": (15, MID_GRAY, 0, 12, 1.1, False),
            "Heading 1": (18, NAVY, 16, 8, 1.0, True),
            "Heading 2": (15.5, TEAL, 12, 6, 1.0, True),
            "Heading 3": (14, NAVY, 9, 4, 1.0, True),
            "Caption": (11.5, MID_GRAY, 4, 4, 1.0, False),
        }
        for name, (size, color, before, after, line, bold) in values.items():
            style = styles[name]
            style.font.name = FONT
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string(color)
            style.font.bold = bold
            r_pr = style._element.get_or_add_rPr()
            r_fonts = r_pr.rFonts
            if r_fonts is None:
                r_fonts = OxmlElement("w:rFonts")
                r_pr.insert(0, r_fonts)
            for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
                r_fonts.set(qn(f"w:{attr}"), FONT)
            p_pr = style._element.get_or_add_pPr()
            if p_pr.find(qn("w:bidi")) is None:
                p_pr.append(OxmlElement("w:bidi"))
            jc = p_pr.find(qn("w:jc"))
            if jc is None:
                jc = OxmlElement("w:jc")
                p_pr.append(jc)
            jc.set(qn("w:val"), "both" if name == "Normal" and self.preset == "narrative_proposal" else "right")
            spacing = p_pr.find(qn("w:spacing"))
            if spacing is None:
                spacing = OxmlElement("w:spacing")
                p_pr.append(spacing)
            spacing.set(qn("w:before"), str(int(before * 20)))
            spacing.set(qn("w:after"), str(int(after * 20)))
            spacing.set(qn("w:line"), str(int(line * 240)))
            spacing.set(qn("w:lineRule"), "auto")

    def _setup_header_footer(self) -> None:
        section = self.doc.sections[0]
        section.different_first_page_header_footer = True
        header = section.header
        hp = header.paragraphs[0]
        set_paragraph_rtl(hp)
        hp.paragraph_format.space_after = Pt(2)
        r = hp.add_run(self.running_title)
        set_run_font(r, 10.5, bold=True, color=MID_GRAY)

        footer = section.footer
        fp = footer.paragraphs[0]
        set_paragraph_rtl(fp, align=WD_ALIGN_PARAGRAPH.CENTER)
        fp.paragraph_format.space_before = Pt(2)
        add_page_number(fp)
        for run in fp.runs:
            set_run_font(run, 9.5, color=MID_GRAY)

    def title(self, text: str, subtitle: str, *, meta: list[str]) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(54)
        p.paragraph_format.space_after = Pt(6)
        set_paragraph_rtl(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(text)
        set_run_font(r, 25, bold=True, color=NAVY)
        set_keep(p, next_=True)

        sp = self.doc.add_paragraph()
        set_paragraph_rtl(sp, align=WD_ALIGN_PARAGRAPH.CENTER)
        sp.paragraph_format.space_after = Pt(24)
        r = sp.add_run(subtitle)
        set_run_font(r, 15, color=TEAL)
        set_keep(sp, next_=True)

        line = self.doc.add_paragraph()
        set_paragraph_rtl(line, align=WD_ALIGN_PARAGRAPH.CENTER)
        line.paragraph_format.space_after = Pt(20)
        r = line.add_run("نام موقت پروژه - نام نهایی پس از توافق تعیین خواهد شد")
        set_run_font(r, 11.5, italic=True, color=MID_GRAY)

        for item in meta:
            mp = self.doc.add_paragraph()
            set_paragraph_rtl(mp, align=WD_ALIGN_PARAGRAPH.CENTER)
            mp.paragraph_format.space_after = Pt(3)
            rr = mp.add_run(item)
            set_run_font(rr, 12.5, color=DARK)

    def h1(self, text: str, *, page_break=False) -> None:
        p = self.doc.add_paragraph(text, style="Heading 1")
        set_paragraph_rtl(p)
        set_keep(p, next_=True, page_break=page_break)
        for run in p.runs:
            set_run_font(run, 18, bold=True, color=NAVY)

    def h2(self, text: str) -> None:
        p = self.doc.add_paragraph(text, style="Heading 2")
        set_paragraph_rtl(p)
        set_keep(p, next_=True)
        for run in p.runs:
            set_run_font(run, 15.5, bold=True, color=TEAL)

    def h3(self, text: str) -> None:
        p = self.doc.add_paragraph(text, style="Heading 3")
        set_paragraph_rtl(p)
        set_keep(p, next_=True)
        for run in p.runs:
            set_run_font(run, 14, bold=True, color=NAVY)

    def para(
        self,
        text: str,
        *,
        bold_prefix: str | None = None,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        color=DARK,
        size=13,
        after: float | None = None,
        keep=False,
    ) -> None:
        p = self.doc.add_paragraph()
        set_paragraph_rtl(p, align=align)
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if bold_prefix and text.startswith(bold_prefix):
            r1 = p.add_run(bold_prefix)
            set_run_font(r1, size, bold=True, color=color)
            r2 = p.add_run(text[len(bold_prefix) :])
            set_run_font(r2, size, color=color)
        else:
            r = p.add_run(text)
            set_run_font(r, size, color=color)
        set_keep(p, lines=keep)

    def bullet(self, text: str, *, color=DARK, size=13) -> None:
        p = self.doc.add_paragraph()
        set_paragraph_rtl(p)
        apply_num(p, self.bullet_num_id)
        r = p.add_run(text)
        set_run_font(r, size, color=color)
        set_keep(p, lines=True)

    def numbered(self, text: str, *, color=DARK, size=13) -> None:
        p = self.doc.add_paragraph()
        set_paragraph_rtl(p)
        apply_num(p, self.decimal_num_id)
        r = p.add_run(text)
        set_run_font(r, size, color=color)
        set_keep(p, lines=True)

    def callout(self, label: str, text: str, *, fill=PALE_TEAL, color=NAVY) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        set_cell_shading(cell, fill)
        set_cell_border(cell, color="B8C8C1", size="5")
        p = cell.paragraphs[0]
        set_paragraph_rtl(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.space_after = Pt(0)
        r1 = p.add_run(f"{label}: ")
        set_run_font(r1, 13, bold=True, color=color)
        r2 = p.add_run(text)
        set_run_font(r2, 13, color=DARK)
        width = section_content_width_dxa(self.doc.sections[0])
        apply_table_geometry(
            table,
            [width],
            table_width_dxa=width,
            indent_dxa=180,
            cell_margins_dxa={"top": 140, "bottom": 140, "start": 180, "end": 180},
        )
        self.doc.add_paragraph().paragraph_format.space_after = Pt(1)

    def table(
        self,
        headers: list[str],
        rows: list[list[str]],
        weights: list[float],
        *,
        status_colors: dict[str, str] | None = None,
        font_size=11.5,
    ) -> None:
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.autofit = False
        table_pr = table._tbl.tblPr
        bidi_visual = OxmlElement("w:bidiVisual")
        table_pr.append(bidi_visual)
        hdr = table.rows[0]
        set_repeat_table_header(hdr)
        for i, text in enumerate(headers):
            cell = hdr.cells[i]
            set_cell_shading(cell, NAVY)
            set_cell_border(cell, color="9AA8B5", size="5")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            set_paragraph_rtl(p, align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font(r, font_size, bold=True, color=WHITE)

        for row_idx, values in enumerate(rows):
            cells = table.add_row().cells
            set_row_no_split(table.rows[-1])
            for i, text in enumerate(values):
                cell = cells[i]
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_border(cell, color="CBD5E1", size="4")
                set_cell_shading(cell, WHITE if row_idx % 2 == 0 else "FAFBFC")
                if status_colors and text in status_colors:
                    set_cell_shading(cell, status_colors[text])
                p = cell.paragraphs[0]
                set_paragraph_rtl(
                    p,
                    align=WD_ALIGN_PARAGRAPH.CENTER if i == len(values) - 1 else WD_ALIGN_PARAGRAPH.RIGHT,
                )
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(text)
                set_run_font(r, font_size, bold=bool(status_colors and text in status_colors), color=DARK)

        width = section_content_width_dxa(self.doc.sections[0])
        raw = [int(width * w / sum(weights)) for w in weights]
        raw[-1] += width - sum(raw)
        apply_table_geometry(
            table,
            raw,
            table_width_dxa=width,
            indent_dxa=TABLE_INDENT_DXA,
            cell_margins_dxa=CELL_MARGINS_DXA,
        )
        self.doc.add_paragraph().paragraph_format.space_after = Pt(1)

    def page_break(self) -> None:
        self.doc.add_page_break()

    def save(self, path: Path, *, title: str, subject: str) -> None:
        self.doc.core_properties.title = title
        self.doc.core_properties.subject = subject
        self.doc.core_properties.author = "تیم توسعه سامانه"
        self.doc.core_properties.keywords = "سامانه مدیریت میدان، کالک، LoRaWAN، رصد نیرو"
        self.doc.core_properties.comments = "نام پروژه موقت است و پس از توافق تعیین می‌شود."
        for paragraph in self.doc.paragraphs:
            set_paragraph_rtl(
                paragraph,
                align=paragraph.alignment if paragraph.alignment is not None else WD_ALIGN_PARAGRAPH.RIGHT,
            )
            for run in paragraph.runs:
                set_run_font(run)
        for table in self.doc.tables:
            if table._tbl.tblPr.find(qn("w:bidiVisual")) is None:
                table._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        set_paragraph_rtl(
                            paragraph,
                            align=paragraph.alignment
                            if paragraph.alignment is not None
                            else WD_ALIGN_PARAGRAPH.RIGHT,
                        )
                        for run in paragraph.runs:
                            set_run_font(run)
        for section in self.doc.sections:
            for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
                for paragraph in part.paragraphs:
                    set_paragraph_rtl(
                        paragraph,
                        align=paragraph.alignment
                        if paragraph.alignment is not None
                        else WD_ALIGN_PARAGRAPH.RIGHT,
                    )
                    for run in paragraph.runs:
                        set_run_font(run)
        self.doc.save(path)


def build_proposal() -> None:
    d = PersianDoc(
        preset="narrative_proposal",
        running_title="پیشنهاد فنی | سامانه مدیریت و رصد میدان",
    )
    d.title(
        "پروپوزال سامانه مدیریت و رصد میدان",
        "پیشنهاد فنی اولیه برای فرماندهی، تبادل کالک و پایش نیروهای حاضر در میدان",
        meta=[
            "مخاطب: مدیران و فرماندهان عملیاتی",
            "ارائه‌دهنده: تیم توسعه سامانه",
            "نسخه اولیه - مرداد ۱۴۰۵",
        ],
    )
    d.callout(
        "موضوع پیشنهاد",
        "توسعه یک سامانه یکپارچه که فرمانده را قادر می‌سازد کالک و دستور عملیاتی را به رده‌های پایین ارسال کند، دریافت و اجرای آن را پیگیری کند و موقعیت و وضعیت نیروهای میدان را روی همان نقشه مشاهده نماید.",
        fill=PALE_BLUE,
    )
    d.para(
        "این سند بر مبنای اطلاعات فعلی تهیه شده است. مدل نهایی تجهیزات، تعداد کاربران، ساختار پوشش رادیویی، نرخ ارسال داده و الزامات امنیتی در جلسه شناخت نیازمندی‌ها قطعی خواهد شد.",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=MID_GRAY,
        size=11.5,
        after=2,
    )

    d.page_break()
    d.h1("۱. خلاصه پیشنهاد")
    d.para(
        "نیاز مطرح‌شده، ایجاد یک چرخه کامل میان مرکز فرماندهی و نیروهای حاضر در میدان است: فرمانده کالک و مأموریت را آماده می‌کند، اطلاعات مرتبط با هر رده به تبلت‌های میدان می‌رسد، افراد دریافت و اجرای دستور را اعلام می‌کنند و هم‌زمان موقعیت و علائم حیاتی آنان از ساعت هوشمند به مرکز بازمی‌گردد."
    )
    d.para(
        "هسته سمت فرماندهی از صفر آغاز نمی‌شود. سامانه موجود در حال حاضر نقشه عملیاتی، کالک، علائم نظامی، ساختار یگان‌ها، موقعیت زمان‌دار، مراحل و وظایف، کاربران، سوابق تغییرات و نقشه آفلاین را در اختیار دارد. توسعه اصلی باقی‌مانده مربوط به کلاینت میدانی، اتصال LoRaWAN، مدیریت تجهیزات، تله‌متری، ارسال مطمئن دستور و هشدارهای عملیاتی است."
    )
    d.callout(
        "جمع‌بندی مدیریتی",
        "پایه نرم‌افزاری فرماندهی و کالک آماده است. پروژه پیشنهادی این پایه را به یک سامانه عملیاتی دوطرفه میان فرمانده و میدان تبدیل می‌کند؛ بدون آنکه اجزای موجود دوباره ساخته شوند.",
    )

    d.h1("۲. برداشت ما از نیاز عملیاتی")
    d.para(
        "سامانه باید به فرمانده رده بالا تصویری واحد از وضعیت مأموریت بدهد. در این تصویر، کالک ابلاغ‌شده، یگان‌ها، افراد، مسیر حرکت، وضعیت ارتباط و هشدارهای مهم در یک محیط نقشه‌ای دیده می‌شوند. رده پایین نیز باید فقط اطلاعات موردنیاز خود را دریافت کند و بتواند با کمترین عملیات، دریافت دستور و وضعیت اجرا را گزارش دهد."
    )
    d.h2("اهداف اصلی")
    for item in [
        "ارسال کالک، دستور، نقطه، مسیر و محدوده عملیاتی از مرکز به رده‌های پایین.",
        "نمایش هر فرد میدان به‌صورت یک آیکون وابسته به یگان و مأموریت او.",
        "دریافت مستقیم موقعیت و علائم حیاتی از ساعت هوشمند دارای LoRaWAN.",
        "دریافت دستورات روی تبلت از طریق مودم LoRaWAN متصل به درگاه OTG.",
        "ثبت تأیید دریافت، مشاهده، شروع اجرا و پایان مأموریت.",
        "ادامه کار در شرایط قطع ارتباط و ارسال مجدد داده پس از برقراری لینک.",
        "ایجاد هشدار برای قطع ارتباط، خروج از محدوده و وضعیت‌های تعریف‌شده.",
    ]:
        d.bullet(item)

    d.h1("۳. جریان عملیاتی پیشنهادی")
    steps = [
        ("۱", "آماده‌سازی عملیات", "فرمانده منطقه، کالک، یگان‌ها، مأموریت‌ها و سطوح دسترسی را در سامانه تعریف می‌کند."),
        ("۲", "آماده‌سازی تبلت", "نقشه پایه و کالک اولیه پیش از عملیات روی تبلت قرار می‌گیرد تا کار وابسته به انتقال فایل حجیم نباشد."),
        ("۳", "ابلاغ مأموریت", "دستورها و تغییرات کم‌حجم از سرور عملیاتی، Network Server و Gateway به مودم OTG تبلت می‌رسند."),
        ("۴", "گزارش میدان", "کاربر دریافت دستور، وضعیت اجرا، گزارش کوتاه و درخواست کمک را از تبلت ارسال می‌کند."),
        ("۵", "پایش فرد", "ساعت هوشمند به‌صورت مستقل موقعیت، علائم حیاتی، زمان ثبت و وضعیت باتری را از طریق LoRaWAN می‌فرستد."),
        ("۶", "نمایش و هشدار", "سامانه داده‌ها را به پرونده فرد متصل می‌کند، آیکون او را به‌روزرسانی می‌کند و قواعد هشدار را اعمال می‌نماید."),
    ]
    d.table(
        ["گام", "عنوان", "شرح عملیات"],
        [[a, b, c] for a, b, c in steps],
        [0.65, 1.5, 4.35],
        font_size=11.5,
    )

    d.h1("۴. معماری مفهومی سامانه")
    d.para(
        "معماری پیشنهادی از چند لایه مستقل تشکیل می‌شود تا نقشه و فرماندهی، ارتباط رادیویی و تجهیزات میدان به یکدیگر وابستگی مستقیم و شکننده نداشته باشند."
    )
    architecture_rows = [
        ["مرکز فرماندهی", "کالک، مأموریت، نقشه، نمایش افراد، هشدارها و گزارش وضعیت"],
        ["سرور عملیاتی", "مدیریت کاربران، مأموریت‌ها، تجهیزات، پیام‌ها، سوابق و API"],
        ["درگاه یکپارچه‌سازی", "تبدیل پیام‌های سامانه به Payload و دریافت داده از Network Server"],
        ["LoRaWAN Network Server", "مدیریت دستگاه‌ها، Join، مسیر پیام، کلیدها و ارتباط Gatewayها"],
        ["Gatewayهای میدان", "تبادل رادیویی با ساعت‌ها و مودم‌های متصل به تبلت"],
        ["ساعت هوشمند", "ارسال مستقیم موقعیت، علائم حیاتی و وضعیت دستگاه"],
        ["تبلت و مودم OTG", "دریافت کالک و دستور، نمایش نقشه آفلاین و ارسال گزارش و تأیید"],
    ]
    d.table(["لایه", "مسئولیت"], architecture_rows, [1.85, 4.65], font_size=11.5)
    d.callout(
        "نکته معماری",
        "ساعت و تبلت دو دستگاه مستقل‌اند، اما هر دو به پرونده یک فرد متصل می‌شوند. آیکون روی نقشه نماینده فرد است؛ نه نماینده ساعت یا مودم.",
        fill=PALE_GOLD,
    )

    d.h1("۵. قابلیت‌های آماده در سامانه موجود")
    d.para(
        "بخش قابل توجهی از نیازهای سمت مرکز در نسخه موجود پیاده شده و قابل نمایش است. این موضوع باعث می‌شود توسعه پیشنهادی روی اتصال میدان و تکمیل گردش عملیات متمرکز شود."
    )
    ready_rows = [
        ["نقشه و کالک عملیاتی", "موجود", "ترسیم و مدیریت اطلاعات مکانی روی نقشه"],
        ["نقشه آفلاین", "موجود", "استفاده بدون وابستگی دائمی به اینترنت"],
        ["علائم نظامی", "موجود", "پشتیبانی از MIL-STD-2525 و APP-6"],
        ["ساختار یگان‌ها", "موجود", "تعریف سلسله‌مراتب، زیرمجموعه‌ها و منابع"],
        ["موقعیت زمان‌دار", "موجود", "ثبت و نمایش تغییر موقعیت یگان در طول زمان"],
        ["مرحله و وظیفه", "موجود", "تعریف مراحل، وظایف و واحدهای مسئول"],
        ["کاربران و دسترسی", "موجود", "نقش‌های کاربری و محدودسازی عملیات"],
        ["سوابق تغییرات", "موجود", "ثبت تاریخچه سناریو و فعالیت‌های اصلی"],
        ["API و ارتباط بلادرنگ", "پایه موجود", "نیازمند تقویت برای تله‌متری پیوسته و مقیاس عملیاتی"],
        ["مدیریت افراد و تجهیزات", "پایه موجود", "نیازمند اتصال هر فرد به ساعت، تبلت و DevEUI"],
    ]
    d.table(
        ["قابلیت", "وضعیت", "کاربرد در پروژه جدید"],
        ready_rows,
        [1.9, 1.15, 3.45],
        status_colors={"موجود": "DDEFE5", "پایه موجود": "FFF1CC"},
        font_size=11,
    )
    d.para(
        "در بیان میزان آمادگی، تفکیک بخش مرکز و میدان ضروری است. بخش اصلی کالک و فرماندهی آماده است؛ اما کلاینت میدان و اتصال تجهیزات باید توسعه داده شوند. از اعلام یک درصد کلی بدون این تفکیک پرهیز می‌شود.",
        color=MID_GRAY,
        size=11.5,
    )

    d.h1("۶. اجزای موردنیاز برای توسعه")
    remaining_rows = [
        ["کلاینت میدانی تبلت", "برنامه نصب‌شونده با نقشه آفلاین، نمایش مأموریت، گزارش و تأیید دریافت"],
        ["رابط USB OTG", "شناسایی مودم، خواندن و نوشتن پیام و کنترل وضعیت اتصال"],
        ["درگاه LoRaWAN", "اتصال Network Server به سامانه از طریق API یا MQTT و مدیریت Codec پیام‌ها"],
        ["ثبت تجهیزات", "نگهداری DevEUI، JoinEUI، نوع دستگاه، مالک، وضعیت و زمان آخرین ارتباط"],
        ["پایگاه تله‌متری", "ذخیره جداگانه موقعیت، علائم حیاتی، باتری و کیفیت ارتباط"],
        ["نمایش زنده فرد", "به‌روزرسانی آیکون، مسیر حرکت، وضعیت ارتباط و آخرین داده معتبر"],
        ["مدیریت پیام", "اولویت، شناسه یکتا، تأیید، تلاش مجدد و جلوگیری از پیام تکراری"],
        ["موتور هشدار", "قواعد خروج از محدوده، قطع ارتباط، بی‌حرکتی و مقادیر اعلامی کارفرما"],
        ["مدیریت امنیت", "کلیدهای دستگاه، ابطال تجهیز، دسترسی کاربران و ثبت رویداد امنیتی"],
    ]
    d.table(["زیرسامانه", "شرح"], remaining_rows, [2.0, 4.5], font_size=11.2)

    d.h1("۷. راهبرد تبادل اطلاعات روی LoRaWAN")
    d.para(
        "LoRaWAN برای پیام‌های کوتاه و دوره‌ای طراحی شده است. بنابراین نقشه پایه، فایل‌های حجیم و کالک کامل نباید هنگام عملیات از این بستر منتقل شوند. این اطلاعات پیش از عملیات روی تبلت قرار می‌گیرند و شبکه رادیویی فقط تغییرات و پیام‌های کم‌حجم را جابه‌جا می‌کند."
    )
    d.h2("پیام‌های مناسب برای این بستر")
    for item in [
        "موقعیت، زمان ثبت، باتری و علائم حیاتی ساعت.",
        "دستور متنی کوتاه و کد مأموریت.",
        "نقطه، مسیر کوتاه یا محدوده فشرده‌شده.",
        "هشدار و اولویت پیام.",
        "تأیید دریافت، مشاهده و وضعیت اجرای مأموریت.",
        "گزارش کوتاه میدانی و درخواست کمک.",
    ]:
        d.bullet(item)
    d.h2("کنترل تحویل پیام")
    d.para(
        "هر پیام عملیاتی باید شناسه، زمان، فرستنده، گیرنده، اولویت و وضعیت تحویل داشته باشد. تبلت پیام را در حافظه محلی ثبت می‌کند و پس از نمایش یا اقدام کاربر، تأیید متناسب را می‌فرستد. در صورت قطع ارتباط، پیام خروجی در صف باقی می‌ماند. سیاست تلاش مجدد و انقضا بر اساس نوع و اولویت پیام تنظیم خواهد شد."
    )
    d.callout(
        "محدودیت آگاهانه",
        "تصویر، ویدئو، فایل بزرگ و نقشه کامل از مسیر LoRaWAN ارسال نمی‌شود. در صورت نیاز به این موارد، باید یک مسیر ارتباطی پرظرفیت یا فرایند بارگذاری پیش از عملیات تعریف شود.",
        fill="FBECEC",
        color=RED,
    )

    d.h1("۸. امنیت، پایداری و ثبت سوابق")
    for item in [
        "استفاده از فعال‌سازی و کلیدهای یکتای هر دستگاه و نگهداری کنترل‌شده اطلاعات امنیتی.",
        "اتصال صریح هر DevEUI به یک فرد و جلوگیری از پذیرش دستگاه ثبت‌نشده.",
        "امکان غیرفعال‌کردن ساعت، تبلت یا مودم مفقودشده.",
        "تفکیک دسترسی فرمانده، اپراتور و نیروی میدان بر اساس رده و مأموریت.",
        "ثبت زمان تولید، دریافت، نمایش و تأیید هر دستور.",
        "تشخیص داده قدیمی، تکراری یا خارج از ترتیب.",
        "ذخیره محلی روی تبلت برای ادامه کار هنگام قطع ارتباط.",
        "پایش آخرین ارتباط Gateway و دستگاه‌ها و نمایش وضعیت نامطمئن به فرمانده.",
    ]:
        d.bullet(item)
    d.para(
        "حدود هشدارهای پزشکی و تفسیر علائم حیاتی باید توسط کارفرما یا مرجع پزشکی مورد تأیید او تعیین شود. سامانه این قواعد را اجرا و ثبت می‌کند، اما تعیین بالینی مقادیر بر عهده تیم نرم‌افزار نیست.",
        color=MID_GRAY,
        size=11.5,
    )

    d.h1("۹. روش پیشنهادی اجرای پروژه")
    phases = [
        ["شناخت و تثبیت نیاز", "بررسی تجهیزات، جریان فرماندهی، پیام‌ها، تعداد کاربران، پوشش و معیارهای پذیرش"],
        ["نمونه اتصال", "دریافت Payload واقعی ساعت و تبادل پیام با مودم OTG و Gateway"],
        ["توسعه زیرسامانه میدان", "ساخت کلاینت تبلت، صف آفلاین، نمایش مأموریت و تأیید دریافت"],
        ["توسعه پایش مرکز", "مدیریت دستگاه، تله‌متری، نمایش آیکون و موتور هشدار"],
        ["یکپارچه‌سازی", "اتصال کامل ساعت، تبلت، Network Server و سامانه فرماندهی"],
        ["آزمون میدانی", "بررسی پوشش، تأخیر، مصرف باتری، قطع و وصل و رفتار سامانه در شرایط واقعی"],
        ["تحویل و آموزش", "مستندات، آموزش کاربران، رفع اشکال پذیرش و تحویل نسخه مصوب"],
    ]
    d.table(["مرحله", "خروجی اصلی"], phases, [1.85, 4.65], font_size=11.3)
    d.para(
        "زمان‌بندی و برآورد مالی پس از دریافت مشخصات تجهیزات و تثبیت دامنه هر مرحله ارائه خواهد شد؛ زیرا مدل ساعت، مودم، Gateway، Network Server و نرخ مورد انتظار ارسال داده مستقیماً بر حجم کار اثر می‌گذارند.",
        color=MID_GRAY,
        size=11.5,
    )

    d.h1("۱۰. اطلاعات موردنیاز از کارفرما")
    requirement_rows = [
        ["ساعت هوشمند", "مدل، سازنده، Firmware، نسخه LoRaWAN، Codec، حسگرها و نمونه Payload"],
        ["تبلت", "مدل، سیستم‌عامل، ظرفیت ذخیره‌سازی، GPS و سیاست نصب برنامه"],
        ["مودم OTG", "مدل، درایور، پروتکل USB/Serial، دستورات، کلاس LoRaWAN و مصرف برق"],
        ["Gateway", "مدل، تعداد، محل نصب، آنتن، فرکانس و نوع Backhaul"],
        ["Network Server", "محصول انتخاب‌شده، محل استقرار، API/MQTT و نحوه مدیریت کلیدها"],
        ["مقیاس عملیات", "تعداد نفرات، تعداد یگان‌ها، وسعت منطقه و تعداد کاربران مرکز"],
        ["الگوی ارسال", "فاصله ارسال موقعیت و علائم حیاتی، تأخیر قابل قبول و اولویت پیام‌ها"],
        ["گردش فرماندهی", "رده‌ها، گیرندگان مجاز، انواع دستور و مراحل تأیید و اجرا"],
        ["هشدارها", "قواعد محدوده، قطع ارتباط، بی‌حرکتی و حدود مورد تأیید علائم حیاتی"],
        ["امنیت و استقرار", "سطح محرمانگی، شبکه داخلی، نگهداری داده، نسخه پشتیبان و ثبت رویداد"],
    ]
    d.table(["موضوع", "اطلاعات موردنیاز"], requirement_rows, [1.75, 4.75], font_size=11.2)

    d.h1("۱۱. خروجی‌های پیشنهادی")
    for item in [
        "نسخه تکمیل‌شده سامانه فرماندهی و مدیریت کالک.",
        "کلاینت نصب‌شونده تبلت میدان با نقشه آفلاین.",
        "ماژول ارتباط OTG با مودم مورد تأیید.",
        "درگاه یکپارچه‌سازی LoRaWAN و Codec پیام‌های ساعت و تبلت.",
        "مدیریت افراد، تجهیزات، شناسه‌ها و وضعیت ارتباط.",
        "داشبورد پایش زنده، مسیر حرکت و هشدارها.",
        "گردش ارسال دستور، تأیید دریافت و وضعیت اجرا.",
        "ثبت سوابق پیام و تله‌متری.",
        "مستندات نصب، پیکربندی، کاربری و پشتیبانی.",
        "آموزش کاربران مرکز و کاربران منتخب میدان.",
    ]:
        d.bullet(item)

    d.h1("۱۲. معیارهای پذیرش اولیه")
    for item in [
        "یک ساعت ثبت‌شده بتواند Payload واقعی موقعیت و علائم حیاتی را به سامانه برساند.",
        "سامانه داده ساعت را به فرد صحیح متصل و آیکون او را روی نقشه به‌روزرسانی کند.",
        "فرمانده بتواند یک دستور یا تغییر کالک را برای فرد یا یگان منتخب ارسال کند.",
        "تبلت پیام را از مودم OTG دریافت و روی نقشه یا صفحه مأموریت نمایش دهد.",
        "وضعیت دریافت و اقدام کاربر از تبلت به مرکز بازگردد.",
        "در قطع ارتباط، پیام خروجی تبلت از بین نرود و پس از اتصال ارسال شود.",
        "داده تکراری یا قدیمی باعث ایجاد موقعیت یا هشدار اشتباه نشود.",
        "کاربر فقط افراد و مأموریت‌های مجاز در رده خود را مشاهده کند.",
    ]:
        d.bullet(item)

    d.h1("۱۳. فرض‌ها و مرزهای پیشنهاد")
    for item in [
        "ساعت هوشمند امکان ارسال مستقیم داده از طریق LoRaWAN را دارد.",
        "تبلت از طریق OTG به مودم LoRaWAN متصل می‌شود.",
        "مودم، ساعت و Gateway دارای مستندات فنی و امکان یکپارچه‌سازی هستند.",
        "کارفرما فرکانس، پوشش، مجوزهای رادیویی و زیرساخت Gateway را تعیین یا تأمین می‌کند.",
        "نوع ارتباط Gateway با مرکز در مرحله شناخت مشخص خواهد شد؛ معماری می‌تواند مرکزی یا محلی/لبه‌ای باشد.",
        "ارسال تصویر، ویدئو و فایل حجیم در دامنه LoRaWAN این پیشنهاد قرار ندارد.",
        "بخش شبیه‌ساز در محدوده موردنظر کارفرما نیست.",
        "نام «سامانه مدیریت و رصد میدان» موقت است.",
    ]:
        d.bullet(item)

    d.h1("۱۴. جمع‌بندی")
    d.para(
        "پیشنهاد حاضر بر یک محصول موجود متکی است که بخش‌های اصلی نقشه، کالک، نمادهای نظامی، ساختار یگان‌ها، موقعیت زمان‌دار، مراحل عملیات، کاربران و نقشه آفلاین را پوشش می‌دهد. توسعه پروژه به‌جای بازسازی این اجزا، بر ایجاد ارتباط دوطرفه با میدان و تبدیل داده تجهیزات به اطلاعات قابل استفاده برای فرمانده متمرکز خواهد شد."
    )
    d.para(
        "پس از دریافت مشخصات واقعی ساعت، مودم OTG، Gateway و Network Server، یک نمونه اتصال ساخته می‌شود تا Payload، تأخیر، مصرف انرژی و چرخه تحویل پیام پیش از تثبیت طراحی نهایی آزموده شود. این مرحله مبنای نهایی‌کردن دامنه، زمان و هزینه خواهد بود."
    )
    d.callout(
        "پیشنهاد اقدام بعدی",
        "برگزاری جلسه مشترک عملیاتی و فنی، دریافت نمونه تجهیزات و مستندات، و اجرای یک نمایش محدود شامل ارسال موقعیت ساعت، نمایش آیکون فرد و ارسال یک دستور به تبلت.",
        fill=PALE_BLUE,
    )

    d.save(
        PROPOSAL_PATH,
        title="پروپوزال سامانه مدیریت و رصد میدان",
        subject="پیشنهاد فنی اولیه برای فرماندهی، تبادل کالک و پایش نیروهای میدان",
    )


def build_speaker_guide() -> None:
    d = PersianDoc(
        preset="compact_reference_guide",
        running_title="راهنمای جلسه | سامانه مدیریت و رصد میدان",
    )
    d.title(
        "متن مطالعه و ارائه",
        "راهنمای دفاع از پیشنهاد سامانه مدیریت و رصد میدان",
        meta=[
            "ویژه جلسه با مدیران و فرماندهان عملیاتی",
            "نام پروژه موقت است",
            "نسخه اولیه - مرداد ۱۴۰۵",
        ],
    )
    d.callout(
        "پیام اصلی جلسه",
        "ما برای ساخت یک نقشه نمایشی نیامده‌ایم. بخش اصلی نقشه، کالک و ساختار فرماندهی را در اختیار داریم و دقیقاً می‌دانیم چگونه ساعت، تبلت و LoRaWAN را به یک چرخه عملیاتی قابل پیگیری وصل کنیم.",
        fill=PALE_BLUE,
    )
    d.para(
        "این متن برای روخوانی خشک نوشته نشده است. یک‌بار کامل مطالعه شود و در جلسه با لحن طبیعی، همراه با نمایش نسخه موجود و مکث برای پرسش‌ها ارائه گردد.",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=MID_GRAY,
        size=11.5,
    )

    d.page_break()
    d.h1("۱. هدف شما در جلسه")
    d.para(
        "قرار نیست در جلسه همه جزئیات فنی را توضیح دهید. باید سه نکته برای کارفرما روشن شود:"
    )
    for item in [
        "جریان عملیات و مسئله فرمانده را درست فهمیده‌اید.",
        "بخش اصلی نرم‌افزار سمت مرکز و کالک هم‌اکنون وجود دارد و قابل نمایش است.",
        "برای اتصال میدان، تجهیزات و LoRaWAN مسیر اجرایی مشخص دارید و محدودیت‌های آن را می‌شناسید.",
    ]:
        d.bullet(item)
    d.callout(
        "موضع مناسب",
        "نگویید کل پروژه آماده است. بگویید هسته فرماندهی و کالک آماده است و کار باقی‌مانده روی اتصال میدان، تجهیزات و تبادل مطمئن داده متمرکز خواهد بود.",
        fill=PALE_GOLD,
    )

    d.h1("۲. متن پیشنهادی ارائه")
    d.h2("شروع جلسه - حدود یک دقیقه")
    d.para(
        "آنچه از نیاز شما برداشت کرده‌ایم، ایجاد یک ارتباط دوطرفه میان فرمانده و نیروهای حاضر در میدان است. فرمانده باید بتواند کالک و مأموریت را آماده و برای رده‌های پایین ارسال کند. از طرف دیگر، وضعیت نیروها فقط به‌صورت گزارش شفاهی یا دستی دریافت نشود؛ موقعیت و وضعیت هر فرد باید روی همان نقشه و در ارتباط با مأموریت او دیده شود."
    )
    d.para(
        "برای این منظور، ساعت هوشمند مستقیماً اطلاعات موقعیت و علائم حیاتی را از طریق LoRaWAN ارسال می‌کند. تبلت نیز با مودم متصل از طریق OTG، ابزار دریافت کالک و دستور و ارسال گزارش میدان خواهد بود."
    )

    d.h2("توضیح جریان عملیات - حدود دو دقیقه")
    d.para(
        "جریان کار را به این صورت در نظر گرفته‌ایم: پیش از عملیات، فرمانده منطقه، کالک، یگان‌ها و مأموریت‌ها را در سامانه تعریف می‌کند. نقشه پایه و کالک اولیه روی تبلت‌های مربوط قرار می‌گیرد. علت این کار آن است که در میدان به انتقال فایل‌های حجیم وابسته نباشیم."
    )
    d.para(
        "هنگام عملیات، فرمانده تغییرات کم‌حجم مانند یک نقطه جدید، مسیر اصلاح‌شده، محدوده، هشدار یا دستور کوتاه را ارسال می‌کند. پیام از سرور عملیاتی وارد شبکه LoRaWAN می‌شود و از طریق Gateway به مودم OTG تبلت می‌رسد. نیروی میدان دریافت دستور، شروع اجرا و پایان کار را اعلام می‌کند."
    )
    d.para(
        "در مسیر برگشت، ساعت هر فرد مستقل از تبلت موقعیت، زمان، وضعیت باتری و علائم حیاتی قابل دسترس را ارسال می‌کند. سامانه این داده را به پرونده همان فرد متصل می‌کند و فرمانده به‌جای مشاهده یک داده خام، آیکون فرد، آخرین موقعیت، زمان ارتباط و وضعیت او را روی کالک می‌بیند."
    )

    d.h2("آنچه هم‌اکنون آماده است - حدود دو دقیقه")
    d.para(
        "نقطه شروع ما یک پروژه خام نیست. در نسخه موجود، نقشه عملیاتی، ابزارهای کالک، علائم نظامی، نقشه آفلاین، ساختار یگان‌ها، ثبت نفرات و تجهیزات، موقعیت زمان‌دار، مراحل عملیات، وظایف، کاربران و تاریخچه تغییرات وجود دارد."
    )
    d.para(
        "برای مثال، هم‌اکنون می‌توانیم یگان را روی نقشه قرار دهیم، نماد مناسب آن را نمایش دهیم، موقعیت را در زمان‌های مختلف ثبت کنیم، مرحله و وظیفه تعریف کنیم و اطلاعات را در شبکه داخلی نگهداری کنیم. این‌ها همان اجزایی هستند که داده میدان باید روی آن‌ها بنشیند."
    )
    d.para(
        "بنابراین برای بخش جدید لازم نیست نقشه، کالک، یگان و زیرساخت کاربران دوباره ساخته شوند. تمرکز توسعه روی دریافت داده ساعت، ارتباط OTG تبلت، مدیریت تجهیزات، نمایش زنده افراد، هشدار و گردش تأیید دستور خواهد بود."
    )

    d.h2("توضیح معماری بدون ورود به جزئیات اضافی - حدود دو دقیقه")
    d.para(
        "هر فرد در سامانه یک پرونده عملیاتی دارد. ساعت و تبلت دو تجهیز آن فرد هستند و هرکدام شناسه مستقل دارند. ساعت منبع اصلی موقعیت و علائم حیاتی است. تبلت ابزار دریافت مأموریت و ارسال گزارش است. آیکونی که فرمانده روی نقشه می‌بیند نماینده فرد است، نه نماینده دستگاه."
    )
    d.para(
        "ساعت پیام خود را به Gateway می‌فرستد. Gateway پیام را به Network Server می‌رساند و از آنجا داده وارد سامانه فرماندهی می‌شود. مسیر ارسال دستور به تبلت برعکس همین جریان است. در نرم‌افزار تبلت نیز یک بخش مخصوص برای ارتباط با مودم OTG ساخته می‌شود."
    )

    d.h2("پاسخ پیش‌دستانه درباره محدودیت LoRaWAN - حدود یک دقیقه")
    d.para(
        "در طراحی، قرار نیست نقشه کامل، تصویر یا فایل حجیم را از LoRaWAN عبور دهیم. نقشه و کالک اصلی پیش از عملیات روی تبلت قرار می‌گیرند. در میدان فقط تغییرات کم‌حجم، مختصات، دستور کوتاه، هشدار، تأیید دریافت و اطلاعات ساعت منتقل می‌شوند. به این ترتیب شبکه برای کاری استفاده می‌شود که با ظرفیت آن سازگار است."
    )

    d.h2("قسمت‌های باقی‌مانده - حدود یک دقیقه")
    d.para(
        "قسمت باقی‌مانده مشخص است: برنامه میدانی تبلت، رابط مودم OTG، اتصال Network Server، ثبت ساعت‌ها و مودم‌ها، پایگاه تله‌متری، نمایش زنده فرد، قواعد هشدار و کنترل تحویل پیام. این توسعه روی هسته فعلی انجام می‌شود و به بازنویسی محصول نیاز ندارد."
    )
    d.para(
        "برای قطعی‌کردن طراحی، باید مدل و مستندات ساعت، مودم، Gateway و Network Server، تعداد نیروها، فاصله ارسال موقعیت، سطح تأخیر قابل قبول و قواعد هشدار را از شما دریافت کنیم."
    )

    d.h2("پایان ارائه - حدود یک دقیقه")
    d.para(
        "پیشنهاد ما این است که جلسه بعدی با حضور نماینده عملیاتی و فنی برگزار شود و یک نمونه از ساعت، مودم OTG و مستندات تجهیزات در اختیار تیم قرار گیرد. نخست یک مسیر محدود را اجرا می‌کنیم: دریافت موقعیت یک ساعت، نمایش آیکون همان فرد روی نقشه و ارسال یک دستور کوتاه به تبلت. پس از تأیید این مسیر، دامنه کامل پروژه، زمان‌بندی و هزینه قطعی می‌شود."
    )
    d.para(
        "مزیت نقطه شروع فعلی این است که نتیجه نمونه اتصال مستقیماً در محیط واقعی کالک و فرماندهی دیده می‌شود؛ نه در یک نرم‌افزار آزمایشی جدا از محصول نهایی."
    )

    d.h1("۳. ترتیب پیشنهادی نمایش نسخه موجود")
    demo_steps = [
        ["۱", "بازکردن نقشه آفلاین", "نشان دهید سامانه بدون نقشه اینترنتی نیز قابل استفاده است."],
        ["۲", "ایجاد یا بازکردن کالک", "نقطه، خط، محدوده و لایه‌های عملیاتی را نمایش دهید."],
        ["۳", "نمایش علائم نظامی", "یک یگان یا فرد نمونه را با نماد مناسب روی نقشه قرار دهید."],
        ["۴", "نمایش ساختار یگان", "ارتباط رده بالا و زیرمجموعه‌ها را نشان دهید."],
        ["۵", "ثبت موقعیت", "توضیح دهید داده GPS ساعت همین موقعیت را به‌صورت خودکار به‌روزرسانی خواهد کرد."],
        ["۶", "مرحله و وظیفه", "یک مأموریت را به یگان نمونه منتسب کنید."],
        ["۷", "کاربران و سوابق", "سطوح دسترسی و ثبت تغییرات را به‌اختصار نشان دهید."],
    ]
    d.table(["ترتیب", "نمایش", "جمله‌ای که باید گفته شود"], demo_steps, [0.65, 1.55, 4.3], font_size=11.2)
    d.callout(
        "نکته نمایش",
        "قسمت شبیه‌ساز را نمایش ندهید و درباره آن صحبت نکنید، چون در محدوده موردنظر کارفرما نیست. نمایش را روی کالک، یگان، نقشه آفلاین و موقعیت متمرکز کنید.",
        fill=PALE_GOLD,
    )

    d.h1("۴. پاسخ به پرسش‌های محتمل")
    qa = [
        ["آیا سامانه کامل آماده است؟", "هسته فرماندهی، کالک و نقشه آماده است. اتصال مستقیم ساعت، تبلت و LoRaWAN و گردش کامل پیام باید روی این هسته توسعه داده شود."],
        ["چقدر از پروژه آماده است؟", "بخش اصلی سمت مرکز آماده است. اگر کل زنجیره مرکز، شبکه و میدان را یکجا حساب کنیم، برآورد فعلی حدود ۵۵ تا ۶۰ درصد است؛ عدد قطعی پس از بررسی تجهیزات اعلام می‌شود."],
        ["آیا کالک کامل با LoRaWAN ارسال می‌شود؟", "خیر. نقشه و کالک اولیه از قبل روی تبلت قرار می‌گیرند. هنگام عملیات فقط تغییرات کم‌حجم و دستورها ارسال می‌شوند."],
        ["اگر تبلت خاموش شود چه می‌شود؟", "ساعت مستقل است و همچنان می‌تواند موقعیت و علائم حیاتی را بفرستد. فقط دریافت دستور و گزارش کاربر تا بازگشت تبلت محدود می‌شود."],
        ["اگر ساعت قطع شود چه می‌شود؟", "سامانه آخرین زمان ارتباط را نشان می‌دهد و پس از عبور از آستانه تعیین‌شده هشدار قطع ارتباط تولید می‌کند."],
        ["مودم چگونه به تبلت وصل می‌شود؟", "از طریق OTG. برنامه تبلت یک ماژول مخصوص برای شناسایی و تبادل پیام با مودم خواهد داشت."],
        ["اگر اینترنت وجود نداشته باشد؟", "LoRaWAN به Gateway و Network Server نیاز دارد. با توجه به ساختار کارفرما، سرور می‌تواند در مرکز یا به‌صورت محلی در قرارگاه مستقر شود. نوع ارتباط Gateway باید در بررسی فنی مشخص شود."],
        ["امنیت پیام‌ها چگونه است؟", "علاوه بر سازوکار امنیتی LoRaWAN، مدیریت کلید یکتای دستگاه، کنترل دسترسی، ابطال تجهیز و ثبت رویداد در طراحی پیش‌بینی می‌شود. سطح نهایی امنیت طبق ضوابط کارفرما قطعی خواهد شد."],
        ["آیا علائم حیاتی قابل اعتماد است؟", "سامانه داده‌ای را که ساعت ارائه می‌کند ثبت و نمایش می‌دهد. مدل حسگر و دقت آن باید توسط کارفرما تأیید شود و حدود هشدار نیز باید از مرجع مورد تأیید کارفرما دریافت شود."],
        ["چه زمانی و با چه هزینه‌ای انجام می‌شود؟", "پس از دریافت مشخصات واقعی تجهیزات، تعداد نیروها و معیارهای پذیرش، زمان و هزینه مرحله‌بندی‌شده ارائه می‌شود. اعلام عدد پیش از بررسی سخت‌افزار می‌تواند گمراه‌کننده باشد."],
    ]
    d.table(["پرسش", "پاسخ پیشنهادی"], qa, [2.05, 4.45], font_size=10.8)

    d.h1("۵. پرسش‌هایی که شما باید از کارفرما بپرسید")
    for item in [
        "مدل دقیق ساعت و نمونه Payload واقعی آن چیست؟",
        "هر چند ثانیه یا دقیقه باید موقعیت و علائم حیاتی ارسال شوند؟",
        "تعداد نیروهای هم‌زمان و وسعت منطقه عملیات چقدر است؟",
        "مدل تبلت، سیستم‌عامل و مدل مودم OTG چیست؟",
        "Network Server انتخاب شده است یا باید پیشنهاد شود؟",
        "Gatewayها چگونه به مرکز متصل می‌شوند و در قطع ارتباط چه رفتاری انتظار می‌رود؟",
        "چه نوع دستورهایی باید به فرد و چه نوع دستورهایی باید به یگان ارسال شوند؟",
        "آیا تأیید دریافت کافی است یا مراحل مشاهده، شروع و پایان نیز لازم است؟",
        "چه کسی حدود هشدار علائم حیاتی را تأیید می‌کند؟",
        "داده‌ها چه مدت نگهداری می‌شوند و چه کسانی مجاز به مشاهده مسیر افراد هستند؟",
        "آیا تصویر یا فایل نیز باید در میدان ارسال شود؟",
        "معیار موفقیت آزمایش میدانی از نظر پوشش، تأخیر و نرخ دریافت پیام چیست؟",
    ]:
        d.bullet(item)

    d.h1("۶. عبارت‌های مناسب و نامناسب")
    phrase_rows = [
        ["نگویید", "کل پروژه آماده است.", "بگویید", "هسته فرماندهی و کالک آماده است و اتصال میدان روی آن توسعه داده می‌شود."],
        ["نگویید", "فقط چند تغییر کوچک لازم دارد.", "بگویید", "دو زیرسامانه مشخص، یعنی کلاینت میدان و درگاه LoRaWAN، به محصول اضافه می‌شود."],
        ["نگویید", "هر نوع داده‌ای روی LoRaWAN ارسال می‌کنیم.", "بگویید", "پیام‌های کم‌حجم و تغییرات کالک روی LoRaWAN منتقل می‌شوند."],
        ["نگویید", "دقت پزشکی ساعت تضمین شده است.", "بگویید", "دقت حسگر و حدود هشدار باید توسط مرجع مورد تأیید کارفرما مشخص شود."],
        ["نگویید", "در هر شرایطی ارتباط لحظه‌ای است.", "بگویید", "تأخیر و نرخ به‌روزرسانی بر اساس پوشش، کلاس دستگاه و ظرفیت شبکه آزموده می‌شود."],
    ]
    d.table(["نوع", "عبارت", "جایگزین", "بیان پیشنهادی"], phrase_rows, [0.75, 1.75, 0.85, 3.15], font_size=10.7)

    d.h1("۷. چک‌لیست پیش از جلسه")
    for item in [
        "نسخه نمایشی سامانه از قبل اجرا و آزمایش شده باشد.",
        "یک سناریوی کوتاه و مرتب برای نمایش آماده باشد.",
        "نقشه آفلاین و نمادها بدون خطا باز شوند.",
        "ترتیب نمایش بیشتر از ده دقیقه نشود.",
        "تصویر یا مشخصات ساعت، تبلت و مودم در صورت دسترسی همراه باشد.",
        "پروپوزال چاپی یا PDF روی یک حافظه جداگانه آماده باشد.",
        "از ورود به قیمت و زمان قطعی پیش از دریافت مشخصات تجهیزات خودداری شود.",
        "پرسش‌های فنی بدون پاسخ یادداشت شوند و پاسخ حدسی داده نشود.",
        "در پایان جلسه، درخواست نمونه تجهیزات و جلسه مشترک عملیاتی-فنی مطرح شود.",
    ]:
        d.bullet(item)

    d.h1("۸. جمع‌بندی کوتاه برای پایان جلسه")
    d.callout(
        "متن پیشنهادی",
        "ما مسئله را به دو بخش جدا نمی‌بینیم؛ کالک فرماندهی و وضعیت میدان باید در یک چرخه به هم متصل باشند. بخش اصلی محیط فرماندهی هم‌اکنون وجود دارد. با اتصال ساعت‌های LoRaWAN، مودم OTG تبلت و Network Server، همین محیط به سامانه مدیریت و رصد میدان تبدیل می‌شود. پیشنهاد ما این است که با تجهیزات واقعی یک نمونه محدود را اجرا کنیم و سپس دامنه نهایی را با معیارهای قابل اندازه‌گیری ببندیم.",
        fill=PALE_BLUE,
    )

    d.save(
        SPEAKER_PATH,
        title="متن مطالعه و ارائه سامانه مدیریت و رصد میدان",
        subject="راهنمای جلسه و دفاع از پیشنهاد برای مدیران و فرماندهان عملیاتی",
    )


def audit_document(path: Path) -> None:
    doc = Document(path)
    if not doc.paragraphs:
        raise RuntimeError(f"Empty document: {path}")
    for p in doc.paragraphs:
        p_pr = p._p.get_or_add_pPr()
        if p_pr.find(qn("w:bidi")) is None:
            raise RuntimeError(f"Non-RTL paragraph in {path}: {p.text[:50]}")
    for table in doc.tables:
        if table._tbl.tblPr.find(qn("w:bidiVisual")) is None:
            raise RuntimeError(f"Non-RTL table in {path}")
    with ZipFile(path) as archive:
        xml = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.endswith(".xml")
        )
    if FONT not in xml:
        raise RuntimeError(f"B Nazanin missing from {path}")
    for forbidden in ("ساجد", "سامانه ما", "TODO", "[[", "]]"):
        if forbidden in xml:
            raise RuntimeError(f"Forbidden placeholder/name {forbidden!r} in {path}")
    print("AUDIT_OK", path.name.encode("unicode_escape").decode("ascii"))


if __name__ == "__main__":
    build_proposal()
    build_speaker_guide()
    audit_document(PROPOSAL_PATH)
    audit_document(SPEAKER_PATH)
